"""Excel Suite definitions, validation, result aggregation, and persistence."""

from __future__ import annotations

import json
import os
import re
import shutil
from dataclasses import dataclass, field

import openpyxl


@dataclass
class Case:
    seq: int
    row: int
    scenario: str
    operation: str
    observe: str
    step_id: str = ""
    step_kind: str = "case"
    step_index: int | None = None
    required_flow: str = ""


@dataclass
class CaseResult:
    seq: int
    verdict: str = "needs_review"
    observed_effect: str = ""
    evidence: list = field(default_factory=list)
    cost_usd: float = 0.0
    tokens: dict = field(default_factory=dict)
    turns: int = 0
    duration_ms: int = 0
    api_ms: int = 0
    api_error_status: int | None = None
    agent_log: list = field(default_factory=list)


_FLOW_BINDING_HINT = re.compile(r"【\s*绑定流程")
_FLOW_BINDING_TAG = re.compile(r"【\s*绑定流程\s*[:：]\s*([^】]*?)\s*】")
_LOGIN_KW = ("帐号", "账号", "登录")
_NOTE_SHEET_KW = ("须知", "需知", "说明", "备注", "注意", "note")


def parse_bound_operation(value: object, row: int) -> tuple[str, str]:
    operation = str(value)
    hints = list(_FLOW_BINDING_HINT.finditer(operation))
    matches = list(_FLOW_BINDING_TAG.finditer(operation))
    if not hints:
        return operation, ""
    if len(hints) != 1 or len(matches) != 1:
        raise ValueError(f"第{row}行绑定流程标记无效或重复；应只写一个“【绑定流程:流程名】”")
    required_flow = matches[0].group(1).strip()
    if not required_flow:
        raise ValueError(f"第{row}行绑定流程标记缺少流程名")
    clean_operation = _FLOW_BINDING_TAG.sub("", operation, count=1).strip()
    if not clean_operation:
        raise ValueError(f"第{row}行绑定流程标记之外缺少操作描述")
    return clean_operation, required_flow


def read_cases(xlsx: str, sheet: str) -> list[Case]:
    worksheet = openpyxl.load_workbook(xlsx)[sheet]
    cases = []
    for row in range(2, worksheet.max_row + 1):
        seq = worksheet.cell(row, 1).value
        if seq is None or "要求" in str(seq):
            continue
        operation = worksheet.cell(row, 2).value
        if operation is None:
            continue
        clean_operation, required_flow = parse_bound_operation(operation, row)
        cases.append(Case(
            int(seq), row, sheet, clean_operation,
            str(worksheet.cell(row, 3).value or ""), required_flow=required_flow,
        ))
    return cases


def read_setup_steps(xlsx: str, sheet: str) -> list[tuple[str, str, str]]:
    worksheet = openpyxl.load_workbook(xlsx)[sheet]
    steps = []
    for row in range(2, worksheet.max_row + 1):
        marker = worksheet.cell(row, 1).value
        if marker is None or "要求" not in str(marker):
            continue
        operation = worksheet.cell(row, 2).value
        if operation in (None, ""):
            continue
        clean_operation, required_flow = parse_bound_operation(operation, row)
        steps.append((clean_operation, str(worksheet.cell(row, 3).value or ""), required_flow))
    return steps


def read_login_requirement(xlsx: str, sheet: str) -> str:
    worksheet = openpyxl.load_workbook(xlsx)[sheet]
    for row in range(2, worksheet.max_row + 1):
        marker = worksheet.cell(row, 1).value
        marker_text = str(marker) if marker is not None else ""
        if "要求" in marker_text and any(keyword in marker_text for keyword in _LOGIN_KW):
            operation = worksheet.cell(row, 2).value
            if operation not in (None, ""):
                return parse_bound_operation(operation, row)[0]
    return ""


def read_notes(xlsx: str, sheet: str) -> str:
    workbook = openpyxl.load_workbook(xlsx)
    fragments: list[str] = []

    def row_text(worksheet, row: int) -> str:
        values = [
            str(worksheet.cell(row, column).value).strip()
            for column in range(1, worksheet.max_column + 1)
            if worksheet.cell(row, column).value not in (None, "")
        ]
        return " ".join(values)

    if sheet in workbook.sheetnames:
        worksheet = workbook[sheet]
        for row in range(2, worksheet.max_row + 1):
            marker = worksheet.cell(row, 1).value
            try:
                int(marker)
                is_case = marker is not None
            except (TypeError, ValueError):
                is_case = False
            if is_case:
                continue
            marker_text = str(marker).strip() if marker is not None else ""
            if "要求" in marker_text:
                continue
            text = row_text(worksheet, row)
            if text:
                fragments.append(text)

    for name in workbook.sheetnames:
        if name == sheet or not any(keyword in name.lower() for keyword in _NOTE_SHEET_KW):
            continue
        worksheet = workbook[name]
        for row in range(1, worksheet.max_row + 1):
            text = row_text(worksheet, row)
            if text:
                fragments.append(text)
    return "\n".join(fragments).strip()


def validate_suite_flow_bindings(cases: list[Case], setup_steps: list, flows: dict) -> None:
    required = {case.required_flow for case in cases if case.required_flow}
    required.update(
        str(step[2]).strip() for step in setup_steps
        if len(step) >= 3 and str(step[2]).strip()
    )
    missing = sorted(name for name in required if name not in flows)
    if missing:
        raise ValueError(f"绑定流程标记引用不存在的流程: {missing}")


def validate_case_flow_bindings(cases: list[Case], flows: dict) -> None:
    validate_suite_flow_bindings(cases, [], flows)


def writeback(xlsx: str, sheet: str, results: list[CaseResult], cases: list[Case]) -> None:
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment
    from PIL import Image as PImage

    workbook = openpyxl.load_workbook(xlsx)
    worksheet = workbook[sheet]
    by_seq = {case.seq: case for case in cases}
    worksheet.column_dimensions["E"].width = 52
    worksheet.column_dimensions["F"].width = 24
    for result in results:
        row = by_seq[result.seq].row
        mark = {
            "pass": "✅", "fail": "❌", "blocked": "⏸️",
            "needs_review": "⚠️", "cancelled": "🚫",
        }.get(result.verdict, "")
        worksheet.cell(
            row, 5, f"{mark}{result.verdict} | {result.observed_effect}",
        ).alignment = Alignment(wrap_text=True, vertical="top")
        worksheet.row_dimensions[row].height = 110
        image = next((
            path for path in result.evidence
            if path.lower().endswith(".png") and os.path.exists(path) and os.path.getsize(path) > 0
        ), None)
        if image:
            try:
                thumbnail = image.replace(".png", "_th.png")
                source = PImage.open(image)
                height = 150
                source = source.resize((int(source.width * height / source.height), height))
                source.save(thumbnail)
                worksheet.add_image(XLImage(thumbnail), f"F{row}")
            except Exception:
                worksheet.cell(row, 6, os.path.basename(image))
        elif result.evidence:
            worksheet.cell(row, 6, " / ".join(result.evidence)[:60]).alignment = Alignment(
                wrap_text=True, vertical="top",
            )
    workbook.save(xlsx)


def aggregate_verdict(results: list[CaseResult], setup_verdicts=None) -> str:
    verdicts = [*(setup_verdicts or []), *(result.verdict for result in results)]
    return next(
        (verdict for verdict in ("fail", "blocked", "needs_review", "pass") if verdict in verdicts),
        "pass",
    )


def persist_results(out_dir: str, out_xlsx: str, xlsx: str, sheet: str, app: str,
                    results: list[CaseResult], cases: list[Case],
                    setup_verdicts: list[str] | None = None, *, final: bool = False) -> dict:
    shutil.copy(xlsx, out_xlsx)
    try:
        os.chmod(out_xlsx, os.stat(out_xlsx).st_mode | 0o200)
    except OSError:
        pass
    writeback(out_xlsx, sheet, results, cases)
    counts = {
        verdict: sum(1 for result in results if result.verdict == verdict)
        for verdict in ("pass", "fail", "blocked", "needs_review", "cancelled")
    }
    summary = {
        "sheet": sheet,
        "app": app,
        "cases": len(cases),
        "verdicts": counts,
        "setup_verdicts": list(setup_verdicts or []),
        "total_cost_usd": round(sum(result.cost_usd for result in results), 4),
        "tokens": {
            "input": sum(result.tokens.get("input_tokens", 0) for result in results),
            "output": sum(result.tokens.get("output_tokens", 0) for result in results),
            "cache_read": sum(result.tokens.get("cache_read_input_tokens", 0) for result in results),
        },
        "per_case": [{
            "seq": result.seq,
            "verdict": result.verdict,
            "cost_usd": round(result.cost_usd, 4),
            "turns": result.turns,
            "duration_s": round(result.duration_ms / 1000),
            "api_s": round(result.api_ms / 1000),
        } for result in results],
    }
    if final:
        summary["verdict"] = aggregate_verdict(results, setup_verdicts)
    with open(os.path.join(out_dir, "summary.json"), "w", encoding="utf-8") as handle:
        json.dump(summary, handle, ensure_ascii=False, indent=2)
    for result in results:
        with open(os.path.join(out_dir, f"case{result.seq}_log.txt"), "w", encoding="utf-8") as handle:
            handle.write("\n\n".join(result.agent_log))
    return summary
