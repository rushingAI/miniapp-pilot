"""Read complete Tencent Docs spreadsheets through the official MCP API."""
from __future__ import annotations

import io
import json
import re
from urllib.parse import urlsplit


MCP_URL = "https://docs.qq.com/openapi/mcp"
_DOC_ID = re.compile(r"^[A-Za-z0-9_-]{8,128}$")
_SHEET_INFO_TOOL = "sheet.get_sheet_info"
_CELL_DATA_TOOL = "sheet.get_cell_data"
_ROW_CHUNK_SIZE = 500


class TencentDocsError(ValueError):
    pass


def parse_sheet_url(url: str) -> tuple[str, str]:
    """Return (file_id, normalized_url) for a Tencent Docs sheet URL."""
    raw = str(url or "").strip()
    try:
        parsed = urlsplit(raw)
    except ValueError as exc:
        raise TencentDocsError("腾讯表格链接格式不正确") from exc
    if parsed.scheme != "https" or parsed.hostname != "docs.qq.com" or parsed.username:
        raise TencentDocsError("只支持 https://docs.qq.com/sheet/... 腾讯表格链接")
    parts = [part for part in parsed.path.split("/") if part]
    if len(parts) != 2 or parts[0] != "sheet" or not _DOC_ID.fullmatch(parts[1]):
        raise TencentDocsError("只支持 https://docs.qq.com/sheet/... 腾讯表格链接")
    file_id = parts[1]
    normalized = f"https://docs.qq.com/sheet/{file_id}"
    if parsed.query:
        normalized += f"?{parsed.query}"
    return file_id, normalized


def _tool_payload(result, operation: str) -> dict:
    """Return one MCP tool's structured payload without exposing cell contents in errors."""
    if getattr(result, "isError", False):
        raise TencentDocsError(f"腾讯文档{operation}失败")
    payload = getattr(result, "structuredContent", None)
    if not isinstance(payload, dict):
        texts = [
            getattr(block, "text", "")
            for block in getattr(result, "content", ())
            if getattr(block, "type", "") == "text"
        ]
        try:
            payload = json.loads("\n".join(texts))
        except (TypeError, json.JSONDecodeError) as exc:
            raise TencentDocsError(f"腾讯文档{operation}未返回可识别的结构化数据") from exc
    if not isinstance(payload, dict):
        raise TencentDocsError(f"腾讯文档{operation}未返回可识别的结构化数据")
    if payload.get("error"):
        raise TencentDocsError(f"腾讯文档{operation}失败")
    return payload


def _sheet_dimension(value, field: str) -> int:
    try:
        dimension = int(value)
    except (TypeError, ValueError) as exc:
        raise TencentDocsError(f"腾讯文档返回的工作表{field}无效") from exc
    if dimension < 0:
        raise TencentDocsError(f"腾讯文档返回的工作表{field}无效")
    return dimension


def _cell_value(cell: dict):
    value_type = str(cell.get("value_type") or "").upper()
    if value_type == "STRING":
        return cell.get("string_value")
    if value_type == "NUMBER":
        return cell.get("number_value")
    if value_type in {"BOOL", "BOOLEAN"}:
        return cell.get("bool_value")
    if value_type:
        string = cell.get("string_value")
        if string not in (None, ""):
            return string
    return None


async def _fetch_complete_workbook(session, file_id: str) -> bytes:
    info_result = await session.call_tool(_SHEET_INFO_TOOL, {"file_id": file_id})
    sheets = _tool_payload(info_result, "读取工作表信息").get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise TencentDocsError("腾讯表格中没有可读取的工作表")

    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        if not isinstance(sheet, dict):
            raise TencentDocsError("腾讯文档返回了无效的工作表信息")
        sheet_id = str(sheet.get("sheet_id") or "").strip()
        sheet_name = str(sheet.get("sheet_name") or "").strip()
        if not sheet_id or not sheet_name:
            raise TencentDocsError("腾讯文档返回了缺少标识或名称的工作表")
        row_count = _sheet_dimension(sheet.get("row_count"), "行数")
        col_count = _sheet_dimension(sheet.get("col_count"), "列数")
        worksheet = workbook.create_sheet(title=sheet_name[:31])
        if row_count == 0 or col_count == 0:
            continue

        for start_row in range(0, row_count, _ROW_CHUNK_SIZE):
            end_row = min(row_count - 1, start_row + _ROW_CHUNK_SIZE - 1)
            cell_result = await session.call_tool(_CELL_DATA_TOOL, {
                "file_id": file_id,
                "sheet_id": sheet_id,
                "start_row": start_row,
                "end_row": end_row,
                "start_col": 0,
                "end_col": col_count - 1,
                "return_csv": False,
                "include_formula": False,
            })
            cells = _tool_payload(cell_result, "读取单元格数据").get("cells")
            if not isinstance(cells, list):
                raise TencentDocsError("腾讯文档返回了无效的单元格数据")
            for cell in cells:
                if not isinstance(cell, dict):
                    raise TencentDocsError("腾讯文档返回了无效的单元格数据")
                try:
                    row = int(cell.get("row"))
                    col = int(cell.get("col"))
                except (TypeError, ValueError) as exc:
                    raise TencentDocsError("腾讯文档返回了无效的单元格坐标") from exc
                if not (start_row <= row <= end_row and 0 <= col < col_count):
                    raise TencentDocsError("腾讯文档返回了超出请求范围的单元格")
                value = _cell_value(cell)
                if value is not None:
                    worksheet.cell(row=row + 1, column=col + 1, value=value)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def fetch_workbook_bytes(url: str, token: str) -> tuple[bytes, dict]:
    """Fetch every worksheet cell through Tencent's range-based read-only tools."""
    file_id, normalized_url = parse_sheet_url(url)
    secret = str(token or "").strip()
    if not secret:
        raise TencentDocsError("请输入腾讯文档 MCP Token")

    import httpx
    from mcp import ClientSession
    from mcp.client.streamable_http import streamable_http_client

    try:
        async with httpx.AsyncClient(
                headers={"Authorization": secret}, timeout=httpx.Timeout(30.0)) as client:
            async with streamable_http_client(MCP_URL, http_client=client) as (read, write, _):
                async with ClientSession(read, write) as session:
                    await session.initialize()
                    tools = await session.list_tools()
                    names = {tool.name for tool in tools.tools}
                    missing = [
                        name for name in (_SHEET_INFO_TOOL, _CELL_DATA_TOOL)
                        if name not in names
                    ]
                    if missing:
                        raise TencentDocsError("当前腾讯文档接口未提供完整表格读取能力")
                    data = await _fetch_complete_workbook(session, file_id)
    except TencentDocsError as exc:
        raise TencentDocsError(str(exc).replace(secret, "***")) from exc
    except Exception as exc:
        message = str(exc).replace(secret, "***")
        raise TencentDocsError(f"连接腾讯文档失败：{message[:300]}") from exc
    return data, {"file_id": file_id, "url": normalized_url}
