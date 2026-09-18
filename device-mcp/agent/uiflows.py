"""读取 xlsx『流程映射』表，供 run_ui_flow 按阶段选择步骤脚本。

本模块只认识通用表结构与声明式匹配规则，不包含任何 app、页面或业务流程知识。
流程映射表结构：流程名｜阶段｜序｜匹配｜脚本名｜下一阶段｜失败后继续。
"""

SHEET_HINT = "流程映射"

_ALIASES = {
    "flow": ("流程名", "流程", "flow"),
    "stage": ("阶段", "stage"),
    "i": ("序", "序号", "优先级", "no"),
    "match": ("匹配", "条件", "match"),
    "script": ("脚本名", "脚本", "script"),
    "next_stage": ("下一阶段", "下阶段", "next"),
    "continue_on_fail": ("失败后继续", "失败继续", "continue"),
}


def _header_map(ws) -> dict:
    out = {}
    for c in range(1, ws.max_column + 1):
        raw = str(ws.cell(1, c).value or "").strip()
        if not raw:
            continue
        exact = next((field for field, names in _ALIASES.items() if raw in names), None)
        if exact:
            out[c] = exact
            continue
        for field, names in _ALIASES.items():
            if any(name in raw for name in sorted(names, key=len, reverse=True)):
                out[c] = field
                break
    return out


def _yes(value) -> bool:
    return str(value or "").strip().lower() in {"是", "true", "yes", "y", "1"}


def read_flows(xlsx: str) -> dict:
    """返回 {流程名: [route,...]}；没有合格映射表时返回空字典。"""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(xlsx, data_only=True)
    except Exception:
        return {}
    flows: dict[str, list[dict]] = {}
    for name in wb.sheetnames:
        if SHEET_HINT not in name:
            continue
        ws = wb[name]
        cols = _header_map(ws)
        if not {"flow", "stage", "match"}.issubset(set(cols.values())):
            continue
        for r in range(2, ws.max_row + 1):
            row = {}
            for c, field in cols.items():
                value = ws.cell(r, c).value
                row[field] = "" if value is None else str(value).strip()
            flow = row.get("flow", "")
            stage = row.get("stage", "")
            match = row.get("match", "")
            if not flow or not stage or not match:
                continue
            try:
                order = float(row.get("i") or 0)
            except ValueError:
                order = 0.0
            flows.setdefault(flow, []).append({
                "stage": stage,
                "i": order,
                "match": match,
                "script": row.get("script", ""),
                "next_stage": row.get("next_stage", ""),
                "continue_on_fail": _yes(row.get("continue_on_fail")),
                "_row": r,
            })
    for routes in flows.values():
        first_row = {}
        for route in routes:
            first_row.setdefault(route["stage"], route["_row"])
        routes.sort(key=lambda route: (first_row[route["stage"]], route["i"], route["_row"]))
        for route in routes:
            route.pop("_row", None)
    return flows
