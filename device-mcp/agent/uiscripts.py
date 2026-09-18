"""读 xlsx 里的『步骤脚本』表 → {脚本名: [步骤,…]}，供 run_ui_steps 按名取用。

**app 无关红线**：本模块只认表格结构（脚本名/序/动作/目标/期望出现/超时），
不认任何具体流程、页面、账号或文案——那些全在人维护的 xlsx 里。
表不存在就返回空 dict，引擎照常跑（与测试数据插件同样的"可插拔"守法）。

表结构（表名含"脚本"即被识别）：
    脚本名 | 序 | 动作 | 目标      | 期望出现   | 超时
    某流程 | 1  | tap_text | 某按钮文字 | 下一页某文字 | 5
    某流程 | 2  | input    | {某变量}   |            | 3
目标/期望的写法见 server.run_ui_steps 的文档。目标里可写 {变量}，由 resolve() 在调用时替换。
"""
import re

SHEET_HINT = "脚本"

# 表头别名 → 内部字段。人怎么写表头都尽量认，认不出的列忽略。
_ALIASES = {
    "script": ("脚本名", "脚本", "名称", "name"),
    "i": ("序", "序号", "步", "步骤", "no"),
    "action": ("动作", "操作", "action"),
    "target": ("目标", "对象", "target"),
    "expect": ("期望出现", "期望", "观测点", "expect"),
    "timeout_s": ("超时", "超时秒", "timeout"),
}
_VAR_RE = re.compile(r"\{([^{}]+)\}")


def _header_map(ws) -> dict:
    """扫第一行表头 → {列号: 内部字段}。"""
    out = {}
    for c in range(1, ws.max_column + 1):
        raw = str(ws.cell(1, c).value or "").strip()
        if not raw:
            continue
        for field, names in _ALIASES.items():
            if any(n == raw or n in raw for n in names):
                out[c] = field
                break
    return out


def read_scripts(xlsx: str) -> dict:
    """读所有表名含"脚本"的 sheet，合并成 {脚本名: [步骤,…]}（按"序"升序）。表不存在→{}。"""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(xlsx, data_only=True)
    except Exception:
        return {}
    scripts: dict[str, list[dict]] = {}
    for name in wb.sheetnames:
        if SHEET_HINT not in name:
            continue
        ws = wb[name]
        cols = _header_map(ws)
        if "script" not in cols.values() or "action" not in cols.values():
            continue                                   # 不是步骤表，跳过（别硬认）
        for r in range(2, ws.max_row + 1):
            row = {}
            for c, field in cols.items():
                v = ws.cell(r, c).value
                row[field] = "" if v is None else str(v).strip()
            sname = row.get("script", "")
            if not sname or not row.get("action"):
                continue                               # 空行/分隔行
            step = {"action": row["action"], "target": row.get("target", ""),
                    "expect": row.get("expect", "")}
            try:
                if row.get("timeout_s"):
                    step["timeout_s"] = float(row["timeout_s"])
            except ValueError:
                pass                                   # 超时列写了非数字 → 用默认值，不因此报废整条脚本
            try:
                step["_i"] = float(row.get("i") or 0)
            except ValueError:
                step["_i"] = 0.0
            scripts.setdefault(sname, []).append(step)
    for name, steps in scripts.items():
        steps.sort(key=lambda s: s["_i"])
        for s in steps:
            s.pop("_i", None)
    return scripts


_VAR_FIELDS = ("target", "expect")   # 这两列都可写 {变量}——期望值常来自用例（如"核对到的应是哪个值"）


def missing_vars(steps: list[dict], variables: dict) -> list[str]:
    """步骤里用到但 variables 没给的变量名（去重保序）。"""
    out, seen = [], set()
    for st in steps:
        for f in _VAR_FIELDS:
            for k in _VAR_RE.findall(str(st.get(f, ""))):
                if k not in variables and k not in seen:
                    seen.add(k)
                    out.append(k)
    return out


def resolve(steps: list[dict], variables: dict) -> list[dict]:
    """把步骤 target/expect 里的 {变量} 换成实际值。不改原列表。未给的变量原样留着（由 missing_vars 先拦）。"""
    def _sub(s: str) -> str:
        return _VAR_RE.sub(lambda m: str(variables.get(m.group(1), m.group(0))), s)
    return [{**st, **{f: _sub(str(st.get(f, ""))) for f in _VAR_FIELDS}} for st in steps]
