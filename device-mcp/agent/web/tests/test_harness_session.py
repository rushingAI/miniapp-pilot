import os
import sys
import asyncio
import json
import time
from types import SimpleNamespace

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from claude_agent_sdk import (  # noqa: E402
    AssistantMessage, ResultMessage, SystemMessage, TextBlock, ToolResultBlock,
    ToolUseBlock, UserMessage,
)
from harness_session import HarnessSession, _mobile_operation_skill, safe_harness_args  # noqa: E402
from model_request import TurnStopped  # noqa: E402
from telemetry import PhaseTimer  # noqa: E402


async def _emit(_event):
    return None


def _activate_parent_turn(session):
    session._turn_epoch += 1
    session._turn_active = True


def test_harness_stop_bounds_and_parallelizes_hung_interrupts():
    class HungClient:
        def __init__(self):
            self.started = asyncio.Event()

        async def interrupt(self):
            self.started.set()
            await asyncio.Event().wait()

    class Excel:
        async def interrupt_turn(self):
            return None

    async def scenario():
        main_client = HungClient()
        operator_client = HungClient()
        session = HarnessSession.__new__(HarnessSession)
        session.emit = _emit
        session.client = main_client
        session.test_excel = Excel()
        session.session_id = "session-1"
        session._operator_client = operator_client
        session._pending_final_action_records = {"pending": {}}
        session._transport_needs_reconnect = False
        session._context_loss_pending = False
        session._options = lambda resume=None: {"resume": resume}

        started = time.perf_counter()
        await asyncio.wait_for(session.stop(grace_timeout_s=0.05), timeout=0.15)

        assert main_client.started.is_set() is True
        assert operator_client.started.is_set() is True
        assert time.perf_counter() - started < 0.15
        assert session.client is None
        assert session._transport_needs_reconnect is True
        assert session._pending_final_action_records == {}

    asyncio.run(scenario())


def test_retired_transport_hooks_cannot_dispatch_after_continue(tmp_path):
    async def scenario():
        session = HarnessSession(
            _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
            provider_env={}, effort="low",
        )
        lifecycle = session._conversation()
        old_options = session._options()
        old_hook = old_options.hooks["PreToolUse"][0].hooks[0]

        lifecycle.generation += 2
        decision = await old_hook({
            "hook_event_name": "PreToolUse",
            "tool_name": "mcp__device__execute_ui_actions",
            "tool_input": {"actions": []},
        })

        output = decision["hookSpecificOutput"]
        assert output["permissionDecision"] == "deny"
        assert "generation" in output["permissionDecisionReason"]

    asyncio.run(scenario())


def test_main_and_operator_hooks_leave_room_for_progress_watchdog(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    timeouts = set()
    for options in (session.options, session._operator_options()):
        matchers = {
            id(matcher): matcher
            for event_matchers in options.hooks.values()
            for matcher in event_matchers
        }.values()
        timeouts.update(matcher.timeout for matcher in matchers)
    assert len(timeouts) == 1
    assert next(iter(timeouts)) > 30 * 60


def test_harness_stop_rebuilds_transport_when_terminal_result_is_missing(monkeypatch):
    class OldClient:
        def __init__(self):
            self.interrupted = False
            self.disconnected = False

        async def interrupt(self):
            self.interrupted = True

        async def disconnect(self):
            self.disconnected = True

    class NewClient:
        def __init__(self, options=None):
            self.options = options
            self.connected = False

        async def connect(self):
            self.connected = True

    class Excel:
        async def interrupt_turn(self):
            return None

    async def scenario():
        old = OldClient()
        session = HarnessSession.__new__(HarnessSession)
        session.emit = _emit
        session.client = old
        session.test_excel = Excel()
        session.session_id = "session-1"
        session._turn_drained = asyncio.Event()
        session._transport_needs_reconnect = False
        session._options = lambda resume=None: {"resume": resume}
        monkeypatch.setattr("harness_session.ClaudeSDKClient", NewClient)

        await session.stop(grace_timeout_s=0)
        await asyncio.sleep(0)
        await session.ensure_transport_ready()

        assert old.interrupted is True
        assert old.disconnected is True
        assert isinstance(session.client, NewClient)
        assert session.client.options == {"resume": "session-1"}
        assert session.client.connected is True

    asyncio.run(scenario())


def test_harness_binds_sdk_session_before_terminal_result(tmp_path):
    class FakeClient:
        async def query(self, _text):
            return None

        async def receive_messages(self):
            yield AssistantMessage(
                content=[], model="k3", usage={}, session_id="session-from-assistant",
            )

    async def scenario():
        session = HarnessSession(
            _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
            provider_env={}, effort="low",
        )
        session.client = FakeClient()
        await session.send("执行测试")
        assert session.session_id == "session-from-assistant"

    asyncio.run(scenario())


def test_harness_stop_waits_for_all_submitted_query_results(tmp_path):
    class FakeClient:
        def __init__(self):
            self.queries = []
            self.messages = asyncio.Queue()
            self.first_result_sent = asyncio.Event()

        async def query(self, text):
            self.queries.append(text)

        async def receive_messages(self):
            while True:
                message = await self.messages.get()
                yield message

        async def interrupt(self):
            await self.messages.put(ResultMessage(
                subtype="success", duration_ms=1, duration_api_ms=1,
                is_error=False, num_turns=1, session_id="session-1",
                total_cost_usd=0.0, usage={},
            ))
            self.first_result_sent.set()

    async def scenario():
        client = FakeClient()
        emitted = []

        async def emit(event):
            emitted.append(event)

        session = HarnessSession(
            emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
            provider_env={}, effort="low",
        )
        session.client = client
        send_task = asyncio.create_task(session.send("初始输入"))
        while not session._accepting_input:
            await asyncio.sleep(0)
        assert await session.append_input({"text": "追加输入"}) is True

        stop_task = asyncio.create_task(session.stop(grace_timeout_s=0.2))
        await client.first_result_sent.wait()
        await asyncio.sleep(0.01)
        returned_after_first_result = stop_task.done()

        await client.messages.put(ResultMessage(
            subtype="success", duration_ms=1, duration_api_ms=1,
            is_error=False, num_turns=1, session_id="session-1",
            total_cost_usd=0.0, usage={},
        ))
        await stop_task
        await send_task

        assert client.queries == ["初始输入", "追加输入"]
        assert returned_after_first_result is False
        drain = next(
            event for event in emitted
            if event.get("type") == "sdk_transport"
            and event.get("action") == "stop_completed"
        )
        assert drain["submitted_queries"] == 2
        assert drain["received_results"] == 2
        assert drain["drain_elapsed_ms"] >= 0

    asyncio.run(scenario())


def test_harness_uses_claude_code_preset_and_strict_profiles(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={"ANTHROPIC_API_KEY": "secret", "ANTHROPIC_MODEL": "k3[1m]"},
        effort="low",
    )

    assert session.options.tools == {"type": "preset", "preset": "claude_code"}
    assert session.options.system_prompt["preset"] == "claude_code"
    assert session.options.setting_sources == []
    assert session.options.skills == []
    assert session.options.strict_mcp_config is True
    assert session.options.cwd == str(session.workspace.work_dir)
    assert set(session.options.mcp_servers) == {"device", "test_excel"}
    assert "mcp__device__ask_human" not in session.options.allowed_tools
    assert "mcp__device__run_test_suite" not in session.options.allowed_tools
    assert "mcp__test_excel__open_test_suite" in session.options.allowed_tools
    assert "不得用 Bash/Python/unzip" in session.options.system_prompt["append"]
    assert "start_seq=0" in session.options.system_prompt["append"]
    assert "按用户要求跳过" in session.options.system_prompt["append"]
    assert "不复用其他 Run 的证据" in session.options.system_prompt["append"]
    assert "不得用 Bash、Read、Glob、Grep 搜索 MiniApp Pilot 产品源码" in session.options.system_prompt["append"]
    assert "当前操作已明确成功且观测点有硬证据才可 pass" in session.options.system_prompt["append"]
    assert "实际落地页与观测点不一致时记 needs_review" in session.options.system_prompt["append"]
    assert "禁止提前执行下一步" in session.options.system_prompt["append"]
    assert "最多两个独立设备动作" in session.options.system_prompt["append"]
    assert "open_test_suite 成功" in session.options.system_prompt["append"]
    assert "不得调用 Agent 或 SendMessage 接管 Suite" in session.options.system_prompt["append"]
    assert "一个 Suite 只派生一个前台 device-operator" in session.options.system_prompt["append"]
    assert "不得直接调用 Device MCP" in session.options.system_prompt["append"]
    assert "fallback 路由" in session.options.system_prompt["append"]
    assert "可由用户解除" in session.options.system_prompt["append"]
    assert "不得调用 record_test_step" in session.options.system_prompt["append"]
    assert "直接领取并返回第一步 current_step" in session.options.system_prompt["append"]
    assert "直接领取并返回下一步 current_step" in session.options.system_prompt["append"]
    assert "继续被中断的测试时" in session.options.system_prompt["append"]


def test_only_custom_agent_is_device_operator_and_cannot_recurse(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    assert session.options.agents == {}
    operator = session._operator_options().agents["device-operator"]
    assert "Agent" in operator.disallowedTools
    assert "Skill" in operator.disallowedTools
    assert "AskUserQuestion" in operator.disallowedTools
    assert operator.background is False
    assert operator.maxTurns == 220
    assert "mcp__device__get_screen" in operator.tools
    assert "mcp__test_excel__record_test_step" in operator.tools


def test_device_operator_contract_makes_guarded_action_batches_the_default(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    operator_options = session._operator_options()
    prompt = operator_options.agents["device-operator"].prompt
    assert "手机操作闭环" in prompt
    assert "一次执行 1–8 个动作" in prompt
    assert "普通界面写操作统一使用 execute_ui_actions" in prompt
    assert "postcondition" in prompt
    assert "tap_image" in prompt
    assert "动作工具返回的画面就是新的当前观察帧" in prompt
    assert "才补充一次 `screenshot`" in prompt
    assert "不循环截图" in prompt
    assert "suite_completed=true" in prompt
    assert "visual_on_miss=true" in prompt
    assert "拿到图后别重复相同查询" in prompt
    assert "allow_bottom_edge=true" in prompt
    assert "后续 `tap_text`" in prompt
    assert "执行到某个 action 时再查找" in prompt
    assert "`find` 唯一命中" in prompt
    assert "bounds_unreliable=false" in prompt
    assert "界面预期实际显示的原文或可靠标识" in prompt
    assert "最终条件能够发现前序漏点" in prompt
    assert "仍优先尝试语义" in prompt
    assert "不得输出工具间进度说明" in prompt
    assert "主 Agent 负责理解用户目标" not in prompt
    assert "主 Agent 负责理解用户目标" not in operator_options.system_prompt["append"]
    assert len(prompt) <= 14_000


def test_native_prompts_require_compact_single_owner_suite_completion(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    operator_prompt = session._operator_options().agents["device-operator"].prompt
    main_prompt = session.options.system_prompt["append"]
    operator_system_prompt = session._operator_options().system_prompt["append"]
    assert '"status":"completed"' in operator_prompt
    assert "不得用散文、Markdown" in operator_prompt
    assert "主 Agent 是面向用户总结的唯一负责人" in main_prompt
    assert "不得再次逐条复述" in main_prompt
    assert "主 Agent 是面向用户总结的唯一负责人" not in operator_system_prompt
    assert "open_test_suite 成功后" not in operator_system_prompt
    assert "单个 JSON outcome" in operator_system_prompt


def test_device_operator_contract_stops_after_a_structured_final_assertion(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    prompt = session._operator_options().agents["device-operator"].prompt
    assert "收口批次" in prompt
    assert "本步骤要求的业务动作已经全部执行或当前已成立" in prompt
    assert "批次后无需新的点击、输入、滑动或分支判断" in prompt
    assert "中间批次，保持 final_assertion=false" in prompt
    assert "最后一个产生副作用的 action 必须直接携带证明 observe 的 postcondition" in prompt
    assert "本批结束当前 Case" in prompt
    assert "final_assertion=true" in prompt
    assert "只有弱条件的收口批次" in prompt
    assert "AUTO_RECORDED" in prompt
    assert "禁止再调用 record_test_step" in prompt
    assert "未确认的视觉路径" in prompt
    assert "screen_change、settle 和单独 absent" in prompt
    assert "不能单独作为最终结果的语义证明" in prompt


def test_device_operator_contract_binds_case_observation_to_final_side_effect(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    prompt = session._operator_options().agents["device-operator"].prompt
    assert "operation 与 observe 共同构成一个完整 Case 目标" in prompt
    assert "最后一个产生副作用的 action" in prompt
    assert "不得把结果检查拆成后续独立 wait" in prompt
    assert "纯观察 Case" in prompt


def test_device_operator_contract_preserves_all_required_observations(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    prompt = session._operator_options().agents["device-operator"].prompt
    assert "多个必须同时成立的结果信号使用 `all`" in prompt
    assert "`any` 只用于一个足以独立证明结果的语义候选与弱保底" in prompt
    assert "收口条件必须覆盖 observe 的全部必要结果" in prompt
    assert "不得自动记录" in prompt


def test_mobile_operation_contract_uses_literal_semantics_before_visual_fallback(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    prompt = session._operator_options().agents["device-operator"].prompt
    assert "界面预期实际显示的原文或可靠标识" in prompt
    assert "缩写、截词或猜测" in prompt
    assert "a11y 可用性未知但候选可靠时仍优先尝试语义" in prompt
    assert "只给迟到的正向语义保留最多约五秒宽限" in prompt


def test_operator_prompt_treats_case_text_as_opaque_natural_language(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    operation = "Tap Continue -> inspect result\n[layout:v2] 文案格式可变化"
    observe = "视觉确认：图标由 ○ 变为 ●；无需固定关键词"

    prompt = session._operator_prompt({
        "suite_id": "suite-format-free",
        "current_step": {
            "step_id": "case-1", "kind": "case", "seq": 1,
            "operation": operation, "observe": observe,
        },
    }, resume=False)

    assert operation.replace("\n", "\\n") in prompt
    assert observe in prompt


def test_harness_does_not_force_final_assertion_on_device_batches(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    _open_single_case(session, "no-final-rewrite.xlsx")
    tool_input = {
        "actions": [{
            "action": "tap_text", "target": "继续",
            "postcondition": {"settle_ms": 200, "checks": []},
        }],
        "final_assertion": False,
    }

    decision = asyncio.run(session._hook({
        "hook_event_name": "PreToolUse", "tool_use_id": "no-rewrite-1",
        "tool_name": "mcp__device__execute_ui_actions",
        "tool_input": tool_input,
        "agent_id": "", "agent_type": "device-operator",
    }))

    assert decision == {}
    assert tool_input["final_assertion"] is False


def test_mobile_operation_skill_is_generic_and_compact():
    skill = _mobile_operation_skill()

    assert len(skill) < 10_000
    for testing_term in (
            "record_test_step", "AUTO_RECORDED", "required_flow", "Suite Registry",
            "Excel", "verdict", "Case 级"):
        assert testing_term not in skill


def _open_single_case(session, filename="auto-record.xlsx", notes=""):
    source = session.workspace.input_dir / filename
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "用例"
    ws.append(["序号", "操作", "观测点", "观测效果"])
    ws.append([1, "点击确认", "显示完成", ""])
    wb.save(source)
    return asyncio.run(session.test_excel.open({
        "filename": filename, "sheet": "用例", "start_seq": 0, "max_cases": 0,
        "notes": notes,
    }))


def test_operator_prompt_preserves_first_step_suite_notes(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    opened = _open_single_case(
        session, "suite-notes.xlsx",
        notes="必须调用 repeat_swipe(count=3)，并沿用至 Suite 完成",
    )

    prompt = session._operator_prompt(opened, resume=False)

    assert "必须调用 repeat_swipe(count=3)" in prompt
    assert "沿用至 Suite 完成" in prompt


def test_structured_final_action_is_recorded_once_without_rewriting_device_output(tmp_path):
    emitted = []

    async def capture(event):
        emitted.append(event)

    session = HarnessSession(
        capture, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    opened = _open_single_case(session)
    suite_id = opened["suite_id"]
    step_id = opened["current_step"]["step_id"]
    state = session.test_excel.suites[suite_id]
    evidence = state.out_dir / "trace" / step_id / "proof.jsonl"
    evidence.write_text("{}\n", encoding="utf-8")
    device_payload = {
        "ok": True,
        "postcondition_satisfied": True,
        "final_assertion_confirmed": True,
        "evidence_paths": [str(evidence)],
        "final_postcondition": {
            "satisfied": True,
            "checks": [{
                "type": "a11y", "state": "present", "satisfied": True,
                "matched": {"text": "显示完成", "resource_id": "result"},
            }],
        },
    }
    response = {"content": [{"type": "text", "text": json.dumps(device_payload)}]}

    asyncio.run(session._hook({
        "hook_event_name": "PreToolUse", "tool_use_id": "batch-1",
        "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
        "agent_id": "", "agent_type": "device-operator",
    }))
    first = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse", "tool_use_id": "batch-1",
        "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
        "tool_response": response,
        "agent_id": "", "agent_type": "device-operator",
    }))
    duplicate = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse", "tool_use_id": "batch-1",
        "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
        "tool_response": response,
        "agent_id": "", "agent_type": "device-operator",
    }))

    output = first["hookSpecificOutput"]
    assert "updatedMCPToolOutput" not in output
    assert "AUTO_RECORDED" in output["additionalContext"]
    assert duplicate == {}
    assert state.status == "completed"
    assert len(state.results) == 1
    assert state.results[0].verdict == "pass"
    assert state.results[0].evidence == [str(evidence)]
    bridge_events = [event for event in emitted if event.get("type") == "completion_bridge"]
    assert bridge_events == [{
        "type": "completion_bridge", "status": "AUTO_RECORDED",
        "suite_id": suite_id, "step_id": step_id, "tool_use_id": "batch-1",
    }]


def test_visual_only_final_action_does_not_advance_the_suite(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    opened = _open_single_case(session, "visual-only.xlsx")
    suite_id = opened["suite_id"]
    step_id = opened["current_step"]["step_id"]
    state = session.test_excel.suites[suite_id]
    evidence = state.out_dir / "trace" / step_id / "screen.png"
    evidence.write_bytes(b"screen")
    response = {"content": [
        {"type": "text", "text": json.dumps({
            "ok": True,
            "postcondition_satisfied": True,
            "final_assertion_confirmed": False,
            "evidence_paths": [str(evidence)],
        })},
        {"type": "image", "data": "aW1hZ2U=", "mimeType": "image/png"},
    ]}

    asyncio.run(session._hook({
        "hook_event_name": "PreToolUse", "tool_use_id": "visual-1",
        "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
        "agent_id": "", "agent_type": "device-operator",
    }))
    result = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse", "tool_use_id": "visual-1",
        "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
        "tool_response": response,
        "agent_id": "", "agent_type": "device-operator",
    }))

    context = result["hookSpecificOutput"]["additionalContext"]
    assert "CURRENT_FRAME_READY" in context
    assert "当前图片就是动作后观察帧" in context
    assert "满足当前步骤观测点时直接记录" in context
    assert state.status == "active"
    assert state.current_step["step_id"] == step_id
    assert state.results == []


def test_visual_frame_hint_is_not_added_for_failed_action(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    response = {"content": [
        {"type": "text", "text": json.dumps({"ok": False, "evidence_paths": []})},
        {"type": "image", "data": "aW1hZ2U=", "mimeType": "image/png"},
    ]}

    result = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse", "tool_use_id": "failed-visual",
        "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
        "tool_response": response,
        "agent_id": "", "agent_type": "device-operator",
    }))

    assert result == {}


def test_auto_record_requires_evidence_and_keeps_the_current_case(tmp_path):
    emitted = []

    async def capture(event):
        emitted.append(event)

    session = HarnessSession(
        capture, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    opened = _open_single_case(session, "missing-evidence.xlsx")
    suite_id = opened["suite_id"]
    step_id = opened["current_step"]["step_id"]
    state = session.test_excel.suites[suite_id]
    response = {"content": [{"type": "text", "text": json.dumps({
        "ok": True,
        "postcondition_satisfied": True,
        "final_assertion_confirmed": True,
        "evidence_paths": [],
    })}]}

    async def scenario():
        await session._hook({
            "hook_event_name": "PreToolUse", "tool_use_id": "missing-1",
            "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
            "agent_id": "", "agent_type": "device-operator",
        })
        return await session._hook({
            "hook_event_name": "PostToolUse", "tool_use_id": "missing-1",
            "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
            "tool_response": response,
            "agent_id": "", "agent_type": "device-operator",
        })

    result = asyncio.run(scenario())

    assert "AUTO_RECORD_SKIPPED" in result["hookSpecificOutput"]["additionalContext"]
    assert state.status == "active"
    assert state.current_step["step_id"] == step_id
    assert state.results == []
    assert [event["status"] for event in emitted
            if event.get("type") == "completion_bridge"] == ["SKIPPED"]


def test_stop_wins_over_auto_record_without_replaying_the_action(tmp_path):
    control = SimpleNamespace(
        stopped=False, stop_reason="用户停止", control_lock=asyncio.Lock(),
    )
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low", control=control,
    )
    opened = _open_single_case(session, "stop-auto-record.xlsx")
    suite_id = opened["suite_id"]
    step_id = opened["current_step"]["step_id"]
    state = session.test_excel.suites[suite_id]
    evidence = state.out_dir / "trace" / step_id / "proof.jsonl"
    evidence.write_text("{}\n", encoding="utf-8")
    response = {"content": [{"type": "text", "text": json.dumps({
        "ok": True,
        "postcondition_satisfied": True,
        "final_assertion_confirmed": True,
        "evidence_paths": [str(evidence)],
    })}]}

    async def scenario():
        await session._hook({
            "hook_event_name": "PreToolUse", "tool_use_id": "stopped-1",
            "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
            "agent_id": "", "agent_type": "device-operator",
        })
        control.stopped = True
        return await session._hook({
            "hook_event_name": "PostToolUse", "tool_use_id": "stopped-1",
            "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
            "tool_response": response,
            "agent_id": "", "agent_type": "device-operator",
        })

    result = asyncio.run(scenario())

    assert "AUTO_RECORD_SKIPPED" in result["hookSpecificOutput"]["additionalContext"]
    assert state.status == "active"
    assert state.current_step["step_id"] == step_id
    assert state.results == []


def test_required_flow_failure_does_not_advance_auto_record(tmp_path):
    import tools as device_tools

    emitted = []

    async def capture(event):
        emitted.append(event)

    session = HarnessSession(
        capture, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    opened = _open_single_case(session, "required-flow.xlsx")
    suite_id = opened["suite_id"]
    step_id = opened["current_step"]["step_id"]
    state = session.test_excel.suites[suite_id]
    state.current_step["required_flow"] = "登录流程"
    device_tools.set_required_flow("登录流程", step_id)
    evidence = state.out_dir / "trace" / step_id / "proof.jsonl"
    evidence.write_text("{}\n", encoding="utf-8")
    response = {"content": [{"type": "text", "text": json.dumps({
        "ok": True,
        "postcondition_satisfied": True,
        "final_assertion_confirmed": True,
        "evidence_paths": [str(evidence)],
    })}]}

    async def scenario():
        await session._hook({
            "hook_event_name": "PreToolUse", "tool_use_id": "flow-1",
            "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
            "agent_id": "", "agent_type": "device-operator",
        })
        return await session._hook({
            "hook_event_name": "PostToolUse", "tool_use_id": "flow-1",
            "tool_name": "mcp__device__execute_ui_actions", "tool_input": {},
            "tool_response": response,
            "agent_id": "", "agent_type": "device-operator",
        })

    try:
        result = asyncio.run(scenario())
    finally:
        device_tools.clear_required_flow()

    assert "AUTO_RECORD_FAILED" in result["hookSpecificOutput"]["additionalContext"]
    assert state.status == "active"
    assert state.current_step["step_id"] == step_id
    assert state.results == []
    assert [event["status"] for event in emitted
            if event.get("type") == "completion_bridge"] == ["FAILED"]


def test_thinking_disabled_mode_uses_sdk_thinking_config_without_invalid_effort(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="none",
    )

    assert session.options.thinking == {"type": "disabled"}
    assert session.options.effort is None
    assert session._operator_options().agents["device-operator"].effort is None


def test_native_hook_rejects_raw_subagent_suite_and_device_mutations(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    noted = []
    observed = []
    session.test_excel.note_tool_actor = lambda name, actor, agent_id: noted.append(
        (name, actor, agent_id)
    )

    async def observe(name, actor, agent_id):
        observed.append((name, actor, agent_id))

    session.test_excel.observe_device_action = observe

    async def scenario():
        await session._hook({
            "hook_event_name": "SubagentStart", "agent_type": "device-operator",
            "agent_id": "operator-1",
        })
        await session._hook({
            "hook_event_name": "PreToolUse", "tool_use_id": "tool-next",
            "tool_name": "mcp__test_excel__next_test_step", "tool_input": {},
            "agent_id": "operator-1",
        })
        await session._hook({
            "hook_event_name": "PreToolUse", "tool_use_id": "tool-device",
            "tool_name": "mcp__device__get_screen", "tool_input": {},
            "agent_id": "operator-1",
        })

    asyncio.run(scenario())

    assert noted == []
    assert observed == []


def test_open_success_deterministically_runs_operator_before_returning_to_main(
        tmp_path, monkeypatch):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    delegated = []

    async def run_operator(payload, *, resume=False):
        delegated.append((payload, resume))
        return {
            "ok": True, "status": "completed", "suite_id": payload["suite_id"],
            "result": "SUITE_COMPLETED suite-1 pass 1/1 child-run",
        }

    monkeypatch.setattr(session, "_run_device_operator", run_operator)
    session.test_excel.suites["suite-1"] = SimpleNamespace(
        status="completed", current_step=None, ready_to_finish=True,
        terminal_reason="全部步骤完成",
    )
    monkeypatch.setattr(session.test_excel, "_completed_payload", lambda _state: {
        "suite_completed": True,
        "summary": {"verdict": "pass"},
        "verdict": "pass",
        "out_xlsx": "/runs/child-run/result.xlsx",
        "run_id": "child-run",
    })
    response = {
        "content": [{"type": "text", "text": json.dumps({
            "suite_id": "suite-1", "run_id": "child-run",
            "current_step": {"step_id": "case-1", "operation": "点击 A"},
        }, ensure_ascii=False)}],
    }

    hook = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse",
        "tool_name": "mcp__test_excel__open_test_suite",
        "tool_use_id": "open-1", "tool_input": {}, "tool_response": response,
        "agent_id": "", "agent_type": "",
    }))

    assert delegated == [({
        "suite_id": "suite-1", "run_id": "child-run",
        "current_step": {"step_id": "case-1", "operation": "点击 A"},
    }, False)]
    output = hook["hookSpecificOutput"]
    assert output["hookEventName"] == "PostToolUse"
    updated = output["updatedMCPToolOutput"]
    payload = json.loads(updated["content"][0]["text"])
    assert payload["deterministic_delegation"]["status"] == "completed"
    assert payload["suite_completed"] is True
    assert payload["summary"] == {"verdict": "pass"}
    assert "current_step" not in payload
    assert "不要调用 Agent 或 SendMessage" in output["additionalContext"]


def test_operator_fallback_finishes_in_the_same_open_hook_and_keeps_registry_cursor(
        tmp_path, monkeypatch):
    """A fallback may change executors, never the Suite owner or parent turn."""
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    registry_step = {"step_id": "case-2", "operation": "继续当前用例"}
    state = SimpleNamespace(
        status="active", current_step=registry_step, ready_to_finish=False,
        terminal_reason="", suite_id="suite-1",
    )
    session.test_excel.suites["suite-1"] = state
    session.test_excel.active_suite_id = "suite-1"
    calls = []

    async def run_adapter(payload, *, resume, route):
        calls.append((route, payload["current_step"], resume))
        if route == "primary":
            return {
                "ok": False, "status": "operator_protocol_error",
                "suite_id": "suite-1", "reason": "invalid_operator_outcome",
            }
        state.status = "completed"
        state.current_step = None
        state.ready_to_finish = True
        session.test_excel.active_suite_id = ""
        return {
            "ok": True, "status": "completed", "suite_id": "suite-1",
            "route": "fallback", "verdict": "pass",
        }

    monkeypatch.setattr(session, "_run_operator_adapter", run_adapter)
    monkeypatch.setattr(session.test_excel, "_completed_payload", lambda _state: {
        "suite_completed": True,
        "summary": {"verdict": "pass"},
        "verdict": "pass",
        "out_xlsx": "/runs/child-run/result.xlsx",
        "run_id": "child-run",
    })
    response = {"content": [{"type": "text", "text": json.dumps({
        "suite_id": "suite-1", "run_id": "child-run",
        "current_step": registry_step,
    }, ensure_ascii=False)}]}

    hook = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse",
        "tool_name": "mcp__test_excel__open_test_suite",
        "tool_use_id": "open-fallback", "tool_input": {}, "tool_response": response,
        "agent_id": "", "agent_type": "",
    }))

    assert calls == [
        ("primary", registry_step, False),
        ("fallback", registry_step, False),
    ]
    output = hook["hookSpecificOutput"]
    payload = json.loads(output["updatedMCPToolOutput"]["content"][0]["text"])
    assert payload["suite_completed"] is True
    assert payload["deterministic_delegation"]["status"] == "completed"
    assert payload["deterministic_delegation"]["route"] == "fallback"
    assert "不要调用 Agent 或 SendMessage" in output["additionalContext"]
    assert "受控降级" not in output["additionalContext"]


def test_open_hook_replaces_stale_first_step_with_registry_cursor(
        tmp_path, monkeypatch):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    registry_step = {"step_id": "case-1", "operation": "进入小程序"}
    session.test_excel.suites["suite-1"] = SimpleNamespace(
        status="active", current_step=registry_step, ready_to_finish=False,
        terminal_reason="",
    )

    async def run_operator(payload, *, resume=False):
        return {
            "ok": False, "status": "delegation_fallback",
            "suite_id": payload["suite_id"], "reason": "CancelledError",
        }

    monkeypatch.setattr(session, "_run_device_operator", run_operator)
    response = {"content": [{"type": "text", "text": json.dumps({
        "suite_id": "suite-1", "run_id": "child-run",
        "current_step": {"step_id": "setup-0", "operation": "核对并切号"},
    }, ensure_ascii=False)}]}

    hook = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse",
        "tool_name": "mcp__test_excel__open_test_suite",
        "tool_use_id": "open-1", "tool_input": {}, "tool_response": response,
        "agent_id": "", "agent_type": "",
    }))

    updated = hook["hookSpecificOutput"]["updatedMCPToolOutput"]
    payload = json.loads(updated["content"][0]["text"])
    assert payload["status"] == "active"
    assert payload["current_step"] == registry_step
    assert payload["current_step"]["step_id"] != "setup-0"


def test_inactivity_timeout_keeps_registry_cursor_authoritative_for_continue(
        tmp_path, monkeypatch):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    registry_step = {"step_id": "case-2", "operation": "继续当前用例"}
    session.test_excel.suites["suite-1"] = SimpleNamespace(
        status="active", current_step=registry_step, ready_to_finish=False,
        terminal_reason="",
    )
    delegated = []

    async def run_operator(payload, *, resume=False):
        delegated.append((payload["current_step"], resume))
        return {
            "ok": False, "status": "inactivity_timeout",
            "suite_id": payload["suite_id"],
            "reason": "operator_inactivity_timeout",
        }

    monkeypatch.setattr(session, "_run_device_operator", run_operator)
    opened = {"content": [{"type": "text", "text": json.dumps({
        "suite_id": "suite-1",
        "current_step": {"step_id": "setup-0", "operation": "核对并切号"},
    }, ensure_ascii=False)}]}

    first_hook = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse",
        "tool_name": "mcp__test_excel__open_test_suite",
        "tool_use_id": "open-1", "tool_input": {}, "tool_response": opened,
        "agent_id": "", "agent_type": "",
    }))
    first_payload = json.loads(
        first_hook["hookSpecificOutput"]["updatedMCPToolOutput"]["content"][0]["text"]
    )

    resumed = {"content": [{"type": "text", "text": json.dumps({
        "status": "resume_current", **registry_step,
    }, ensure_ascii=False)}]}
    asyncio.run(session._hook({
        "hook_event_name": "PostToolUse",
        "tool_name": "mcp__test_excel__next_test_step",
        "tool_use_id": "next-1", "tool_input": {"suite_id": "suite-1"},
        "tool_response": resumed, "agent_id": "", "agent_type": "",
    }))

    assert first_payload["status"] == "active"
    assert first_payload["current_step"] == registry_step
    assert first_payload["deterministic_delegation"]["status"] == "inactivity_timeout"
    assert "不要调用 Agent 自动重试" in first_hook["hookSpecificOutput"]["additionalContext"]
    assert delegated == [
        ({"step_id": "setup-0", "operation": "核对并切号"}, False),
        (registry_step, True),
    ]


def test_open_hook_reports_interrupted_registry_state_without_stale_step(
        tmp_path, monkeypatch):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    registry_step = {"step_id": "case-2", "operation": "继续当前用例"}
    session.test_excel.suites["suite-1"] = SimpleNamespace(
        status="interrupted", current_step=registry_step, ready_to_finish=False,
        terminal_reason="开始新测试",
    )

    async def run_operator(payload, *, resume=False):
        return {
            "ok": False, "status": "delegation_fallback",
            "suite_id": payload["suite_id"], "reason": "CancelledError",
        }

    monkeypatch.setattr(session, "_run_device_operator", run_operator)
    response = {"content": [{"type": "text", "text": json.dumps({
        "suite_id": "suite-1",
        "current_step": {"step_id": "setup-0", "operation": "核对并切号"},
    }, ensure_ascii=False)}]}

    hook = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse",
        "tool_name": "mcp__test_excel__open_test_suite",
        "tool_use_id": "open-1", "tool_input": {}, "tool_response": response,
        "agent_id": "", "agent_type": "",
    }))

    updated = hook["hookSpecificOutput"]["updatedMCPToolOutput"]
    payload = json.loads(updated["content"][0]["text"])
    assert payload["status"] == "interrupted"
    assert payload["terminal_reason"] == "开始新测试"
    assert payload["current_step"] == registry_step


def test_open_hook_accepts_the_real_sdk_mcp_content_array_shape(tmp_path, monkeypatch):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    delegated = []

    async def run_operator(payload, *, resume=False):
        delegated.append((payload["suite_id"], resume))
        return {"ok": True, "status": "completed", "suite_id": payload["suite_id"]}

    monkeypatch.setattr(session, "_run_device_operator", run_operator)
    # Claude Code PostToolUse supplies an MCP result as the content-block array itself,
    # not the CallToolResult {content: [...]} wrapper used by our original unit fixture.
    response = [{"type": "text", "text": json.dumps({
        "suite_id": "suite-real-shape",
        "current_step": {"step_id": "case-1", "operation": "点击 A"},
    })}]

    hook = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse",
        "tool_name": "mcp__test_excel__open_test_suite",
        "tool_use_id": "open-real", "tool_input": {}, "tool_response": response,
        "agent_id": "", "agent_type": "",
    }))

    assert delegated == [("suite-real-shape", False)]
    output = hook["hookSpecificOutput"]
    assert output["hookEventName"] == "PostToolUse"
    assert isinstance(output["updatedMCPToolOutput"], list)
    updated = json.loads(output["updatedMCPToolOutput"][0]["text"])
    assert updated["deterministic_delegation"]["status"] == "completed"


def test_unreadable_open_result_reports_protocol_error_without_background_fallback(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    hook = asyncio.run(session._hook({
        "hook_event_name": "PostToolUse",
        "tool_name": "mcp__test_excel__open_test_suite",
        "tool_use_id": "open-bad-shape", "tool_input": {},
        "tool_response": [{"type": "image", "data": "ignored"}],
        "agent_id": "", "agent_type": "",
    }))

    context = hook["hookSpecificOutput"]["additionalContext"]
    assert "operator_protocol_error" in context
    assert "不得调用 Agent、SendMessage" in context
    assert "不得声称已自动移交" in context


def test_deterministic_operator_uses_selected_agent_profile(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    options = session._operator_options()

    assert options.extra_args["agent"] == "device-operator"
    assert set(options.agents) == {"device-operator"}
    assert "Agent" in options.agents["device-operator"].disallowedTools


def test_operator_outcome_uses_structured_fields_instead_of_magic_words(
        tmp_path, monkeypatch):
    class FakeOperatorClient:
        def __init__(self, options):
            self.options = options

        async def connect(self):
            return None

        async def query(self, _prompt):
            return None

        async def receive_messages(self):
            yield AssistantMessage(
                content=[TextBlock("需要用户确认支付密码后再继续")],
                model="k3", usage={}, session_id="operator-session-1",
            )
            yield ResultMessage(
                subtype="success", duration_ms=1, duration_api_ms=1,
                is_error=False, num_turns=1, session_id="operator-session-1",
                total_cost_usd=0.0, usage={},
                structured_output={
                    "status": "needs_user_input", "suite_id": "suite-1",
                    "question": "请输入支付密码", "current_step_id": "case-2",
                },
            )

        async def disconnect(self):
            return None

    monkeypatch.setattr("harness_session.ClaudeSDKClient", FakeOperatorClient)
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.test_excel.active_suite_id = "suite-1"
    session.test_excel.suites["suite-1"] = SimpleNamespace(
        status="active", current_step={"step_id": "case-2"},
    )
    _activate_parent_turn(session)

    result = asyncio.run(session._run_operator_adapter({
        "suite_id": "suite-1", "current_step": {"step_id": "case-2"},
    }, route="primary"))

    assert result["status"] == "needs_user_input"
    assert result["question"] == "请输入支付密码"
    assert result["current_step_id"] == "case-2"


def test_operator_prose_cannot_be_misclassified_as_a_domain_outcome(
        tmp_path, monkeypatch):
    class FakeOperatorClient:
        def __init__(self, options):
            self.options = options

        async def connect(self):
            return None

        async def query(self, _prompt):
            return None

        async def receive_messages(self):
            yield AssistantMessage(
                content=[TextBlock("当前需要用户确认后再继续")],
                model="k3", usage={}, session_id="operator-session-1",
            )
            yield ResultMessage(
                subtype="success", duration_ms=1, duration_api_ms=1,
                is_error=False, num_turns=1, session_id="operator-session-1",
                total_cost_usd=0.0, usage={}, result="当前需要用户确认后再继续",
            )

        async def disconnect(self):
            return None

    monkeypatch.setattr("harness_session.ClaudeSDKClient", FakeOperatorClient)
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.test_excel.active_suite_id = "suite-1"
    session.test_excel.suites["suite-1"] = SimpleNamespace(
        status="active", current_step={"step_id": "case-2"},
    )
    _activate_parent_turn(session)

    result = asyncio.run(session._run_operator_adapter({
        "suite_id": "suite-1", "current_step": {"step_id": "case-2"},
    }, route="primary"))

    assert result["status"] == "operator_protocol_error"
    assert result["reason"] == "outcome_not_json_object"


def test_revoked_operator_lease_denies_late_suite_mutation(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session._turn_epoch = 7
    session._turn_active = True
    lease = session._execution_leases.issue(
        turn_epoch=7, suite_id="suite-1", operator_id="primary-suite-1",
        route="primary",
    )
    options = session._operator_options(execution_lease=lease)
    hook = options.hooks["PreToolUse"][0].hooks[0]
    session._execution_leases.revoke(lease)

    decision = asyncio.run(hook({
        "hook_event_name": "PreToolUse",
        "tool_name": "mcp__device__execute_ui_actions",
        "tool_input": {"actions": []},
        "agent_id": "", "agent_type": "device-operator",
    }))

    output = decision["hookSpecificOutput"]
    assert output["permissionDecision"] == "deny"
    assert "执行凭证已失效" in output["permissionDecisionReason"]


def test_active_suite_rejects_raw_agent_and_sendmessage_control(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.turn_mode = "continue_active_suite"
    session.test_excel.active_suite_id = "suite-1"

    async def scenario(name):
        return await session._hook({
            "hook_event_name": "PreToolUse", "tool_name": name,
            "tool_input": {}, "agent_id": "", "agent_type": "",
        })

    for name in ("Agent", "SendMessage"):
        decision = asyncio.run(scenario(name))["hookSpecificOutput"]
        assert decision["permissionDecision"] == "deny"
        assert "Harness 前台执行" in decision["permissionDecisionReason"]

    for name in (
            "mcp__device__execute_ui_actions",
            "mcp__test_excel__record_test_step",
            "mcp__test_excel__finish_test_suite"):
        decision = asyncio.run(scenario(name))["hookSpecificOutput"]
        assert decision["permissionDecision"] == "deny"
        assert "前台 operator" in decision["permissionDecisionReason"]


def test_schema_validation_failure_is_exported_to_parent_and_child_journals(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    opened = _open_single_case(session, "schema-failure.xlsx")
    suite_id = opened["suite_id"]
    state = session.test_excel.suites[suite_id]
    tool_input = {
        "actions": [{"action": "tap_text", "target": "继续", "settle_ms": 300}],
    }

    async def scenario():
        await session._hook({
            "hook_event_name": "PreToolUse", "tool_use_id": "bad-schema-1",
            "tool_name": "mcp__device__execute_ui_actions",
            "tool_input": tool_input, "agent_id": "", "agent_type": "device-operator",
        })
        await session._hook({
            "hook_event_name": "PostToolUseFailure", "tool_use_id": "bad-schema-1",
            "tool_name": "mcp__device__execute_ui_actions",
            "tool_input": tool_input,
            "error": "Input validation error: Additional properties are not allowed ('settle_ms' was unexpected)",
            "agent_id": "", "agent_type": "device-operator",
        })

    asyncio.run(scenario())

    parent = session.trace_dir / "execution_journal.jsonl"
    child = state.out_dir / "trace" / "execution_journal.jsonl"
    for path in (parent, child):
        records = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
        failure = next(item for item in records if item["event"] == "tool_schema_failure")
        assert failure["unexpected_fields"] == ["settle_ms"]
        assert failure["action_executed"] is False
        assert "继续" not in json.dumps(records, ensure_ascii=False)


def test_fourth_device_write_is_blocked_after_three_unchanged_failures(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    opened = _open_single_case(session, "no-progress.xlsx")
    state = session.test_excel.suites[opened["suite_id"]]
    image = state.out_dir / "trace" / opened["current_step"]["step_id"] / "after.png"
    image.write_bytes(b"same-screen")

    async def fail(index):
        tool_input = {
            "actions": [{
                "action": "tap_image", "target": f"{100 + index},800",
                "postcondition": {"settle_ms": 100, "checks": []},
            }],
        }
        decision = await session._hook({
            "hook_event_name": "PreToolUse", "tool_use_id": f"attempt-{index}",
            "tool_name": "mcp__device__execute_ui_actions",
            "tool_input": tool_input, "agent_id": "", "agent_type": "device-operator",
        })
        if decision:
            return decision
        response = {"content": [{"type": "text", "text": json.dumps({
            "ok": False,
            "failed_at": {
                "reason": "校验超时", "action_executed": True,
                "after_image": {"path": str(image)},
            },
        })}]}
        await session._hook({
            "hook_event_name": "PostToolUse", "tool_use_id": f"attempt-{index}",
            "tool_name": "mcp__device__execute_ui_actions",
            "tool_input": tool_input, "tool_response": response,
            "agent_id": "", "agent_type": "device-operator",
        })
        return {}

    async def scenario():
        assert await fail(1) == {}
        assert await fail(2) == {}
        assert await fail(3) == {}
        return await fail(4)

    fourth = asyncio.run(scenario())["hookSpecificOutput"]
    assert fourth["permissionDecision"] == "deny"
    assert "连续三次" in fourth["permissionDecisionReason"]


def test_deterministic_operator_runner_finishes_on_its_result_message(
        tmp_path, monkeypatch):
    instances = []

    class FakeOperatorClient:
        def __init__(self, options):
            self.options = options
            self.prompt = ""
            instances.append(self)

        async def connect(self):
            return None

        async def query(self, prompt):
            self.prompt = prompt

        async def receive_messages(self):
            yield AssistantMessage(
                content=[TextBlock("SUITE_COMPLETED suite-1 pass 1/1 child-run")],
                model="k3", usage={}, session_id="operator-session-1",
            )
            yield ResultMessage(
                subtype="success", duration_ms=1, duration_api_ms=1,
                is_error=False, num_turns=1, session_id="operator-session-1",
                total_cost_usd=0.0, usage={},
            )

        async def disconnect(self):
            return None

    monkeypatch.setattr("harness_session.ClaudeSDKClient", FakeOperatorClient)
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.test_excel.suites["suite-1"] = type("State", (), {"status": "completed"})()
    _activate_parent_turn(session)

    result = asyncio.run(session._run_device_operator({
        "suite_id": "suite-1", "current_step": {"step_id": "case-1"},
    }))

    assert result["ok"] is True
    assert result["status"] == "completed"
    assert result["operator_session_id"] == "operator-session-1"
    assert "result" not in result
    assert "suite_id=suite-1" in instances[0].prompt


def test_resumed_operator_receives_queued_user_answer_in_initial_prompt(
        tmp_path, monkeypatch):
    instances = []

    class FakeOperatorClient:
        def __init__(self, options):
            self.options = options
            self.prompt = ""
            instances.append(self)

        async def connect(self):
            return None

        async def query(self, prompt):
            self.prompt = prompt

        async def receive_messages(self):
            yield ResultMessage(
                subtype="success", duration_ms=1, duration_api_ms=1,
                is_error=False, num_turns=1, session_id="operator-session-1",
                total_cost_usd=0.0, usage={},
                structured_output={
                    "status": "needs_user_input", "suite_id": "suite-1",
                    "question": "是否继续？", "current_step_id": "case-1",
                },
            )

        async def disconnect(self):
            return None

    monkeypatch.setattr("harness_session.ClaudeSDKClient", FakeOperatorClient)
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.test_excel.active_suite_id = "suite-1"
    session.test_excel.suites["suite-1"] = SimpleNamespace(
        status="active", current_step={"step_id": "case-1"},
    )
    session._operator_suite_id = "suite-1"
    session._operator_session_id = "operator-session-1"
    session._operator_inputs.put_nowait({
        "suite_id": "suite-1",
        "text": "用户对“是否继续？”的答复：E2E_CONTINUE",
    })
    _activate_parent_turn(session)

    asyncio.run(session._run_operator_adapter({
        "suite_id": "suite-1", "current_step": {"step_id": "case-1"},
    }, resume=True, route="primary"))

    assert "用户对“是否继续？”的答复：E2E_CONTINUE" in instances[0].prompt
    assert session._operator_inputs.empty()


def test_deterministic_operator_records_actionable_model_and_local_runtime_boundaries(
        tmp_path, monkeypatch):
    class FakeProcess:
        pid = 4321

    class FakeTransport:
        _process = FakeProcess()

    class FakeOperatorClient:
        def __init__(self, options):
            self.options = options
            self._transport = FakeTransport()

        async def connect(self):
            return None

        async def query(self, _prompt):
            return None

        async def receive_messages(self):
            yield SystemMessage(subtype="init", data={
                "session_id": "safe-session", "model": "k3",
                "claude_code_version": "2.1.0", "permissionMode": "default",
                "tools": ["one", "two"], "mcp_servers": [{"name": "device"}],
                "agents": ["device-operator"], "skills": [], "plugins": [],
            })
            yield AssistantMessage(
                content=[ToolUseBlock(
                    "tool-1", "mcp__device__execute_ui_actions",
                    {"actions": [{"action": "tap_text", "target": "private-target"}]},
                )],
                model="k3", usage={"input_tokens": 120, "cache_read_input_tokens": 80},
                session_id="operator-session-1",
            )
            yield UserMessage(content=[ToolResultBlock(
                "tool-1",
                content=[
                    {"type": "text", "text": "private-result"},
                    {"type": "image", "data": "base64-private-image"},
                ],
                is_error=False,
            )])
            yield AssistantMessage(
                content=[TextBlock("SUITE_COMPLETED suite-1 pass 1/1 child-run")],
                model="k3", usage={"output_tokens": 30},
                session_id="operator-session-1",
            )
            yield ResultMessage(
                subtype="success", duration_ms=2500, duration_api_ms=2200,
                is_error=False, num_turns=2, session_id="operator-session-1",
                total_cost_usd=0.01,
                usage={
                    "input_tokens": 120, "output_tokens": 30,
                    "cache_read_input_tokens": 80,
                    "cache_creation_input_tokens": 10,
                },
            )

        async def disconnect(self):
            return None

    monkeypatch.setattr("harness_session.ClaudeSDKClient", FakeOperatorClient)
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.test_excel.suites["suite-1"] = type("State", (), {
        "status": "completed", "current_step": {"step_id": "case-1"},
    })()
    _activate_parent_turn(session)

    result = asyncio.run(session._run_device_operator({
        "suite_id": "suite-1",
        "current_step": {
            "step_id": "case-1", "operation": "private-operation",
            "observe": "private-observation",
        },
    }))

    assert result["status"] == "completed"
    phase_path = tmp_path / "attempt" / "trace" / "interactive" / "phase_timing.jsonl"
    phases = [json.loads(line) for line in phase_path.read_text(encoding="utf-8").splitlines()]
    names = [item["phase"] for item in phases]
    required_in_order = [
        "operator_started",
        "operator_connect_started", "operator_connect_completed",
        "operator_query_started", "model_request_submitted", "operator_query_submitted",
        "model_first_message", "operator_system_message",
        "model_first_assistant_message", "model_first_tool_use",
        "model_first_assistant_content", "operator_tool_use_emitted",
        "model_first_tool_result", "operator_tool_result_received",
        "model_first_text", "operator_text_emitted",
        "model_request_finished", "operator_result_received",
        "operator_disconnect_started", "operator_disconnect_completed",
    ]
    positions = [names.index(name) for name in required_in_order]
    assert positions == sorted(positions)
    assert names.count("model_last_message") == 5
    started = phases[0]
    assert started["scope"] == "device_operator"
    assert started["suite_id"] == "suite-1"
    assert started["prompt_chars"] > started["current_step_chars"] > 0
    assert started["operator_system_chars"] > 0
    assert started["operator_profile_chars"] > 0
    assert started["observer_cpu_ms"] >= 0

    first_message = next(item for item in phases if item["phase"] == "model_first_message")
    first_assistant = next(
        item for item in phases if item["phase"] == "model_first_assistant_message"
    )
    assert first_message["message_type"] == "SystemMessage"
    assert first_assistant["message_type"] == "AssistantMessage"
    assert first_assistant["model"] == "k3"
    assert first_assistant["input_tokens"] == 120
    assert first_assistant["cache_read_input_tokens"] == 80
    system_message = next(item for item in phases if item["phase"] == "operator_system_message")
    assert system_message["sdk_model"] == "k3"
    assert system_message["cli_version"] == "2.1.0"
    assert system_message["initialized_tool_count"] == 2
    assert system_message["initialized_mcp_server_count"] == 1
    assert system_message["initialized_agent_count"] == 1

    tool_use = next(item for item in phases if item["phase"] == "operator_tool_use_emitted")
    assert tool_use["tool_name"] == "execute_ui_actions"
    assert tool_use["tool_input_bytes"] > 0
    assert tool_use["step_id"] == "case-1"
    tool_result = next(
        item for item in phases if item["phase"] == "operator_tool_result_received"
    )
    assert tool_result["text_chars"] == len("private-result")
    assert tool_result["image_count"] == 1
    assert tool_result["image_encoded_chars"] == len("base64-private-image")
    assert tool_result["step_id"] == "case-1"

    serialized = "\n".join(json.dumps(item, ensure_ascii=False) for item in phases)
    assert "private-target" not in serialized
    assert "private-result" not in serialized
    assert "private-operation" not in serialized
    assert "private-observation" not in serialized

    results_path = tmp_path / "attempt" / "trace" / "interactive" / "results.jsonl"
    operator_result = json.loads(results_path.read_text(encoding="utf-8").splitlines()[-1])
    assert operator_result == {
        "type": "operator_result",
        "ts": operator_result["ts"],
        "suite_id": "suite-1",
        "route": "primary",
        "lease_generation": 1,
        "operator_session_fingerprint": operator_result["operator_session_fingerprint"],
        "duration_ms": 2500,
        "api_ms": 2200,
        "num_turns": 2,
        "cost_usd": 0.01,
        "usage": {
            "input_tokens": 120, "output_tokens": 30,
            "cache_read_input_tokens": 80,
            "cache_creation_input_tokens": 10,
        },
        "effort": "low",
    }


def test_operator_hook_boundaries_share_tool_use_id_without_business_payload(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    phase_path = tmp_path / "attempt" / "trace" / "interactive" / "phase_timing.jsonl"
    session._operator_phase_timer = PhaseTimer(
        str(phase_path),
        {"scope": "device_operator", "suite_id": "suite-1", "operator_id": "op-1"},
    )
    session._operator_cli_pid = 4321

    asyncio.run(session._hook({
        "hook_event_name": "PreToolUse", "tool_name": "mcp__device__find",
        "tool_use_id": "tool-1", "tool_input": {"text": "private-target"},
        "agent_id": "", "agent_type": "device-operator",
    }))
    asyncio.run(session._hook({
        "hook_event_name": "PostToolUseFailure", "tool_name": "mcp__device__find",
        "tool_use_id": "tool-1", "tool_input": {"text": "private-target"},
        "error": "private-error", "agent_id": "", "agent_type": "device-operator",
    }))

    records = [json.loads(line) for line in phase_path.read_text(encoding="utf-8").splitlines()]
    assert [record["phase"] for record in records] == [
        "operator_hook_pre_received", "operator_hook_failure_received",
    ]
    assert {record["tool_use_id"] for record in records} == {"tool-1"}
    assert {record["tool_name"] for record in records} == {"find"}
    serialized = json.dumps(records, ensure_ascii=False)
    assert "private-target" not in serialized
    assert "private-error" not in serialized


def test_operator_can_run_over_thirty_minutes_when_progress_keeps_refreshing(
        tmp_path, monkeypatch):
    import harness_session as module

    class FakeClock:
        def __init__(self):
            self.value = 0.0

        def __call__(self):
            return self.value

        def advance(self, seconds):
            self.value += seconds

    clock = FakeClock()

    class FakeOperatorClient:
        def __init__(self, options):
            self.options = options

        async def connect(self):
            clock.advance(10 * 60)

        async def query(self, _prompt):
            clock.advance(10 * 60)

        async def receive_messages(self):
            for index in range(3):
                clock.advance(10 * 60)
                yield AssistantMessage(
                    content=[TextBlock(f"progress-{index}")],
                    model="k3", usage={}, session_id="operator-session-1",
                )
            clock.advance(10 * 60)
            yield ResultMessage(
                subtype="success", duration_ms=1, duration_api_ms=1,
                is_error=False, num_turns=1, session_id="operator-session-1",
                total_cost_usd=0.0, usage={},
            )

        async def disconnect(self):
            return None

    monkeypatch.setattr(module, "ClaudeSDKClient", FakeOperatorClient)
    monkeypatch.setattr(module, "_monotonic", clock)
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.test_excel.suites["suite-1"] = type("State", (), {"status": "completed"})()
    _activate_parent_turn(session)

    result = asyncio.run(session._run_device_operator({
        "suite_id": "suite-1", "current_step": {"step_id": "case-1"},
    }))

    assert clock.value > 30 * 60
    assert result["status"] == "completed"


def test_operator_is_interrupted_after_continuous_inactivity(tmp_path, monkeypatch):
    import harness_session as module

    instances = []

    class FakeOperatorClient:
        def __init__(self, options):
            self.options = options
            self.interrupted = False
            instances.append(self)

        async def connect(self):
            return None

        async def query(self, _prompt):
            return None

        async def receive_messages(self):
            await asyncio.Future()
            if False:
                yield None

        async def interrupt(self):
            self.interrupted = True

        async def disconnect(self):
            return None

    monkeypatch.setattr(module, "ClaudeSDKClient", FakeOperatorClient)
    monkeypatch.setattr(module, "_OPERATOR_INACTIVITY_TIMEOUT_SECONDS", 0.02)
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    _activate_parent_turn(session)

    result = asyncio.run(asyncio.wait_for(session._run_device_operator({
        "suite_id": "suite-1", "current_step": {"step_id": "case-1"},
    }), timeout=0.5))

    assert result["status"] == "inactivity_timeout"
    assert result["reason"] == "operator_inactivity_timeout"
    assert instances[0].interrupted is True


def test_operator_tool_events_refresh_inactivity_deadline(tmp_path, monkeypatch):
    import harness_session as module

    class FakeOperatorClient:
        queue = None

        def __init__(self, options):
            self.options = options
            self.queue = asyncio.Queue()

        async def connect(self):
            return None

        async def query(self, _prompt):
            return None

        async def receive_messages(self):
            yield await self.queue.get()

        async def disconnect(self):
            return None

    async def scenario():
        monkeypatch.setattr(module, "ClaudeSDKClient", FakeOperatorClient)
        monkeypatch.setattr(module, "_OPERATOR_INACTIVITY_TIMEOUT_SECONDS", 0.03)
        session = HarnessSession(
            _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
            provider_env={}, effort="low",
        )
        session.test_excel.suites["suite-1"] = type(
            "State", (), {"status": "completed"},
        )()
        _activate_parent_turn(session)
        runner = asyncio.create_task(session._run_device_operator({
            "suite_id": "suite-1", "current_step": {"step_id": "case-1"},
        }))
        while session._operator_client is None:
            await asyncio.sleep(0)
        for index in range(4):
            await asyncio.sleep(0.015)
            await session._hook({
                "hook_event_name": "PostToolUse",
                "tool_name": "mcp__test_excel__record_test_step",
                "tool_use_id": f"record-{index}", "tool_response": {},
                "agent_id": "", "agent_type": "device-operator",
            })
        await session._operator_client.queue.put(ResultMessage(
            subtype="success", duration_ms=1, duration_api_ms=1,
            is_error=False, num_turns=1, session_id="operator-session-1",
            total_cost_usd=0.0, usage={},
        ))
        return await asyncio.wait_for(runner, timeout=0.2)

    result = asyncio.run(scenario())
    assert result["status"] == "completed"


def test_user_stop_wakes_operator_watchdog_without_advancing_registry_cursor(
        tmp_path, monkeypatch):
    import harness_session as module

    class Control:
        stopped = False
        stop_reason = ""

    class FakeOperatorClient:
        connected = None

        def __init__(self, options):
            self.options = options
            self.connected = asyncio.Event()
            self.interrupted = False

        async def connect(self):
            self.connected.set()

        async def query(self, _prompt):
            return None

        async def receive_messages(self):
            await asyncio.Future()
            if False:
                yield None

        async def interrupt(self):
            self.interrupted = True

        async def disconnect(self):
            return None

    class FakeMainClient:
        async def interrupt(self):
            return None

    async def scenario():
        monkeypatch.setattr(module, "ClaudeSDKClient", FakeOperatorClient)
        control = Control()
        session = HarnessSession(
            _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
            provider_env={}, effort="low", control=control,
        )
        registry_step = {"step_id": "case-2", "operation": "继续当前用例"}
        state = SimpleNamespace(
            suite_id="suite-1", status="active", current_step=registry_step,
            ready_to_finish=False, terminal_reason="",
        )
        session.test_excel.active_suite_id = "suite-1"
        session.test_excel.suites["suite-1"] = state
        monkeypatch.setattr(session.test_excel, "_persist_state", lambda _state: None)

        async def emit_state(_state, _event):
            return None

        monkeypatch.setattr(session.test_excel, "_emit", emit_state)
        session.client = FakeMainClient()
        _activate_parent_turn(session)
        runner = asyncio.create_task(session._run_device_operator({
            "suite_id": "suite-1", "current_step": registry_step,
        }))
        while session._operator_client is None:
            await asyncio.sleep(0)
        await session._operator_client.connected.wait()

        control.stopped = True
        control.stop_reason = "user_stop"
        await asyncio.wait_for(session.stop(grace_timeout_s=0), timeout=0.2)
        cancelled = False
        try:
            await asyncio.wait_for(runner, timeout=0.2)
        except asyncio.CancelledError:
            cancelled = True

        assert cancelled is True
        assert state.status == "active"
        assert state.current_step == registry_step

    asyncio.run(scenario())


def test_resume_cursor_result_deterministically_resumes_operator(tmp_path, monkeypatch):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    delegated = []

    async def run_operator(payload, *, resume=False):
        delegated.append((payload["suite_id"], resume))
        return {"ok": True, "status": "paused", "suite_id": payload["suite_id"]}

    monkeypatch.setattr(session, "_run_device_operator", run_operator)
    # next_test_step currently returns the resumed step without repeating suite_id;
    # the authoritative identity is still present in the tool input.
    response = {"content": [{"type": "text", "text": json.dumps({
        "status": "resume_current",
        "step_id": "case-2", "operation": "点击 B", "observe": "进入 B 页",
    })}]}

    asyncio.run(session._hook({
        "hook_event_name": "PostToolUse",
        "tool_name": "mcp__test_excel__next_test_step",
        "tool_use_id": "next-1", "tool_input": {"suite_id": "suite-1"},
        "tool_response": response,
        "agent_id": "", "agent_type": "",
    }))

    assert delegated == [("suite-1", True)]


def test_selected_top_level_operator_is_attributed_as_device_operator(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    noted = []
    session.test_excel.note_tool_actor = lambda name, actor, agent_id: noted.append(
        (name, actor, agent_id)
    )

    asyncio.run(session._hook({
        "hook_event_name": "PreToolUse", "tool_use_id": "next-1",
        "tool_name": "mcp__test_excel__next_test_step", "tool_input": {},
        "agent_id": "", "agent_type": "device-operator",
    }))

    assert noted == [(
        "mcp__test_excel__next_test_step", "device_operator", "deterministic-operator",
    )]


def test_native_hook_denies_device_action_until_suite_cursor_is_revalidated(tmp_path):
    emitted = []

    async def emit(event):
        emitted.append(event)

    session = HarnessSession(
        emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.test_excel.device_action_block_reason = lambda: (
        "恢复测试后必须先调用同一 suite_id 的 next_test_step"
    )

    decision = asyncio.run(session._hook({
        "hook_event_name": "PreToolUse", "tool_use_id": "tool-device",
        "tool_name": "mcp__device__get_screen", "tool_input": {},
        "agent_id": "",
    }))

    output = decision["hookSpecificOutput"]
    assert output["permissionDecision"] == "deny"
    assert "next_test_step" in output["permissionDecisionReason"]
    assert emitted[-1]["denied"] is True


def test_native_hook_denies_online_data_refresh_during_active_suite(tmp_path):
    observed = []
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.test_excel.active_suite_id = "suite-active"

    async def observe(*args):
        observed.append(args)

    session.test_excel.observe_device_action = observe
    decision = asyncio.run(session._hook({
        "hook_event_name": "PreToolUse", "tool_use_id": "tool-refresh",
        "tool_name": "mcp__device__prepare_test_data", "tool_input": {},
        "agent_id": "",
    }))

    output = decision["hookSpecificOutput"]
    assert output["permissionDecision"] == "deny"
    assert "active Suite" in output["permissionDecisionReason"]
    assert observed == []


def test_continue_requires_next_before_observing_device(tmp_path):
    source = tmp_path / "cases.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "用例"
    ws.append(["序号", "操作", "观测点", "观测效果"])
    ws.append([1, "点击 A", "显示结果", ""])
    wb.save(source)

    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    imported = session.workspace.import_file(source)
    opened = asyncio.run(session.test_excel.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    suite_id = opened["suite_id"]
    asyncio.run(session.test_excel.next({"suite_id": suite_id}))

    injected = asyncio.run(session.prepare_turn("continue_active_suite"))

    assert "先调用同一 suite_id 的 next_test_step" in injected
    assert "next_test_step" in session.test_excel.device_action_block_reason()
    resumed = asyncio.run(session.test_excel.next({"suite_id": suite_id}))
    assert resumed["status"] == "resume_current"
    assert session.test_excel.device_action_block_reason() == ""
    asyncio.run(session.test_excel.finalize_active("test cleanup", status="interrupted"))


def test_start_new_suite_choice_defers_old_suite_transition_until_open_preflight(tmp_path):
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.test_excel.active_suite_id = "suite-active"
    finalized = []

    async def finalize(*args, **kwargs):
        finalized.append((args, kwargs))

    session.test_excel.finalize_active = finalize
    prefix = asyncio.run(session.prepare_turn("start_new_suite"))

    assert finalized == []
    assert session.test_excel.active_suite_id == "suite-active"
    assert "先预检新 Excel" in prefix


def test_attachment_is_copied_into_run_input(tmp_path):
    upload = tmp_path / "brief.txt"
    upload.write_text("compare products", encoding="utf-8")
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    imported = session.import_attachment({"filename": "brief.txt"})

    assert imported["sha256"]
    assert imported["path"].startswith(str(session.workspace.input_dir))


def test_attachment_preserves_sheet_metadata_for_model_selection(tmp_path):
    upload = tmp_path / "cases.xlsx"
    upload.write_bytes(b"xlsx-placeholder")
    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )

    imported = session.import_attachment({
        "filename": "cases.xlsx",
        "sheets": ["执行须知", "v1创建任务上传身份证刷脸通过"],
    })

    assert imported["sheets"] == ["执行须知", "v1创建任务上传身份证刷脸通过"]


def test_send_injects_available_sheets_and_v1_selection_hint(tmp_path):
    upload = tmp_path / "cases.xlsx"
    upload.write_bytes(b"xlsx-placeholder")
    captured = []

    class FakeClient:
        async def query(self, text):
            captured.append(text)

        async def receive_response(self):
            if False:
                yield None

    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.client = FakeClient()

    asyncio.run(session.send("执行v1", file={
        "filename": "cases.xlsx",
        "sheets": ["执行须知", "v1创建任务上传身份证刷脸通过"],
    }))

    assert "v1创建任务上传身份证刷脸通过" in captured[0]
    assert "名称以 v1 开头" in captured[0]


def test_native_harness_forwards_all_tool_result_text_to_chat(tmp_path):
    class FakeClient:
        async def query(self, _text):
            return None

        async def receive_response(self):
            yield AssistantMessage(
                content=[ToolUseBlock("tool-1", "mcp__device__get_screen", {})],
                model="k3", usage={},
            )
            yield UserMessage(content=[ToolResultBlock(
                "tool-1", content=[{"type": "text", "text": "子 Agent 完整输出"}],
                is_error=False,
            )])
            yield ResultMessage(
                subtype="success", duration_ms=1, duration_api_ms=1,
                is_error=False, num_turns=1, session_id="session-1",
                total_cost_usd=0.0, usage={},
            )

    async def scenario():
        emitted = []

        async def emit(event):
            emitted.append(event)

        session = HarnessSession(
            emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
            provider_env={}, effort="low",
        )
        session.client = FakeClient()
        await session.send("执行测试")
        return emitted

    emitted = asyncio.run(scenario())
    assert [event["text"] for event in emitted if event["type"] == "chat_delta"] == [
        "子 Agent 完整输出",
    ]
    completed = [event for event in emitted if event["type"] == "tool_event"][-1]
    assert completed["status"] == "ok"


def test_native_harness_records_complete_model_request_lifecycle(tmp_path):
    class FakeClient:
        async def query(self, _text):
            return None

        async def receive_response(self):
            yield AssistantMessage(content=[], model="k3", usage={})
            yield ResultMessage(
                subtype="success", duration_ms=1, duration_api_ms=1,
                is_error=False, num_turns=1, session_id="session-1",
                total_cost_usd=0.0, usage={},
            )

    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.client = FakeClient()
    phase_path = tmp_path / "attempt" / "trace" / "interactive" / "phase_timing.jsonl"
    timer = PhaseTimer(str(phase_path), {"scope": "interactive_turn", "turn_id": "t1"})

    asyncio.run(session.send("执行测试", phase_timer=timer))

    records = [json.loads(line) for line in phase_path.read_text(encoding="utf-8").splitlines()]
    lifecycle = [record for record in records if record["phase"].startswith("model_")]
    assert [record["phase"] for record in lifecycle] == [
        "model_request_submitted", "model_first_message", "model_first_assistant_message",
        "model_last_message",
        "model_last_message", "model_request_finished",
    ]
    assert len({record["request_id"] for record in lifecycle}) == 1
    assert lifecycle[-1]["end_reason"] == "result"


def test_native_suite_append_is_delivered_to_device_operator_not_parent(tmp_path):
    class FakeClient:
        def __init__(self):
            self.queries = []

        async def query(self, text):
            self.queries.append(text)

    emitted = []

    async def emit(event):
        emitted.append(event)

    async def scenario():
        session = HarnessSession(
            emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
            provider_env={}, effort="low",
        )
        session.client = FakeClient()
        session._accepting_input = True
        session.test_excel.active_suite_id = "suite-1"
        session.permission_policy.device_actor_id = "operator-1"
        session.test_excel.note_tool_actor = lambda *_args: None
        command = {
            "text": "请先打开系统设置校验位置权限",
            "client_message_id": "message-1", "input_seq": 1,
            "turn_id": "turn-1", "suite_id": "suite-1", "step_id": "setup-0",
        }

        assert await session.append_input(command) is True
        assert session.client.queries == []

        parent_hook_output = await session._hook({
            "hook_event_name": "PreToolUse", "tool_name": "Read",
            "tool_input": {"file_path": str(session.workspace.input_dir)},
            "tool_use_id": "tool-parent", "agent_id": "",
        })
        assert parent_hook_output == {}

        hook_output = await session._hook({
            "hook_event_name": "PreToolUse",
            "tool_name": "mcp__test_excel__next_test_step",
            "tool_input": {}, "tool_use_id": "tool-1", "agent_id": "operator-1",
        })
        context = hook_output["hookSpecificOutput"]["additionalContext"]
        assert "用户刚刚追加了以下要求" in context
        assert "请先打开系统设置校验位置权限" in context
        submitted = [
            event for event in emitted
            if event.get("type") == "turn_input" and event.get("status") == "submitted"
        ]
        assert len(submitted) == 1

    asyncio.run(scenario())


def test_native_explicit_model_failure_stops_and_preserves_active_step(tmp_path):
    source = tmp_path / "cases.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "用例"
    ws.append(["序号", "操作", "观测点", "观测效果"])
    ws.append([1, "点击 A", "显示结果", ""])
    wb.save(source)

    class Control:
        def __init__(self):
            self.stopped = False
            self.stop_reason = ""

    class FakeClient:
        def __init__(self):
            self.interrupted = False

        async def query(self, _text):
            return None

        async def interrupt(self):
            self.interrupted = True

        async def receive_response(self):
            yield AssistantMessage(
                content=[ToolUseBlock("agent-1", "Agent", {})], model="k3", usage={},
            )
            yield UserMessage(content=[ToolResultBlock(
                "agent-1",
                content="API Error: 400 litellm.BadRequestError: There are no healthy deployments for this model",
                is_error=False,
            )])

    emitted = []

    async def emit(event):
        emitted.append(event)

    control = Control()
    session = HarnessSession(
        emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low", control=control,
    )
    imported = session.workspace.import_file(source)
    opened = asyncio.run(session.test_excel.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    suite_id = opened["suite_id"]
    asyncio.run(session.test_excel.next({"suite_id": suite_id}))
    current_step = dict(session.test_excel.suites[suite_id].current_step)
    client = FakeClient()
    session.client = client

    try:
        asyncio.run(session.send("执行测试"))
        assert False, "provider failure must stop the native turn"
    except TurnStopped:
        pass

    assert client.interrupted is True
    assert session.test_excel.suites[suite_id].status == "active"
    assert session.test_excel.suites[suite_id].current_step == current_step
    # The supervisor owns the single turn_stopped event; the session raises a
    # structured stop without inventing an additional pause state.
    asyncio.run(session.test_excel.finalize_active("test cleanup", status="interrupted"))


def test_sensitive_builtin_tool_inputs_are_not_exposed_in_ui_summary():
    assert "secret" not in safe_harness_args("Bash", {"command": "echo secret"})
    assert "body" not in safe_harness_args("Write", {"content": "body"})
    assert "token" not in safe_harness_args("WebFetch", {"url": "https://x/?token=abc"})


def test_harness_stop_interrupts_turn_without_finalizing_active_suite(tmp_path):
    class FakeClient:
        def __init__(self):
            self.interrupted = False

        async def interrupt(self):
            self.interrupted = True

    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    fake = FakeClient()
    session.client = fake
    finalized = []

    async def forbidden_finalize(*args, **kwargs):
        finalized.append((args, kwargs))

    session.test_excel.finalize_active = forbidden_finalize
    asyncio.run(session.stop(grace_timeout_s=0.05))

    assert fake.interrupted is True
    assert finalized == []


def test_harness_stop_interrupts_deterministic_operator_and_main_turn(tmp_path):
    class FakeClient:
        def __init__(self):
            self.interrupted = False

        async def interrupt(self):
            self.interrupted = True

    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    main = FakeClient()
    operator = FakeClient()
    session.client = main
    session._operator_client = operator

    asyncio.run(session.stop(grace_timeout_s=0.05))

    assert main.interrupted is True
    assert operator.interrupted is True


def test_native_ask_user_question_returns_answers_keyed_by_question(tmp_path):
    prompts = []

    async def ask(prompt):
        prompts.append(prompt)
        return "继续"

    session = HarnessSession(
        _emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low", ask_user=ask,
    )
    session.test_excel.active_suite_id = "suite-1"
    result = asyncio.run(session._can_use_tool("AskUserQuestion", {
        "questions": [{
            "question": "是否继续？", "header": "确认", "multiSelect": False,
            "options": [{"label": "继续", "description": "执行下一步"}],
        }],
    }, None))

    assert result.updated_input["answers"] == {"是否继续？": "继续"}
    assert "继续" in prompts[0]
    queued = session._operator_inputs.get_nowait()
    assert queued["suite_id"] == "suite-1"
    assert "是否继续？" in queued["text"]
    assert "继续" in queued["text"]


def test_effort_change_reconnects_with_public_resume_and_same_workspace(monkeypatch, tmp_path):
    import harness_session as module
    instances = []

    class FakeClient:
        def __init__(self, options):
            self.options = options
            self.connected = False
            self.disconnected = False
            instances.append(self)

        async def connect(self): self.connected = True
        async def disconnect(self): self.disconnected = True

    emitted = []

    async def emit(event): emitted.append(event)

    monkeypatch.setattr(module, "ClaudeSDKClient", FakeClient)
    session = HarnessSession(
        emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    asyncio.run(session.connect())
    old_workspace = session.workspace.root
    session.session_id = "sdk-session-1"

    asyncio.run(session.change_effort("high"))

    assert instances[0].disconnected is True
    assert instances[1].options.resume == "sdk-session-1"
    assert instances[1].options.effort == "high"
    assert session.workspace.root == old_workspace
    assert emitted[-1]["transport_reconnected"] is True


def test_provider_reconfigure_preserves_sdk_session(monkeypatch, tmp_path):
    import harness_session as module
    instances = []

    class OldClient:
        async def disconnect(self):
            return None

    class NewClient:
        def __init__(self, options=None):
            self.options = options
            instances.append(self)

        async def connect(self):
            return None

    async def emit(_event):
        return None

    monkeypatch.setattr(module, "ClaudeSDKClient", NewClient)
    session = HarnessSession(
        emit, attempt_dir=str(tmp_path / "attempt"), uploads_dir=str(tmp_path),
        provider_env={}, effort="low",
    )
    session.client = OldClient()
    session.session_id = "sdk-session-provider"

    asyncio.run(session.reconfigure_provider({"ANTHROPIC_API_KEY": "next-key"}, model="k3"))

    assert session.session_id == "sdk-session-provider"
    assert instances[-1].options.resume == "sdk-session-provider"
