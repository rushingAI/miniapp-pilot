"""run_ui_steps：固定 UI 流程执行器（按步骤跑 + 逐步校验 + 失败即停）与 uiscripts 读表。
纯逻辑单测——stub 掉点击/输入/读屏/读窗口，不碰真机。
重点守：失败即停不硬闯、不重试；代码零 app 知识（流程全在 xlsx）。"""
import asyncio
import json
import os
import re
import sys
import inspect

import pytest
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
import server  # noqa: E402
import tools  # noqa: E402
import uiscripts  # noqa: E402

FAST = {"poll_ms": 1, "default_timeout_s": 0.05}


def test_execute_ui_actions_interface_teaches_the_supported_plan_language():
    action = tools.execute_ui_actions.input_schema["properties"]["actions"]["items"]["properties"]["action"]
    target = tools.execute_ui_actions.input_schema["properties"]["actions"]["items"]["properties"]["target"]

    assert action["enum"] == [
        "tap_text", "tap_text_retry", "tap_xy", "input", "clear_input",
        "swipe_up", "swipe_down", "press_key", "wait", "ok_if_text",
        "launch_app", "stop_app", "tap_image",
    ]
    assert "xy:540,1746" in target["description"]
    assert "pct:50,75" in target["description"]
    item = tools.execute_ui_actions.input_schema["properties"]["actions"]["items"]
    assert item["properties"]["allow_bottom_edge"]["type"] == "boolean"
    assert "有副作用动作" in tools.execute_ui_actions.description
    assert "不返回内联图片" in tools.execute_ui_actions.description
    assert "语义条件优先" in tools.execute_ui_actions.description
    assert "视觉保底" in tools.execute_ui_actions.description
    assert "count=1" in tools.find.description
    assert "matches[0].c" in tools.find.description
    assert "unrel" in tools.find.description


def test_execute_ui_actions_schema_exposes_one_to_eight_actions_with_flat_postconditions():
    schema = tools.execute_ui_actions.input_schema
    actions = schema["properties"]["actions"]
    item_properties = actions["items"]["properties"]

    assert actions["minItems"] == 1
    assert actions["maxItems"] == 8
    assert "settle_ms" not in schema["properties"]
    assert "expect" not in item_properties
    assert schema["properties"]["final_assertion"]["type"] == "boolean"
    postcondition = item_properties["postcondition"]
    assert postcondition["properties"]["mode"]["enum"] == ["all", "any"]
    assert postcondition["properties"]["checks"]["minItems"] == 0
    assert postcondition["properties"]["checks"]["maxItems"] == 4
    check = postcondition["properties"]["checks"]["items"]
    assert check["properties"]["type"]["enum"] == ["a11y", "focus", "screen_change"]
    assert check["properties"]["selector"]["enum"] == ["text", "resource_id", "content_desc"]
    assert check["properties"]["match"]["enum"] == ["exact", "contains", "suffix"]
    assert check["properties"]["state"]["enum"] == ["present", "absent"]


def test_execute_ui_steps_requires_every_check_in_an_all_postcondition(monkeypatch):
    monkeypatch.setattr(server, "get_focused_app", lambda: "org.example.tasklist/Main")
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [
            {"text": "确定", "content_desc": "", "resource_id": "confirm",
             "center": [540, 1800], "bounds": [400, 1740, 680, 1860],
             "bounds_unreliable": False},
            {"text": "任务列表（2项）", "content_desc": "", "resource_id": "title",
             "center": [540, 120], "bounds": [0, 80, 1080, 160],
             "bounds_unreliable": False},
        ],
    ))
    monkeypatch.setattr(server, "tap_element", lambda **_kwargs: {
        "ok": True, "matched": 1, "clicked_index": 0,
    })

    result = server.execute_ui_steps([{
        "action": "tap_text",
        "target": "~:确定",
        "postcondition": {
            "mode": "all", "settle_ms": 0, "timeout_s": 0.05,
            "checks": [
                {"type": "a11y", "selector": "text", "value": "任务列表（2项）",
                 "match": "contains", "state": "present"},
                {"type": "focus", "value": "org.example.tasklist", "state": "present"},
            ],
        },
    }], poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    assert result["done_steps"] == 1
    step = result["log"][0]
    assert step["action_executed"] is True
    assert step["postcondition_satisfied"] is True
    assert [check["satisfied"] for check in step["postcondition"]["checks"]] == [True, True]


def test_execute_ui_steps_interrupts_inside_a_long_postcondition(monkeypatch):
    stop_checks = {"count": 0}

    def should_stop():
        stop_checks["count"] += 1
        return stop_checks["count"] > 1

    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (1080, 2340, []))

    started = __import__("time").perf_counter()
    result = server.execute_ui_steps([{
        "action": "wait",
        "postcondition": {
            "mode": "all", "settle_ms": 0, "timeout_s": 0.2,
            "checks": [{"type": "a11y", "selector": "text", "value": "完成",
                        "match": "contains", "state": "present"}],
        },
    }], poll_ms=10, settle_ms=0, should_stop=should_stop)

    assert __import__("time").perf_counter() - started < 0.1
    assert result["ok"] is False
    assert result["stopped"] is True
    assert result["done_steps"] == 0


def test_private_stop_callbacks_do_not_expand_public_device_signatures():
    for function in (
        server.tap_wait_tap, server.wait_digits_leave, server.repeat_swipe,
        server.drag, server.tap_digits,
    ):
        assert "should_stop" not in inspect.signature(function).parameters


def test_tool_stop_callback_stays_stopped_after_continue_resets_boolean(monkeypatch):
    control = type("Control", (), {"stopped": False, "control_seq": 7})()
    monkeypatch.setattr(tools, "CONTROL", control)
    should_stop = tools._stop_callback()

    control.stopped = True
    control.control_seq = 8
    assert should_stop() is True

    control.stopped = False
    assert should_stop() is True


def test_repeat_swipe_stop_preserves_completed_prefix(monkeypatch):
    calls = []
    checks = {"count": 0}

    def should_stop():
        checks["count"] += 1
        return checks["count"] > 2

    monkeypatch.setattr(server, "_wm_size", lambda: (1000, 2000))
    monkeypatch.setattr(server, "_adb", lambda *args, **_kwargs: calls.append(args))

    result = server._repeat_swipe_impl(
        count=4, duration_ms=10, interval_ms=20, should_stop=should_stop,
    )

    assert result["ok"] is False
    assert result["stopped"] is True
    assert result["completed_count"] == 1
    assert len(calls) == 1


def test_tap_digits_stop_never_sends_the_next_digit(monkeypatch):
    taps = []
    checks = {"count": 0}

    def should_stop():
        checks["count"] += 1
        return checks["count"] > 2

    elements = [
        {"text": digit, "clickable": True, "bounds_unreliable": False,
         "center": [int(digit) + 10, 100]}
        for digit in "0123456789"
    ]
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (100, 200, elements))
    monkeypatch.setattr(server, "_adb", lambda *args, **_kwargs: taps.append(args))

    result = server._tap_digits_impl("1234", should_stop=should_stop)

    assert result["ok"] is False
    assert result["stopped"] is True
    assert result["completed_count"] == 1
    assert len(taps) == 1


def test_execute_ui_steps_resolves_a_later_text_target_after_the_previous_action(monkeypatch):
    page = {"name": "home"}
    taps = []

    def elements(**_kwargs):
        texts = {
            "home": ["任务"],
            "tasks": ["收藏"],
            "favorites": ["收藏页"],
        }[page["name"]]
        return 1080, 2340, [
            {"text": text, "content_desc": "", "resource_id": "",
             "center": [540, 120], "bounds_unreliable": False}
            for text in texts
        ]

    def tap_element(**kwargs):
        text = kwargs.get("text")
        taps.append(text)
        if text == "任务" and page["name"] == "home":
            page["name"] = "tasks"
        elif text == "收藏" and page["name"] == "tasks":
            page["name"] = "favorites"
        else:
            return {"ok": False, "error": "target not on current page"}
        return {"ok": True, "matched": 1, "clicked_index": 0}

    monkeypatch.setattr(server, "_screen_elements", elements)
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "tap_element", tap_element)

    result = server.execute_ui_steps([
        {
            "action": "tap_text", "target": "任务",
            "postcondition": {
                "mode": "all", "settle_ms": 0, "timeout_s": 0.05,
                "checks": [{"type": "a11y", "selector": "text", "value": "收藏",
                            "match": "exact", "state": "present"}],
            },
        },
        {
            "action": "tap_text", "target": "收藏",
            "postcondition": {
                "mode": "all", "settle_ms": 0, "timeout_s": 0.05,
                "checks": [{"type": "a11y", "selector": "text", "value": "收藏页",
                            "match": "exact", "state": "present"}],
            },
        },
    ], poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    assert result["done_steps"] == 2
    assert taps == ["任务", "收藏"]


def test_execute_ui_actions_wait_uses_any_postcondition_instead_of_expect(monkeypatch):
    captured = []
    monkeypatch.setattr(server, "execute_ui_steps", lambda actions, *_args: (
        captured.extend(actions) or {
            "ok": True, "done_steps": 1, "total_steps": 1, "log": [],
        }
    ))

    result = json.loads(asyncio.run(tools.execute_ui_actions.handler({
        "actions": [{
            "action": "wait",
            "postcondition": {
                "mode": "any", "settle_ms": 0, "timeout_s": 3,
                "checks": [
                    {"type": "a11y", "selector": "text", "value": "完成",
                     "match": "contains", "state": "present"},
                    {"type": "a11y", "selector": "resource_id", "value": "done",
                     "match": "suffix", "state": "present"},
                ],
            },
        }],
        "settle_ms": 0,
    }))["content"][0]["text"])

    assert result["ok"] is True
    assert captured[0]["postcondition"]["mode"] == "any"
    assert "expect" not in captured[0]


def test_execute_ui_steps_combines_explicit_screen_change_with_a11y(tmp_path, monkeypatch):
    captures = []

    def save(path):
        color = (0, 0, 0) if not captures else (255, 255, 255)
        Image.new("RGB", (20, 20), color).save(path)
        captures.append(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    monkeypatch.setattr(server, "_adb", lambda *_args: None)
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "目标页", "content_desc": "", "resource_id": "title",
            "center": [540, 120], "bounds": [0, 80, 1080, 160],
            "bounds_unreliable": False,
        }],
    ))

    result = server.execute_ui_steps([{
        "action": "tap_xy", "target": "xy:10,20",
        "postcondition": {
            "mode": "all", "settle_ms": 0, "timeout_s": 0.05,
            "checks": [
                {"type": "screen_change", "diff_threshold": 3},
                {"type": "a11y", "selector": "text", "value": "目标页",
                 "match": "exact", "state": "present"},
            ],
        },
    }], poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    check = result["log"][0]["postcondition"]["checks"][0]
    assert check["satisfied"] is True
    assert check["diff"] > 3
    assert len(captures) >= 2


def test_execute_ui_steps_tap_text_forwards_the_selected_match_index(monkeypatch):
    seen = []
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "产品", "content_desc": "", "resource_id": "product",
            "center": [200, 300], "bounds": [100, 250, 300, 350],
            "bounds_unreliable": False,
        }],
    ))
    monkeypatch.setattr(server, "tap_element", lambda **kwargs: (
        seen.append(kwargs) or {"ok": True, "matched": 3, "clicked_index": kwargs["index"]}
    ))

    result = server.execute_ui_steps([
        {"action": "tap_text", "target": "产品", "index": 2},
    ], poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    assert seen == [{"text": "产品", "contains": False, "index": 2}]


def test_execute_ui_steps_clear_input_forwards_max_chars_and_returns_lengths(monkeypatch):
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "clear_input", lambda **kwargs: {
        "ok": True, "before_len": 12, "after_len": 0, "max_chars": kwargs["max_chars"],
    })

    result = server.execute_ui_steps([
        {"action": "clear_input", "target": "rid:account", "max_chars": 64},
    ], poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    assert result["log"][0]["action_result"] == {
        "ok": True, "before_len": 12, "after_len": 0, "max_chars": 64,
    }


def test_execute_ui_steps_launch_app_accepts_a_visible_app_name_target(monkeypatch):
    calls = []
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "launch_app", lambda **kwargs: (
        calls.append(kwargs) or {"ok": True, "package": "com.example", "via": "launcher_text"}
    ))

    result = server.execute_ui_steps([
        {"action": "launch_app", "target": "app:示例应用"},
    ], poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    assert calls == [{"app_name": "示例应用", "timeout_s": 5.0}]
    assert result["log"][0]["action_result"]["via"] == "launcher_text"


@pytest.mark.parametrize("target", [
    "xy:-1,10", "xy:nan,10", "xy:inf,10", "xy:1080,10",
    "pct:-1,50", "pct:50,101",
])
def test_execute_ui_actions_rejects_every_invalid_tap_xy_before_execution(target, monkeypatch):
    monkeypatch.setattr(server, "_wm_size", lambda: (1080, 2340))
    monkeypatch.setattr(server, "execute_ui_steps", lambda *_args: (_ for _ in ()).throw(
        AssertionError("预检失败后不能进入执行器")
    ))

    result = json.loads(asyncio.run(tools.execute_ui_actions.handler({
        "actions": [{"action": "tap_xy", "target": target}],
        "settle_ms": 0,
    }))["content"][0]["text"])

    assert result["preflight_rejected"] is True
    assert result["done_steps"] == 0
    assert result["failed_at"]["action_executed"] is False


def test_execute_ui_actions_returns_an_after_frame_without_a_run_trace(monkeypatch):
    tools.set_trace_dir(None)
    monkeypatch.setattr(server, "execute_ui_steps", lambda actions, *_args: {
        "ok": True, "done_steps": len(actions), "total_steps": len(actions), "log": [{
            "i": 1, "action": "press_key", "action_executed": True,
            "postcondition_satisfied": True,
        }],
    })

    def save(path):
        Image.new("RGB", (108, 234), (20, 40, 60)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)

    tool_result = asyncio.run(tools.execute_ui_actions.handler({
        "actions": [{"action": "press_key", "target": "back"}],
        "settle_ms": 0,
    }))
    result = json.loads(tool_result["content"][0]["text"])

    assert result["ok"] is True
    assert result["observation_source"] == "batch_after"
    assert any(item["type"] == "image" for item in tool_result["content"])
    assert result["evidence_paths"] == [result["path"]]
    assert result["steps"][0]["action_executed"] is True
    assert result["steps"][0]["postcondition_satisfied"] is True


def test_execute_ui_actions_tap_image_without_a_current_frame_returns_one(monkeypatch):
    tools.set_trace_dir(None)
    tools.LAST_SCREENSHOT_META = None
    monkeypatch.setattr(server, "execute_ui_steps", lambda *_args: (_ for _ in ()).throw(
        AssertionError("缺少坐标来源画面时不得执行动作")
    ))

    def save(path):
        Image.new("RGB", (108, 234), (30, 60, 90)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    try:
        tool_result = asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{"action": "tap_image", "target": "20,100"}],
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.LAST_SCREENSHOT_META = None

    assert result["preflight_rejected"] is True
    assert result["observation_source"] == "failure_current_screen"
    assert any(item["type"] == "image" for item in tool_result["content"])


def test_execute_ui_actions_persists_the_last_postcondition_as_final_assertion(
        tmp_path, monkeypatch):
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "任务列表（2项）", "content_desc": "", "resource_id": "title",
            "center": [540, 120], "bounds": [0, 80, 1080, 160],
            "bounds_unreliable": False,
        }],
    ))

    def save(path):
        Image.new("RGB", (108, 234), (20, 40, 60)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    async def no_sleep(_seconds):
        return None
    monkeypatch.setattr(tools.asyncio, "sleep", no_sleep)
    tools.set_trace_dir(str(tmp_path))
    tools.LAST_SCREENSHOT_META = {
        "device_wh": [1080, 2340], "image_wh": [108, 234],
        "scale": 10.0, "path": str(tmp_path / "stale.png"),
    }
    try:
        tool_result = asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{
                "action": "wait",
                "postcondition": {
                    "mode": "all", "settle_ms": 0, "timeout_s": 0.05,
                    "checks": [{
                        "type": "a11y", "selector": "text", "value": "任务列表（2项）",
                        "match": "contains", "state": "present",
                    }],
                },
            }],
            "final_assertion": True,
            "settle_ms": 0,
            "poll_ms": 1,
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["final_assertion_confirmed"] is True
    assert {os.path.basename(path) for path in result["evidence_paths"]} >= {
        "uisteps.jsonl", "final_assertions.jsonl",
    }
    assert any(path.endswith("_action_01_after.png") for path in result["evidence_paths"])
    assert all(item["type"] != "image" for item in tool_result["content"])
    assert tools.LAST_SCREENSHOT_META is None
    records = [json.loads(line) for line in (tmp_path / "final_assertions.jsonl").read_text().splitlines()]
    assert records[-1]["source"] == "execute_ui_actions"
    assert records[-1]["result"]["postcondition_satisfied"] is True


def test_weak_absent_postcondition_keeps_visual_fallback_for_final_assertion(
        tmp_path, monkeypatch):
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (1080, 2340, []))

    def save(path):
        Image.new("RGB", (108, 234), (20, 40, 60)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    tools.set_trace_dir(str(tmp_path))
    try:
        tool_result = asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{
                "action": "wait",
                "postcondition": {
                    "mode": "all", "settle_ms": 0, "timeout_s": 0.01,
                    "checks": [{
                        "type": "a11y", "selector": "text", "value": "旧页面",
                        "match": "contains", "state": "absent",
                    }],
                },
            }],
            "final_assertion": True,
            "poll_ms": 1,
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["ok"] is True
    assert result["postcondition_satisfied"] is True
    assert result["final_assertion_confirmed"] is False
    assert any(item["type"] == "image" for item in tool_result["content"])
    assert not (tmp_path / "final_assertions.jsonl").exists()


def test_screen_change_only_keeps_visual_fallback_for_final_assertion(
        tmp_path, monkeypatch):
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "execute_ui_steps", lambda *_args, **_kwargs: {
        "ok": True, "done_steps": 1, "total_steps": 1,
        "log": [{
            "i": 1, "action": "press_key", "action_executed": True,
            "postcondition_satisfied": True,
            "postcondition": {
                "satisfied": True,
                "checks": [{
                    "type": "screen_change", "state": "present", "satisfied": True,
                    "diff": 8.5, "threshold": 3.0,
                }],
            },
        }],
    })

    def save(path):
        Image.new("RGB", (108, 234), (20, 40, 60)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    tools.set_trace_dir(str(tmp_path))
    try:
        tool_result = asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{
                "action": "press_key", "target": "enter",
                "postcondition": {
                    "mode": "all", "settle_ms": 0, "timeout_s": 1,
                    "checks": [{"type": "screen_change", "diff_threshold": 3}],
                },
            }],
            "final_assertion": True,
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["ok"] is True
    assert result["postcondition_satisfied"] is True
    assert result["final_assertion_confirmed"] is False
    assert any(item["type"] == "image" for item in tool_result["content"])
    assert not (tmp_path / "final_assertions.jsonl").exists()


def test_final_assertion_prefers_delayed_semantic_match_over_early_screen_change(
        tmp_path, monkeypatch):
    reads = {"n": 0}
    taps = []
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "_wm_size", lambda: (1080, 2340))
    monkeypatch.setattr(server, "_adb", lambda *args: taps.append(args))

    def screen(**_kwargs):
        reads["n"] += 1
        elements = []
        if reads["n"] >= 5:
            elements.append({
                "text": "任务列表（2项）", "content_desc": "", "resource_id": "title",
                "center": [540, 120], "bounds": [0, 80, 1080, 160],
                "bounds_unreliable": False,
            })
        return 1080, 2340, elements

    monkeypatch.setattr(server, "_screen_elements", screen)
    monkeypatch.setattr(server, "_capture_postcondition_screen", lambda: {
        "path": str(tmp_path / "condition.png"), "bytes": 1, "secure": False,
    })
    monkeypatch.setattr(server, "_postcondition_screen_diff", lambda *_args: 9.0)

    def save(path):
        Image.new("RGB", (108, 234), (20, 40, 60)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    tools.set_trace_dir(str(tmp_path))
    try:
        tool_result = asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{
                "action": "tap_xy", "target": "xy:10,20",
                "postcondition": {
                    "mode": "any", "settle_ms": 0, "timeout_s": 0.08,
                    "checks": [
                        {"type": "a11y", "selector": "text",
                         "value": "任务列表（2项）", "match": "contains",
                         "state": "present"},
                        {"type": "screen_change", "diff_threshold": 3},
                    ],
                },
            }],
            "final_assertion": True,
            "poll_ms": 1,
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["ok"] is True
    assert result["postcondition_satisfied"] is True
    assert result["final_assertion_confirmed"] is True
    assert reads["n"] >= 5
    assert len(taps) == 1
    assert all(item["type"] != "image" for item in tool_result["content"])


def test_final_assertion_keeps_weak_visual_success_when_semantic_never_appears(
        tmp_path, monkeypatch):
    taps = []
    screen_captures = {"n": 0}
    monkeypatch.setattr(server, "_FINAL_SEMANTIC_GRACE_S", 0.01, raising=False)
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "_wm_size", lambda: (1080, 2340))
    monkeypatch.setattr(server, "_adb", lambda *args: taps.append(args))
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (1080, 2340, []))
    def capture_condition():
        screen_captures["n"] += 1
        return {
            "path": str(tmp_path / f"condition-{screen_captures['n']}.png"),
            "bytes": 1, "secure": False,
        }

    monkeypatch.setattr(server, "_capture_postcondition_screen", capture_condition)
    monkeypatch.setattr(server, "_postcondition_screen_diff", lambda *_args: 9.0)

    def save(path):
        Image.new("RGB", (108, 234), (20, 40, 60)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    tools.set_trace_dir(str(tmp_path))
    try:
        tool_result = asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{
                "action": "tap_xy", "target": "xy:10,20",
                "postcondition": {
                    "mode": "any", "settle_ms": 0, "timeout_s": 0.08,
                    "checks": [
                        {"type": "a11y", "selector": "text", "value": "目标页",
                         "match": "contains", "state": "present"},
                        {"type": "screen_change", "diff_threshold": 3},
                    ],
                },
            }],
            "final_assertion": True,
            "poll_ms": 1,
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["ok"] is True
    assert result["postcondition_satisfied"] is True
    assert result["final_assertion_confirmed"] is False
    assert 8 <= result["final_postcondition"]["waited_ms"] < 50
    assert screen_captures["n"] == 2
    assert len(taps) == 1
    assert any(item["type"] == "image" for item in tool_result["content"])
    assert not (tmp_path / "final_assertions.jsonl").exists()


def test_final_semantic_grace_covers_a_late_accessibility_observation(monkeypatch):
    clock = {"now": 0.0}

    monkeypatch.setattr(server.time, "monotonic", lambda: clock["now"])
    monkeypatch.setattr(server.time, "sleep", lambda seconds: clock.__setitem__(
        "now", clock["now"] + seconds,
    ))
    monkeypatch.setattr(server, "get_focused_app", lambda: "pkg/Activity")
    monkeypatch.setattr(server, "_discard_condition_capture", lambda _capture: None)
    monkeypatch.setattr(server, "_capture_postcondition_screen", lambda: {
        "path": "after.png", "bytes": 1, "secure": False,
    })
    monkeypatch.setattr(server, "_postcondition_screen_diff", lambda *_args: 9.0)

    def elements(**_kwargs):
        text = "目标页" if clock["now"] >= 4.0 else ""
        return 1080, 2340, [{
            "text": text, "resource_id": "", "content_desc": "",
            "center": [540, 1000], "bounds": [400, 900, 680, 1100],
        }]

    monkeypatch.setattr(server, "_screen_elements", elements)

    result = server._evaluate_postcondition(
        {
            "mode": "any", "settle_ms": 0, "timeout_s": 10,
            "checks": [
                {"type": "a11y", "selector": "text", "value": "目标页",
                 "match": "contains", "state": "present"},
                {"type": "screen_change", "diff_threshold": 3},
            ],
        },
        poll_s=1.0,
        screen_before={"path": "before.png", "bytes": 1, "secure": False},
        prefer_positive_semantic=True,
    )

    assert result["satisfied"] is True
    assert result["waited_ms"] == 4000
    assert result["checks"][0]["matched_count"] == 1


def test_positive_focus_requires_an_actual_transition_for_final_assertion(
        tmp_path, monkeypatch):
    focus = {"value": "pkg/OldActivity"}
    monkeypatch.setattr(server, "get_focused_app", lambda: focus["value"])

    def press_key(_key):
        focus["value"] = "pkg/NewActivity"
        return {"ok": True}

    monkeypatch.setattr(server, "press_key", press_key)

    def save(path):
        Image.new("RGB", (108, 234), (20, 40, 60)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    tools.set_trace_dir(str(tmp_path))
    try:
        result = json.loads(asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{
                "action": "press_key", "target": "enter",
                "postcondition": {
                    "mode": "all", "settle_ms": 0, "timeout_s": 0.01,
                    "checks": [{
                        "type": "focus", "value": "NewActivity", "state": "present",
                    }],
                },
            }],
            "final_assertion": True,
            "poll_ms": 1,
        }))["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["final_assertion_confirmed"] is True
    assert (tmp_path / "final_assertions.jsonl").is_file()
    assert "_final_focus_before" not in (tmp_path / "uisteps.jsonl").read_text(
        encoding="utf-8"
    )


def test_same_window_focus_match_is_not_a_deterministic_final_assertion(
        tmp_path, monkeypatch):
    monkeypatch.setattr(server, "get_focused_app", lambda: "pkg/SameActivity")
    monkeypatch.setattr(server, "press_key", lambda _key: {"ok": True})

    def save(path):
        Image.new("RGB", (108, 234), (20, 40, 60)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    tools.set_trace_dir(str(tmp_path))
    try:
        tool_result = asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{
                "action": "press_key", "target": "enter",
                "postcondition": {
                    "mode": "all", "settle_ms": 0, "timeout_s": 0.01,
                    "checks": [{
                        "type": "focus", "value": "SameActivity", "state": "present",
                    }],
                },
            }],
            "final_assertion": True,
            "poll_ms": 1,
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["ok"] is True
    assert result["postcondition_satisfied"] is True
    assert result["final_assertion_confirmed"] is False
    assert any(item["type"] == "image" for item in tool_result["content"])
    assert not (tmp_path / "final_assertions.jsonl").exists()


def test_obsolete_single_action_tools_are_not_publicly_exposed():
    removed = {
        "tap", "tap_screenshot", "tap_element", "wait_for", "input_text",
        "clear_input", "press_key", "launch_app", "stop_app",
    }
    native_names = {getattr(item, "name", "") for item in tools._ALL}
    fastmcp_names = set(server.mcp._tool_manager._tools)

    assert removed.isdisjoint(native_names)
    assert removed.isdisjoint(fastmcp_names)
    assert all(not hasattr(tools, name) for name in removed)


def test_postcondition_any_and_absent_are_evaluated_from_one_screen_read(monkeypatch):
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "完成", "content_desc": "", "resource_id": "page_done",
            "center": [540, 120], "bounds": [0, 80, 1080, 160],
            "bounds_unreliable": False,
        }],
    ))

    result = server.execute_ui_steps([{
        "action": "wait",
        "postcondition": {
            "mode": "all", "settle_ms": 0, "timeout_s": 0.05,
            "checks": [
                {"type": "a11y", "selector": "text", "value": "错误",
                 "match": "exact", "state": "absent"},
                {"type": "a11y", "selector": "resource_id", "value": "missing",
                 "match": "suffix", "state": "absent"},
            ],
        },
    }, {
        "action": "wait",
        "postcondition": {
            "mode": "any", "settle_ms": 0, "timeout_s": 0.05,
            "checks": [
                {"type": "a11y", "selector": "text", "value": "不存在",
                 "match": "exact", "state": "present"},
                {"type": "a11y", "selector": "resource_id", "value": "done",
                 "match": "suffix", "state": "present"},
            ],
        },
    }], poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    assert [step["postcondition_satisfied"] for step in result["log"]] == [True, True]


def test_postcondition_uses_the_action_default_timeout_when_omitted(monkeypatch):
    reads = {"n": 0}
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")

    def delayed_screen(**_kwargs):
        reads["n"] += 1
        elements = [] if reads["n"] == 1 else [{
            "text": "完成", "content_desc": "", "resource_id": "done",
            "center": [540, 120], "bounds": [0, 80, 1080, 160],
            "bounds_unreliable": False,
        }]
        return 1080, 2340, elements

    monkeypatch.setattr(server, "_screen_elements", delayed_screen)

    result = server.execute_ui_steps([{
        "action": "wait",
        "postcondition": {
            "mode": "all", "settle_ms": 0,
            "checks": [{"type": "a11y", "selector": "text", "value": "完成",
                        "match": "exact", "state": "present"}],
        },
    }], default_timeout_s=0.05, poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    assert reads["n"] >= 2


def test_screen_change_accepts_an_explicit_zero_threshold(tmp_path, monkeypatch):
    before = tmp_path / "before.png"
    after = tmp_path / "after.png"
    before.write_bytes(b"before")
    after.write_bytes(b"after")
    monkeypatch.setattr(server, "_capture_postcondition_screen", lambda: {
        "path": str(after), "bytes": 1, "secure": False,
    })
    monkeypatch.setattr(server, "_postcondition_screen_diff", lambda *_args: 1.0)

    result = server._evaluate_postcondition({
        "mode": "all", "settle_ms": 0, "timeout_s": 0,
        "checks": [{"type": "screen_change", "diff_threshold": 0}],
    }, poll_s=0.001, screen_before={
        "path": str(before), "bytes": 1, "secure": False,
    })

    assert result["satisfied"] is True
    assert result["checks"][0]["threshold"] == 0


def test_side_effect_postcondition_honors_explicit_webview_timeout(monkeypatch):
    monkeypatch.setattr(server, "_WEBVIEW_POSTCONDITION_TIMEOUT_S", 0.01)
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Web")
    reads = 0

    def screen(**_kwargs):
        nonlocal reads
        reads += 1
        elements = [{
            "text": "", "content_desc": "https://example.test/page.html",
            "resource_id": "", "class": "WebView",
        }]
        if reads >= 15:
            elements.append({
                "text": "目标", "content_desc": "", "resource_id": "",
                "center": [540, 120], "bounds_unreliable": False,
            })
        return 1080, 2340, elements

    monkeypatch.setattr(server, "_screen_elements", screen)
    monkeypatch.setattr(server, "_adb", lambda *_args: None)

    result = server.execute_ui_steps([{
        "action": "tap_xy", "target": "xy:10,20",
        "postcondition": {
            "mode": "all", "settle_ms": 0, "timeout_s": 0.08,
            "checks": [{"type": "a11y", "selector": "text", "value": "目标",
                        "match": "contains", "state": "present"}],
        },
    }], poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    assert result["log"][0]["action_executed"] is True
    assert result["log"][0]["postcondition_satisfied"] is True
    assert reads >= 15


def test_explicit_wait_postcondition_honors_its_webview_timeout(monkeypatch):
    monkeypatch.setattr(server, "_WEBVIEW_POSTCONDITION_TIMEOUT_S", 0.01)
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Web")
    reads = 0

    def screen(**_kwargs):
        nonlocal reads
        reads += 1
        elements = [{
            "text": "", "content_desc": "https://example.test/page.html",
            "resource_id": "", "class": "WebView",
        }]
        if reads >= 15:
            elements.append({"text": "目标", "content_desc": "", "resource_id": ""})
        return 1080, 2340, elements

    monkeypatch.setattr(server, "_screen_elements", screen)

    result = server.execute_ui_steps([{
        "action": "wait",
        "postcondition": {
            "mode": "all", "settle_ms": 0, "timeout_s": 0.08,
            "checks": [{"type": "a11y", "selector": "text", "value": "目标",
                        "match": "contains", "state": "present"}],
        },
    }], poll_ms=1, settle_ms=0)

    assert result["ok"] is True
    assert reads >= 15


def test_implicit_wait_postcondition_keeps_the_webview_quick_probe_cap(monkeypatch):
    monkeypatch.setattr(server, "_WEBVIEW_POSTCONDITION_TIMEOUT_S", 0.01)
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Web")
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "", "content_desc": "https://example.test/page.html",
            "resource_id": "", "class": "WebView",
        }],
    ))

    result = server.execute_ui_steps([{
        "action": "wait",
        "postcondition": {
            "mode": "all", "settle_ms": 0,
            "checks": [{"type": "a11y", "selector": "text", "value": "目标",
                        "match": "contains", "state": "present"}],
        },
    }], default_timeout_s=0.08, poll_ms=1, settle_ms=0)

    assert result["ok"] is False
    observed = result["failed_at"]["postcondition"]
    assert observed["webview_surface"] is True
    assert observed["effective_timeout_s"] == 0.01
    assert observed["waited_ms"] < 50


def test_screen_change_on_secure_screen_is_explicitly_unavailable(monkeypatch):
    monkeypatch.setattr(server, "save_screenshot", lambda path: {
        "path": path, "bytes": 0, "secure": True,
    })
    monkeypatch.setattr(server, "_adb", lambda *_args: None)
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Secure")

    result = server.execute_ui_steps([{
        "action": "tap_xy", "target": "xy:10,20",
        "postcondition": {
            "mode": "all", "settle_ms": 0, "timeout_s": 0.01,
            "checks": [{"type": "screen_change", "diff_threshold": 3}],
        },
    }], poll_ms=1, settle_ms=0)

    assert result["ok"] is False
    assert result["failed_at"]["action_executed"] is True
    check = result["failed_at"]["postcondition"]["checks"][0]
    assert check["satisfied"] is False
    assert check["unavailable"] == "secure_screen"


def test_failed_postcondition_reports_executed_action_without_replay(monkeypatch):
    taps = []
    monkeypatch.setattr(server, "_adb", lambda *args: taps.append(args))
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (1080, 2340, []))

    result = server.execute_ui_steps([{
        "action": "tap_xy", "target": "xy:10,20",
        "postcondition": {
            "mode": "all", "settle_ms": 0, "timeout_s": 0.01,
            "checks": [{"type": "a11y", "selector": "text", "value": "目标页",
                        "match": "contains", "state": "present"}],
        },
    }], poll_ms=1, settle_ms=0)

    assert result["ok"] is False
    assert result["done_steps"] == 0
    assert result["failed_at"]["action_executed"] is True
    assert result["failed_at"]["postcondition_satisfied"] is False
    assert len(taps) == 1


def test_settle_only_final_request_succeeds_with_visual_fallback(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")

    def save(path):
        Image.new("RGB", (108, 234), (20, 40, 60)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    tools.set_trace_dir(str(tmp_path))
    try:
        tool_result = asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{
                "action": "wait",
                "postcondition": {
                    "mode": "all", "settle_ms": 1, "timeout_s": 1, "checks": [],
                },
            }],
            "final_assertion": True,
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["ok"] is True
    assert result["postcondition_satisfied"] is True
    assert result["final_assertion_confirmed"] is False
    assert any(item["type"] == "image" for item in tool_result["content"])
    assert not (tmp_path / "final_assertions.jsonl").exists()


def test_execute_ui_actions_allows_a_visually_confirmed_app_key_above_absolute_bottom(
        monkeypatch):
    captured = []
    tools.LAST_SCREENSHOT_META = {
        "image_wh": [108, 234], "device_wh": [1080, 2340],
        "scale": 10.0, "path": "/tmp/current-frame.png",
    }
    monkeypatch.setattr(
        server, "execute_ui_steps",
        lambda actions, *_args: captured.extend(actions) or {
            "ok": True, "done_steps": len(actions), "total_steps": len(actions), "log": [],
        },
    )
    try:
        result = json.loads(asyncio.run(tools.execute_ui_actions.handler({
            "actions": [
                {"action": "tap_image", "target": "13,225", "allow_bottom_edge": True},
                {"action": "tap_image", "target": "54,180"},
            ],
            "settle_ms": 0,
        }))["content"][0]["text"])
    finally:
        tools.LAST_SCREENSHOT_META = None

    assert result["ok"] is True
    assert result["bottom_edge_override_count"] == 1
    assert captured[0] == {"action": "tap_xy", "target": "xy:130,2250"}
    assert "allow_bottom_edge" not in captured[0]


def test_execute_ui_actions_converts_current_image_coordinates_inside_the_batch(monkeypatch):
    captured = []
    tools.LAST_SCREENSHOT_META = {
        "image_wh": [108, 234],
        "device_wh": [1080, 2340],
        "scale": 10.0,
        "path": "/tmp/current-frame.png",
    }
    monkeypatch.setattr(
        server,
        "execute_ui_steps",
        lambda actions, *_args: captured.extend(actions) or {
            "ok": True, "done_steps": len(actions), "total_steps": len(actions), "log": [],
        },
    )
    try:
        result = json.loads(asyncio.run(tools.execute_ui_actions.handler({
            "actions": [
                {"action": "tap_image", "target": "20,100"},
                {"action": "tap_image", "target": "80,200"},
            ],
            "settle_ms": 0,
        }))["content"][0]["text"])
    finally:
        tools.LAST_SCREENSHOT_META = None

    assert result["ok"] is True
    assert result["image_tap_count"] == 2
    assert captured == [
        {"action": "tap_xy", "target": "xy:200,1000"},
        {"action": "tap_xy", "target": "xy:800,2000"},
    ]


def test_execute_ui_actions_records_tap_image_as_the_declared_action(monkeypatch):
    tools.set_trace_dir(None)
    tools.LAST_SCREENSHOT_META = {
        "image_wh": [108, 234], "device_wh": [1080, 2340],
        "scale": 10.0, "path": "/tmp/current-frame.png",
    }
    monkeypatch.setattr(server, "execute_ui_steps", lambda *_args: {
        "ok": True, "done_steps": 1, "total_steps": 1, "log": [{
            "i": 1, "action": "tap_xy", "target": "xy:200,1000", "ok": True,
            "action_executed": True, "postcondition_satisfied": True,
            "action_result": {"action": "tap", "x": 200, "y": 1000},
        }],
    })
    async def no_frame(_tag):
        return "", {"bytes": 0, "secure": True}
    monkeypatch.setattr(tools, "_capture_current_frame", no_frame)
    try:
        result = json.loads(asyncio.run(tools.execute_ui_actions.handler({
            "actions": [{"action": "tap_image", "target": "20,100"}],
            "settle_ms": 0,
        }))["content"][0]["text"])
    finally:
        tools.LAST_SCREENSHOT_META = None

    assert result["steps"][0]["action"] == "tap_image"
    assert result["steps"][0]["target"] == "20,100"
    assert result["steps"][0]["action_result"] == {
        "action": "tap", "x": 200, "y": 1000,
    }


def test_execute_ui_actions_preflights_every_action_before_touching_the_device(monkeypatch):
    pressed = []
    tools.CONTROL = None
    tools.set_trace_dir(None)
    monkeypatch.setattr(
        server, "press_key", lambda key: pressed.append(key) or {"ok": True},
    )

    result = json.loads(asyncio.run(tools.execute_ui_actions.handler({
        "actions": [
            {"action": "press_key", "target": "back"},
            {"action": "tap_text_retry", "target": "设置"},
        ],
        "settle_ms": 0,
    }))["content"][0]["text"])

    assert result["ok"] is False
    assert result["preflight_rejected"] is True
    assert result["failed_at"]["i"] == 2
    assert "postcondition" in result["failed_at"]["reason"]
    assert pressed == []


def test_execute_ui_actions_rejects_empty_wait_before_touching_the_device(monkeypatch):
    pressed = []
    tools.CONTROL = None
    tools.set_trace_dir(None)
    monkeypatch.setattr(
        server, "press_key", lambda key: pressed.append(key) or {"ok": True},
    )

    result = json.loads(asyncio.run(tools.execute_ui_actions.handler({
        "actions": [
            {"action": "press_key", "target": "back"},
            {"action": "wait", "timeout_s": 3},
        ],
        "settle_ms": 0,
    }))["content"][0]["text"])

    assert result["ok"] is False
    assert result["preflight_rejected"] is True
    assert result["failed_at"]["i"] == 2
    assert "postcondition" in result["failed_at"]["reason"]
    assert pressed == []


def test_find_webview_miss_returns_visual_fallback_in_the_same_call(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "", "content_desc": "https://example.test/page.html",
            "resource_id": "", "class": "android.webkit.WebView", "clickable": False,
        }],
    ))
    screenshots = []

    def save(path):
        screenshots.append(path)
        Image.new("RGB", (108, 234), (30, 60, 90)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    tools.set_trace_dir(str(tmp_path / "case-webview"))
    try:
        tool_result = asyncio.run(tools.find.handler({"text": "批量定投"}))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["count"] == 0
    assert result["webview_surface"] is True
    assert result["observation_source"] == "find_webview_fallback"
    assert len(screenshots) == 1
    assert any(item["type"] == "image" for item in tool_result["content"])


def test_find_explicit_miss_returns_visual_fallback_without_webview(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "普通原生页面", "content_desc": "", "resource_id": "title",
            "class": "android.widget.TextView", "clickable": False,
        }],
    ))
    screenshots = []

    def save(path):
        screenshots.append(path)
        Image.new("RGB", (108, 234), (30, 60, 90)).save(path)
        return {"path": path, "bytes": os.path.getsize(path), "secure": False}

    monkeypatch.setattr(server, "save_screenshot", save)
    tools.set_trace_dir(str(tmp_path / "case-native"))
    try:
        tool_result = asyncio.run(tools.find.handler({
            "text": "批量定投", "visual_on_miss": True,
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["count"] == 0
    assert result.get("webview_surface") is not True
    assert result["visual_fallback_requested"] is True
    assert result["visual_fallback_used"] is True
    assert result["observation_source"] == "find_visual_fallback"
    assert len(screenshots) == 1
    assert any(item["type"] == "image" for item in tool_result["content"])


def test_confirmed_batch_postcondition_blocks_a_redundant_screenshot(monkeypatch, tmp_path):
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "测试完成", "content_desc": "", "resource_id": "",
            "center": [540, 1200], "bounds_unreliable": False,
        }],
    ))
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    screenshots = []
    monkeypatch.setattr(server, "save_screenshot", lambda path: screenshots.append(path) or {
        "path": path, "bytes": 0, "secure": True,
    })
    tools.set_trace_dir(str(tmp_path / "case-1"))

    confirmed = json.loads(asyncio.run(tools.execute_ui_actions.handler({
        "actions": [{
            "action": "wait",
            "postcondition": {"mode": "all", "settle_ms": 0, "timeout_s": 0.01,
                              "checks": [{"type": "a11y", "selector": "text",
                                          "value": "测试完成", "match": "contains",
                                          "state": "present"}]},
        }],
        "final_assertion": True,
        "settle_ms": 0,
        "poll_ms": 1,
    }))["content"][0]["text"])
    blocked = json.loads(asyncio.run(tools.screenshot.handler({
        "path": str(tmp_path / "redundant.png"),
    }))["content"][0]["text"])

    assert confirmed["final_assertion_confirmed"] is True
    assertion_path = next(path for path in confirmed["evidence_paths"]
                          if path.endswith("final_assertions.jsonl"))
    assert os.path.isfile(assertion_path)
    assert blocked == {
        "ok": False,
        "redundant_screenshot_blocked": True,
        "reason": "最终观测点已经由 execute_ui_actions 的结构化命中确认；请直接记录步骤结果",
    }
    assert len(screenshots) == 3


def test_confirmed_final_find_blocks_a_redundant_screenshot(monkeypatch, tmp_path):
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "结果通过", "content_desc": "", "resource_id": "result",
            "center": [540, 1200], "bounds": [0, 1100, 1080, 1300],
            "clickable": False, "bounds_unreliable": False,
        }],
    ))
    screenshots = []
    monkeypatch.setattr(server, "save_screenshot", lambda path: screenshots.append(path) or {
        "path": path, "bytes": 0, "secure": True,
    })
    tools.set_trace_dir(str(tmp_path / "case-find"))

    confirmed = json.loads(asyncio.run(tools.find.handler({
        "text": "结果通过", "final_assertion": True,
    }))["content"][0]["text"])
    blocked = json.loads(asyncio.run(tools.screenshot.handler({
        "path": str(tmp_path / "redundant-find.png"),
    }))["content"][0]["text"])

    assert confirmed["final_assertion_confirmed"] is True
    assert blocked["redundant_screenshot_blocked"] is True
    assert "find" in blocked["reason"]
    assert screenshots == []


def test_final_assertion_without_structured_evidence_keeps_screenshot_fallback(monkeypatch, tmp_path):
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "完成", "content_desc": "", "resource_id": "",
            "center": [540, 1200], "bounds_unreliable": False,
        }],
    ))
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    screenshots = []
    monkeypatch.setattr(server, "save_screenshot", lambda path: screenshots.append(path) or {
        "path": path, "bytes": 0, "secure": True,
    })
    monkeypatch.setattr(tools, "_log_final_assertion", lambda *_args: "")
    trace_dir = tmp_path / "case-1"
    tools.set_trace_dir(str(trace_dir))

    asyncio.run(tools.execute_ui_actions.handler({
        "actions": [{
            "action": "wait",
            "postcondition": {"mode": "all", "settle_ms": 0, "timeout_s": 0.01,
                              "checks": [{"type": "a11y", "selector": "text",
                                          "value": "完成", "match": "contains",
                                          "state": "present"}]},
        }],
        "final_assertion": True,
        "settle_ms": 0,
        "poll_ms": 1,
    }))
    result = json.loads(asyncio.run(tools.screenshot.handler({}))["content"][0]["text"])
    tools.set_trace_dir(None)

    assert result["secure"] is True
    assert len(screenshots) == 4
    assert os.path.dirname(screenshots[-1]) == str(trace_dir)


def test_final_assertion_screenshot_guard_ends_at_the_next_step(monkeypatch, tmp_path):
    monkeypatch.setattr(server, "_screen_elements", lambda **_kwargs: (
        1080, 2340, [{
            "text": "当前步骤完成", "content_desc": "", "resource_id": "",
            "center": [540, 1200], "bounds": [0, 1100, 1080, 1300],
            "clickable": False, "bounds_unreliable": False,
        }],
    ))
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    screenshots = []
    monkeypatch.setattr(server, "save_screenshot", lambda path: screenshots.append(path) or {
        "path": path, "bytes": 0, "secure": True,
    })
    tools.set_trace_dir(str(tmp_path / "case-1"))

    asyncio.run(tools.find.handler({
        "text": "当前步骤完成", "final_assertion": True,
    }))
    tools.set_trace_dir(str(tmp_path / "case-2"))
    result = json.loads(asyncio.run(tools.screenshot.handler({}))["content"][0]["text"])
    tools.set_trace_dir(None)

    assert result["secure"] is True
    assert len(screenshots) == 1
    assert os.path.dirname(screenshots[0]) == str(tmp_path / "case-2")


def test_execute_ui_actions_runs_a_temporary_batch_and_writes_evidence(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(
        server, "save_screenshot",
        lambda path: (
            Image.new("RGB", (108, 234), (30, 60, 90)).save(path)
            or {"path": path, "bytes": os.path.getsize(path), "secure": False}
        ),
    )
    async def no_sleep(_seconds):
        return None
    monkeypatch.setattr(tools.asyncio, "sleep", no_sleep)
    monkeypatch.setattr(server, "press_key", lambda key: {"ok": True, "key": key})
    tools.set_trace_dir(str(tmp_path))
    try:
        tool_result = asyncio.run(tools.execute_ui_actions.handler({
            "actions": [
                {"action": "press_key", "target": "back"},
                {"action": "press_key", "target": "back"},
            ],
            "settle_ms": 0,
        }))
        result = json.loads(tool_result["content"][0]["text"])
    finally:
        tools.set_trace_dir(None)

    assert result["ok"] is True
    assert result["done_steps"] == 2
    assert any(path.endswith("uisteps.jsonl") for path in result["evidence_paths"])
    action_after_paths = [
        path for path in result["evidence_paths"]
        if re.search(r"_action_0[12]_after\.png$", path)
    ]
    assert len(action_after_paths) == 2
    step_record = json.loads((tmp_path / "uisteps.jsonl").read_text().splitlines()[-1])
    assert [step["after_image"]["path"] for step in step_record["steps"]] == action_after_paths
    assert result["path"] == action_after_paths[-1]
    assert result["observation_source"] == "batch_after"
    assert any(item["type"] == "image" for item in tool_result["content"])


def test_execute_ui_actions_stops_before_starting_the_next_action(monkeypatch):
    class Control:
        stopped = False

    control = Control()
    calls = 0

    def evaluate_postcondition(*_args, **_kwargs):
        nonlocal calls
        calls += 1
        control.stopped = True
        return {
            "satisfied": True, "mode": "all", "checks": [{"satisfied": True}],
            "waited_ms": 0, "focused_app": "com.example/Main",
        }

    monkeypatch.setattr(server, "_evaluate_postcondition", evaluate_postcondition)
    monkeypatch.setattr(server, "get_focused_app", lambda: "com.example/Main")
    monkeypatch.setattr(tools, "CONTROL", control)
    async def no_frame(_tag):
        return "", {"bytes": 0, "secure": True}
    monkeypatch.setattr(tools, "_capture_current_frame", no_frame)

    result = json.loads(asyncio.run(tools.execute_ui_actions.handler({
        "actions": [
            {"action": "wait", "postcondition": {
                "mode": "all", "settle_ms": 0, "timeout_s": 1,
                "checks": [{"type": "focus", "value": "ready"}],
            }},
            {"action": "wait", "postcondition": {
                "mode": "all", "settle_ms": 0, "timeout_s": 1,
                "checks": [{"type": "focus", "value": "never-started"}],
            }},
        ],
        "settle_ms": 0,
    }))["content"][0]["text"])

    assert result["stopped"] is True
    assert result["done_steps"] == 1
    assert calls == 1


def _stub(monkeypatch, screen, acts):
    """screen = {"focus":…, "texts":[…]}；acts 收集实际做过的动作。"""
    monkeypatch.setattr(server, "get_focused_app", lambda: screen["focus"])
    monkeypatch.setattr(
        server, "_screen_elements",
        lambda filter_useful=True, max_elements=300: (
            1080, 2340,
            [{"text": t, "content_desc": "", "resource_id": ""} for t in screen["texts"]]
            + [{"text": "", "content_desc": "", "resource_id": r} for r in screen.get("rids", [])]))
    monkeypatch.setattr(server, "_adb", lambda *a: acts.append(("adb",) + a))

    def _te(resource_id="", text="", index=0, contains=False, content_desc=""):
        acts.append(("tap_element", resource_id or text or content_desc))
        return {"ok": text in screen["texts"] or bool(resource_id or content_desc), "matched": 1}
    monkeypatch.setattr(server, "tap_element", _te)
    monkeypatch.setattr(server, "input_text", lambda t: (acts.append(("input", t)), {"ok": True})[1])
    monkeypatch.setattr(server, "clear_input",
                        lambda resource_id="", text="", max_chars=40:
                        (acts.append(("clear", resource_id or text)), {"ok": True, "after_len": 0})[1])
    monkeypatch.setattr(server, "swipe", lambda *a: acts.append(("swipe",) + a))
    monkeypatch.setattr(server, "press_key", lambda k: (acts.append(("key", k)), {"ok": True})[1])


# ---------- 顺利路径 ----------

def test_runs_all_steps_in_order(monkeypatch):
    screen = {"focus": "pkg/A", "texts": ["设置", "退出登录"]}
    acts = []
    _stub(monkeypatch, screen, acts)
    out = server.run_ui_steps([
        {"action": "tap_text", "target": "设置", "expect": "退出登录"},
        {"action": "clear_input", "target": "rid:acct"},
        {"action": "input", "target": "someone"},
        {"action": "press_key", "target": "back"},
    ], **FAST)
    assert out["ok"] is True
    assert out["done_steps"] == 4 and out["total_steps"] == 4
    assert [a[0] for a in acts] == ["tap_element", "clear", "input", "key"]
    assert acts[2] == ("input", "someone")


def test_tool_returns_persisted_ui_steps_as_evidence(monkeypatch, tmp_path):
    monkeypatch.setattr(tools, "TRACE_DIR", str(tmp_path))
    monkeypatch.setattr(tools, "UI_SCRIPTS", {
        "核对": [{"action": "wait", "expect": "完成"}],
    })
    monkeypatch.setattr(server, "execute_ui_steps", lambda *args, **kwargs: {
        "ok": True, "done_steps": 1, "total_steps": 1,
        "log": [{"i": 1, "action": "wait", "ok": True}],
    })

    result = asyncio.run(tools.run_ui_steps.handler({
        "script": "核对", "vars": {}, "default_timeout_s": 1,
        "poll_ms": 1, "settle_ms": 0,
    }))
    payload = json.loads(result["content"][0]["text"])

    assert payload["evidence_path"] == str(tmp_path / "uisteps.jsonl")
    assert (tmp_path / "uisteps.jsonl").is_file()


def test_expect_focus_form(monkeypatch):
    screen = {"focus": "pkg/LoginUI", "texts": []}
    acts = []
    _stub(monkeypatch, screen, acts)
    out = server.run_ui_steps([{"action": "wait", "expect": "focus:LoginUI"}], **FAST)
    assert out["ok"] is True


def test_target_prefixes(monkeypatch):
    screen = {"focus": "pkg/A", "texts": ["确定"], "rids": ["pkg:id/ok_btn"]}
    acts = []
    _stub(monkeypatch, screen, acts)
    out = server.run_ui_steps([
        {"action": "tap_text", "target": "xy:540,1746"},
        {"action": "tap_text", "target": "rid:ok_btn"},
        {"action": "tap_xy", "target": "100,200"},
    ], **FAST)
    assert out["ok"] is True
    assert ("adb", "shell", "input", "tap", "540", "1746") in acts
    assert ("adb", "shell", "input", "tap", "100", "200") in acts
    assert ("tap_element", "ok_btn") in acts


def test_desc_prefix_uses_exact_content_description(monkeypatch):
    """桌面图标等控件可能只有 content-desc；步骤表应能精确声明，不退回坐标。"""
    screen = {"focus": "pkg/A", "texts": []}
    acts = []
    _stub(monkeypatch, screen, acts)
    monkeypatch.setattr(server, "_screen_elements", lambda **k: (
        1080, 2340,
        [{"text": "", "content_desc": "目标入口", "resource_id": ""}],
    ))

    out = server.run_ui_steps(
        [{"action": "tap_text", "target": "desc:目标入口"}],
        default_timeout_s=0, poll_ms=1, settle_ms=0,
    )

    assert out["ok"] is True
    assert ("tap_element", "目标入口") in acts


def test_row_prefix_uses_text_y_and_taps_viewport_midline(monkeypatch):
    """文字节点不可点击时，row: 仍以 a11y 文本定位 y，但点击整行中部，不写死设备坐标。"""
    screen = {"focus": "pkg/A", "texts": ["设置", "通用"]}
    acts = []
    _stub(monkeypatch, screen, acts)
    monkeypatch.setattr(server, "_screen_elements", lambda **kwargs: (
        1080, 2340,
        [
            {"text": "设置", "content_desc": "", "resource_id": "android:id/title",
             "center": [197, 1505]},
            {"text": "通用", "content_desc": "", "resource_id": "", "center": [89, 1253]},
        ],
    ))

    out = server.run_ui_steps(
        [{"action": "tap_text", "target": "row:设置", "expect": "通用"}],
        default_timeout_s=0, poll_ms=1, settle_ms=0,
    )

    assert out["ok"] is True
    assert ("adb", "shell", "input", "tap", "540", "1505") in acts
    assert not [action for action in acts if action[0] == "tap_element"]


def test_row_prefix_missing_target_never_taps(monkeypatch):
    screen = {"focus": "pkg/A", "texts": []}
    acts = []
    _stub(monkeypatch, screen, acts)

    out = server.run_ui_steps(
        [{"action": "tap_text", "target": "row:设置"}],
        default_timeout_s=0, poll_ms=1, settle_ms=0,
    )

    assert out["ok"] is False
    assert not [action for action in acts if action[0] == "adb"]


def test_tap_text_retry_clicks_twice_only_while_target_remains(monkeypatch):
    """首击被吞时可再点一次；第二击必须由“目标仍在”这个页面证据约束。"""
    screen = {"focus": "pkg/A", "texts": ["设置"]}
    acts = []
    _stub(monkeypatch, screen, acts)

    def adb(*args):
        acts.append(("adb",) + args)

    def tap_element(resource_id="", text="", index=0, contains=False, content_desc=""):
        acts.append(("tap_element", resource_id or text or content_desc))
        screen["texts"] = ["通用"]
        return {"ok": True, "matched": 1}

    monkeypatch.setattr(server, "_adb", adb)
    monkeypatch.setattr(server, "tap_element", tap_element)
    monkeypatch.setattr(server, "_screen_elements", lambda **kwargs: (
        1080, 2340,
        [{"text": text, "content_desc": "", "resource_id": "android:id/title",
          "center": [197, 1505]} for text in screen["texts"]],
    ))

    out = server.run_ui_steps(
        [{"action": "tap_text_retry", "target": "row:设置", "expect": "通用",
          "timeout_s": 0.02}],
        default_timeout_s=0.02, poll_ms=1, settle_ms=0,
    )

    assert out["ok"] is True
    assert [action for action in acts if action[0] == "adb"] == [
        ("adb", "shell", "input", "tap", "540", "1505"),
    ]
    assert [action for action in acts if action[0] == "tap_element"] == [
        ("tap_element", "设置"),
    ]
    assert out["log"][0]["attempts"] == 2


def test_tap_text_retry_does_not_click_again_after_target_disappears(monkeypatch):
    """首击后目标消失但期望也未出现，说明页面已变化；此时禁止盲目第二击。"""
    screen = {"focus": "pkg/A", "texts": ["设置"]}
    acts = []
    _stub(monkeypatch, screen, acts)

    def adb(*args):
        acts.append(("adb",) + args)
        screen["texts"] = ["另一个页面"]

    monkeypatch.setattr(server, "_adb", adb)
    monkeypatch.setattr(server, "_screen_elements", lambda **kwargs: (
        1080, 2340,
        [{"text": text, "content_desc": "", "resource_id": "android:id/title",
          "center": [197, 1505]} for text in screen["texts"]],
    ))

    out = server.run_ui_steps(
        [{"action": "tap_text_retry", "target": "row:设置", "expect": "通用",
          "timeout_s": 0.01}],
        default_timeout_s=0.01, poll_ms=1, settle_ms=0,
    )

    assert out["ok"] is False
    assert len([action for action in acts if action[0] == "adb"]) == 1
    assert "目标已不在当前页" in out["failed_at"]["reason"]


# ---------- 失败即停 ----------

def test_stops_at_failed_expect_and_does_not_continue(monkeypatch):
    screen = {"focus": "pkg/A", "texts": ["设置"]}          # "退出登录" 永远不出现
    acts = []
    _stub(monkeypatch, screen, acts)
    out = server.run_ui_steps([
        {"action": "tap_text", "target": "设置", "expect": "退出登录"},
        {"action": "input", "target": "不该被执行"},
    ], **FAST)
    assert out["ok"] is False
    assert out["done_steps"] == 0 and out["total_steps"] == 2
    assert out["failed_at"]["i"] == 1 and "校验超时" in out["failed_at"]["reason"]
    assert ("input", "不该被执行") not in acts        # 关键：后续步骤一步都没做
    assert "中间态" in out["hint"]


def test_stops_when_target_not_found(monkeypatch):
    screen = {"focus": "pkg/A", "texts": []}
    acts = []
    _stub(monkeypatch, screen, acts)
    out = server.run_ui_steps([
        {"action": "tap_text", "target": "并不存在的按钮"},
        {"action": "input", "target": "不该被执行"},
    ], **FAST)
    assert out["ok"] is False and out["failed_at"]["i"] == 1
    assert "等待目标" in out["failed_at"]["reason"]
    assert ("input", "不该被执行") not in acts


def test_missing_target_is_never_clicked_blindly(monkeypatch):
    screen = {"focus": "pkg/A", "texts": []}
    acts = []
    _stub(monkeypatch, screen, acts)
    server.run_ui_steps([{"action": "tap_text", "target": "缺失"}], **FAST)
    assert len([a for a in acts if a[0] == "tap_element"]) == 0     # 轮询只看不点


def test_waits_for_delayed_target_then_clicks_once(monkeypatch):
    screen = {"focus": "pkg/A", "texts": []}
    acts = []
    _stub(monkeypatch, screen, acts)
    polls = {"n": 0}

    def delayed_screen(filter_useful=True, max_elements=300):
        polls["n"] += 1
        if polls["n"] == 3:
            screen["texts"] = ["目标按钮"]
        return 1080, 2340, [
            {"text": t, "content_desc": "", "resource_id": ""} for t in screen["texts"]]

    monkeypatch.setattr(server, "_screen_elements", delayed_screen)
    out = server.run_ui_steps(
        [{"action": "tap_text", "target": "目标按钮", "timeout_s": 0.1}],
        default_timeout_s=0.1, poll_ms=1, settle_ms=0)

    assert out["ok"] is True
    assert [a for a in acts if a[0] == "tap_element"] == [("tap_element", "目标按钮")]


# ---------- ok_if_text 条件短路：核对→已达成就收工 / 未达成才继续 ----------

def test_short_circuits_when_condition_already_met(monkeypatch):
    screen = {"focus": "pkg/A", "texts": ["我", "微信号：目标账号"]}   # 现状已经满足
    acts = []
    _stub(monkeypatch, screen, acts)
    out = server.run_ui_steps([
        {"action": "tap_text", "target": "我", "expect": "微信号"},
        {"action": "ok_if_text", "target": "目标账号"},
        {"action": "input", "target": "不该被执行"},
    ], **FAST)
    assert out["ok"] is True
    assert out["short_circuit"] is True and out["stopped_at_step"] == 2
    assert ("input", "不该被执行") not in acts        # 关键：已达成 → 后面的修复流程一步不做


def test_continues_when_condition_not_met(monkeypatch):
    screen = {"focus": "pkg/A", "texts": ["微信号：别的账号"]}      # 现状不满足 → 要往下走
    acts = []
    _stub(monkeypatch, screen, acts)
    out = server.run_ui_steps([
        {"action": "ok_if_text", "target": "目标账号", "timeout_s": 0.02},
        {"action": "input", "target": "该被执行"},
    ], **FAST)
    assert out["ok"] is True and not out.get("short_circuit")
    assert ("input", "该被执行") in acts


def test_expect_supports_variables_for_final_recheck():
    """复核步骤靠 expect={账号}——所以 expect 也必须做变量替换，否则等的是字面量 '{账号}'。"""
    steps = [{"action": "tap_text", "target": "我", "expect": "{账号}"}]
    assert uiscripts.missing_vars(steps, {}) == ["账号"]
    assert uiscripts.resolve(steps, {"账号": "abc"})[0]["expect"] == "abc"


def test_app_lifecycle_actions(monkeypatch):
    screen = {"focus": "pkg/A", "texts": ["通讯录"]}
    acts = []
    _stub(monkeypatch, screen, acts)
    monkeypatch.setattr(server, "stop_app", lambda p: acts.append(("stop", p)))
    monkeypatch.setattr(server, "launch_app", lambda package="", **kwargs: acts.append(("launch", package)) or {"ok": True})
    out = server.run_ui_steps([
        {"action": "stop_app", "target": "com.example"},
        {"action": "launch_app", "target": "com.example", "expect": "通讯录"},
    ], **FAST)
    assert out["ok"] is True
    assert acts == [("stop", "com.example"), ("launch", "com.example")]


def test_press_key_wake_is_idempotent_wake_key(monkeypatch):
    calls = []
    monkeypatch.setattr(server, "_adb", lambda *args: calls.append(args))
    assert server.press_key("wake") == {"ok": True, "key": "wake"}
    assert calls == [("shell", "input", "keyevent", "224")]


def test_launch_app_step_stops_when_launch_fails(monkeypatch):
    screen = {"focus": "pkg/A", "texts": []}
    acts = []
    _stub(monkeypatch, screen, acts)
    monkeypatch.setattr(
        server,
        "launch_app",
        lambda package="", **kwargs: {"ok": False, "reason": "not_installed", "package": package},
    )

    out = server.run_ui_steps([
        {"action": "launch_app", "target": "com.missing"},
        {"action": "input", "target": "不应执行"},
    ], **FAST)

    assert out["ok"] is False
    assert out["failed_at"]["action"] == "launch_app"
    assert out["failed_at"]["action_executed"] is False
    assert "not_installed" in out["failed_at"]["reason"]
    assert ("input", "不应执行") not in acts


def test_rejects_unknown_action(monkeypatch):
    screen = {"focus": "pkg/A", "texts": []}
    acts = []
    _stub(monkeypatch, screen, acts)
    out = server.run_ui_steps([{"action": "自爆", "target": ""}], **FAST)
    assert out["ok"] is False and "未知动作" in out["failed_at"]["reason"]
    assert acts == []


# ---------- uiscripts 读表 / 变量 ----------

def _write_xlsx(path, rows, sheet="步骤脚本"):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["脚本名", "序", "动作", "目标", "期望出现", "超时"])
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def test_reads_and_sorts_steps(tmp_path):
    p = _write_xlsx(tmp_path / "t.xlsx", [
        ["切号", 2, "input", "{账号}", "", 3],
        ["切号", 1, "tap_text", "设置", "退出登录", 5],
        ["别的", 1, "press_key", "back", "", ""],
    ])
    s = uiscripts.read_scripts(str(p))
    assert set(s) == {"切号", "别的"}
    assert [x["action"] for x in s["切号"]] == ["tap_text", "input"]      # 按"序"排好
    assert s["切号"][0]["timeout_s"] == 5.0
    assert "timeout_s" not in s["别的"][0]                               # 超时留空 → 用默认值


def test_missing_sheet_returns_empty(tmp_path):
    p = _write_xlsx(tmp_path / "t2.xlsx", [["切号", 1, "press_key", "back", "", ""]], sheet="用例表")
    assert uiscripts.read_scripts(str(p)) == {}                          # 没有脚本表 → 引擎照常跑


def test_variable_substitution_and_missing_detection():
    steps = [{"action": "input", "target": "{账号}", "expect": ""},
             {"action": "input", "target": "{密码}", "expect": ""}]
    assert uiscripts.missing_vars(steps, {"账号": "a"}) == ["密码"]
    got = uiscripts.resolve(steps, {"账号": "a", "密码": "b"})
    assert [g["target"] for g in got] == ["a", "b"]
    assert steps[0]["target"] == "{账号}"                                # 不改原列表


# ---------- 接线 ----------

def test_registered_in_all():
    names = [getattr(t, "name", "") for t in tools._ALL]
    assert "run_ui_steps" in names and "execute_ui_actions" in names
    assert "clear_input" not in names


def test_tool_lists_available_scripts_when_name_unknown():
    tools.set_ui_scripts({"切号": [{"action": "press_key", "target": "back"}]})
    tools.CONTROL = None
    out = json.loads(asyncio.run(tools.run_ui_steps.handler({"script": "不存在"}))["content"][0]["text"])
    assert out["ok"] is False and out["available"] == ["切号"]
    tools.set_ui_scripts({})


def test_tool_can_list_available_scripts_without_fake_script_name():
    tools.set_ui_scripts({"切号": [{"action": "press_key", "target": "back"}]})
    tools.CONTROL = None

    schema = tools.run_ui_steps.input_schema
    assert schema["type"] == "object"
    assert "script" not in schema.get("required", [])
    out = json.loads(asyncio.run(tools.run_ui_steps.handler({}))["content"][0]["text"])

    assert out == {"ok": True, "listed": True, "available": ["切号"]}
    tools.set_ui_scripts({})


def test_tool_refuses_when_vars_missing():
    tools.set_ui_scripts({"切号": [{"action": "input", "target": "{账号}", "expect": ""}]})
    tools.CONTROL = None
    out = json.loads(asyncio.run(tools.run_ui_steps.handler({"script": "切号"}))["content"][0]["text"])
    assert out["ok"] is False and out["needed"] == ["账号"]              # 缺账号就不跑，别拿 "{账号}" 当字面量输进去
    tools.set_ui_scripts({})


def test_tool_resolves_vars_and_calls_executor(monkeypatch):
    seen = {}
    monkeypatch.setattr(tools.server, "execute_ui_steps",
                        lambda steps, default_timeout_s=5.0, poll_ms=400, settle_ms=500,
                        should_stop=None, **kwargs: (seen.update(
                            steps=steps, dto=default_timeout_s, settle=settle_ms,
                        ), {"ok": True})[1])
    tools.set_ui_scripts({"切号": [{"action": "input", "target": "{账号}", "expect": ""}]})
    tools.CONTROL = None
    out = json.loads(asyncio.run(tools.run_ui_steps.handler(
        {"script": "切号", "vars": {"账号": "wzyyqp"}}))["content"][0]["text"])
    assert out["ok"] is True
    assert seen["steps"][0]["target"] == "wzyyqp"
    assert seen["settle"] == 500                      # 动作前沉降默认值要透传到执行器
    tools.set_ui_scripts({})


def test_settles_before_device_actions_but_not_before_waits(monkeypatch):
    """碰设备的动作前要沉降(治'元素刚出现还不可交互就点→落空')；纯等待/判断步不该白等。"""
    screen = {"focus": "pkg/A", "texts": ["按钮"]}
    acts, naps = [], []
    _stub(monkeypatch, screen, acts)
    monkeypatch.setattr(server.time, "sleep", lambda s: naps.append(round(s, 3)))
    server.run_ui_steps([
        {"action": "tap_text", "target": "按钮"},
        {"action": "wait"},
        {"action": "ok_if_text", "target": "不存在", "timeout_s": 0},
    ], default_timeout_s=0, poll_ms=1, settle_ms=500)
    assert naps.count(0.5) == 1                       # 只有 tap 那一步沉降过


def test_stopped_short_circuits(monkeypatch):
    called = []
    monkeypatch.setattr(tools.server, "execute_ui_steps", lambda *a, **k: called.append(1))
    tools.set_ui_scripts({"切号": [{"action": "press_key", "target": "back"}]})

    class C:
        stopped = True
    tools.CONTROL = C()
    out = json.loads(asyncio.run(tools.run_ui_steps.handler({"script": "切号"}))["content"][0]["text"])
    tools.CONTROL = None
    tools.set_ui_scripts({})
    assert out.get("stopped") is True and called == []


# ---------- 可插拔红线 ----------

FORBIDDEN = ["微信", "wechat", "任务清单", "支付密码", "CheckPwd",
             "FaceFlash", "AppBrand", "135790", "123456", "退出登录", "同意并登录", "paytest"]


def _block(path: str, marker: str) -> str:
    src = open(path, encoding="utf-8").read()
    i = src.index(marker)
    m = re.search(
        r"\n(?:@(mcp\.tool|tool)\(|def [A-Za-z_][A-Za-z0-9_]*\()",
        src[i + len(marker):],
    )
    return src[i:i + len(marker) + (m.start() if m else len(src))]


def test_executor_and_script_reader_have_no_app_knowledge():
    """流程知识必须全在 xlsx——代码里连"退出登录""同意并登录"这种流程文案都不许出现。"""
    blocks = {
        "server.execute_ui_steps": _block(server.__file__, "def execute_ui_steps("),
        "server.clear_input": _block(server.__file__, "def clear_input("),
        "tools.run_ui_steps": _block(tools.__file__, '"run_ui_steps",'),
        "uiscripts(整个模块)": open(uiscripts.__file__, encoding="utf-8").read(),
    }
    for name, block in blocks.items():
        hits = [w for w in FORBIDDEN if w.lower() in block.lower()]
        assert not hits, f"{name} 里出现 app 专属词: {hits}（流程知识应只在 xlsx『步骤脚本』表）"


# ---------- tap_wait_tap 的 tap1 选择器（纯增量，不传则行为不变） ----------

def test_tap_wait_tap_tap1_selector_used_and_reported(monkeypatch):
    """给了 tap1_target 就按 a11y 定位第一下，并回报 tap1_via 供离线区分走的哪条路。"""
    calls = []
    monkeypatch.setattr(server, "get_focused_app", lambda: "pkg/TriggerActivity")
    monkeypatch.setattr(server, "_adb", lambda *a: calls.append(("adb",) + a))
    monkeypatch.setattr(server, "tap_element",
                        lambda resource_id="", text="", index=0, contains=False:
                        (calls.append(("sel", resource_id or text)), {"ok": True})[1])
    out = server.tap_wait_tap(-1, -1, "TriggerActivity", 54, 134,
                              timeout_s=0.05, poll_ms=1, settle_ms=1, tap1_target="下一步")
    assert ("sel", "下一步") in calls          # 用了选择器
    assert out["tap1_via"] == "text"


def test_tap_wait_tap_without_target_still_uses_xy(monkeypatch):
    """不传 tap1_target → 老行为：按坐标点，且如实标 tap1_via=xy。"""
    calls = []
    monkeypatch.setattr(server, "get_focused_app", lambda: "pkg/TriggerActivity")
    monkeypatch.setattr(server, "_adb", lambda *a: calls.append(a))
    monkeypatch.setattr(server, "tap_element", lambda **k: (_ for _ in ()).throw(AssertionError("不该走选择器")))
    out = server.tap_wait_tap(539, 1881, "TriggerActivity", 54, 134,
                              timeout_s=0.05, poll_ms=1, settle_ms=1)
    assert ("shell", "input", "tap", "539", "1881") in calls
    assert out["tap1_via"] == "xy"


def test_tap_wait_tap_selector_miss_stops_without_tapping(monkeypatch):
    """选择器没定位到 → 不点、不进等待阶段，如实回 tap1_ok=false 交回大脑（别硬改坐标点）。"""
    calls = []
    monkeypatch.setattr(server, "get_focused_app", lambda: "pkg/SomePage")
    monkeypatch.setattr(server, "_adb", lambda *a: calls.append(a))
    monkeypatch.setattr(server, "tap_element",
                        lambda resource_id="", text="", index=0, contains=False: {"ok": False, "error": "未找到匹配元素"})
    out = server.tap_wait_tap(-1, -1, "TriggerActivity", 54, 134, tap1_target="下一步")
    assert out["matched"] is False and out["tap1_ok"] is False
    assert not any("tap" in c for c in calls)     # 一次点击都没发


# ---------- 坐标按屏幕百分比表达（换分辨率不废） ----------

def test_pct_point_scales_with_screen(monkeypatch):
    monkeypatch.setattr(server, "_wm_size", lambda: (1080, 2340))
    assert server._point("pct:5,5.73") == (54, 134)      # 本机换算回原坐标
    assert server._point("540,1746") == (540, 1746)      # 绝对像素照旧
    monkeypatch.setattr(server, "_wm_size", lambda: (1440, 3120))
    assert server._point("pct:5,5.73") == (72, 179)      # 换分辨率自动跟随


def test_run_ui_steps_accepts_pct_target(monkeypatch):
    screen = {"focus": "pkg/A", "texts": []}
    acts = []
    _stub(monkeypatch, screen, acts)
    monkeypatch.setattr(server, "_wm_size", lambda: (1080, 2340))
    out = server.run_ui_steps([{"action": "tap_xy", "target": "pct:50,50"}], **FAST)
    assert out["ok"] is True
    assert ("adb", "shell", "input", "tap", "540", "1170") in acts


def test_tap_wait_tap_tap2_pct_resolved_once(monkeypatch):
    """tap2 支持 pct（只读 wm size、不读屏 → 相机页上安全），并回报实际落点。"""
    calls = []
    monkeypatch.setattr(server, "get_focused_app", lambda: "pkg/TriggerActivity")
    monkeypatch.setattr(server, "_adb", lambda *a: calls.append(a))
    monkeypatch.setattr(server, "_wm_size", lambda: (1080, 2340))
    monkeypatch.setattr(server, "_screen_elements",
                        lambda **k: (_ for _ in ()).throw(AssertionError("tap2 阶段不许读屏")))
    out = server.tap_wait_tap(-1, -1, "TriggerActivity", 0, 0,
                              timeout_s=0.05, poll_ms=1, settle_ms=1, tap2_target="pct:5,5.73")
    assert out["tap2_xy"] == [54, 134]
    assert ("shell", "input", "tap", "54", "134") in calls


# ---------- 返回体瘦身：成功只回摘要，失败才带流水 ----------

def test_success_returns_summary_without_step_flow(monkeypatch):
    monkeypatch.setattr(tools.server, "execute_ui_steps",
                        lambda *a, **k: {"ok": True, "done_steps": 22, "total_steps": 22,
                                         "log": [{"i": i} for i in range(22)]})
    tools.set_ui_scripts({"s": [{"action": "press_key", "target": "back"}]})
    tools.CONTROL = None
    txt = asyncio.run(tools.run_ui_steps.handler({"script": "s"}))["content"][0]["text"]
    out = json.loads(txt)
    assert out["ok"] is True and out["done_steps"] == 22
    assert "log" not in out                     # 成功时不把 22 步流水喂给模型
    assert len(txt) < 200                       # 返回体小(此前实测 ≈4070 字符)
    tools.set_ui_scripts({})


def test_failure_still_returns_step_flow(monkeypatch):
    monkeypatch.setattr(tools.server, "execute_ui_steps",
                        lambda *a, **k: {"ok": False, "done_steps": 3, "total_steps": 22,
                                         "failed_at": {"i": 4}, "log": [{"i": i} for i in range(4)]})
    tools.set_ui_scripts({"s": [{"action": "press_key", "target": "back"}]})
    tools.CONTROL = None
    out = json.loads(asyncio.run(tools.run_ui_steps.handler({"script": "s"}))["content"][0]["text"])
    assert out["ok"] is False and len(out["log"]) == 4   # 失败要带流水,大脑靠它判断卡在哪
    tools.set_ui_scripts({})


# ---------- 文字子串匹配（治"wait_for 等到了、tap_element 却点不中"） ----------

def test_tap_element_exact_vs_contains(monkeypatch):
    """tap_element 默认精确匹配(u2 语义)；contains=True 走 textContains。
    真机实证：刷脸协议勾选项只有一整段"您知悉并同意…"，精确匹配"同意"必然落空。"""
    seen = {}

    class _Sel:
        def __init__(self, n): self.count = n
        def __getitem__(self, i): return self
        def click(self): seen["clicked"] = True

    class _Dev:
        def __call__(self, **kw):
            seen["kw"] = kw
            return _Sel(0 if "text" in kw else 1)     # 精确=0 命中，子串=1 命中

    monkeypatch.setattr(server, "_dev", lambda: _Dev())
    assert server.tap_element(text="同意")["ok"] is False           # 精确 → 找不到
    assert "text" in seen["kw"]
    assert server.tap_element(text="同意", contains=True)["ok"] is True
    assert "textContains" in seen["kw"]                             # 走了子串


def test_tap_element_content_desc_is_exact(monkeypatch):
    seen = {}

    class _Sel:
        count = 1
        def __getitem__(self, index): return self
        def click(self): seen["clicked"] = True

    class _Dev:
        def __call__(self, **kwargs):
            seen["kwargs"] = kwargs
            return _Sel()

    monkeypatch.setattr(server, "_dev", lambda: _Dev())
    assert server.tap_element(content_desc="目标入口")["ok"] is True
    assert seen["kwargs"] == {"description": "目标入口"}


def test_contains_prefix_in_step_target(monkeypatch):
    """步骤表里写 '~:片段' 即子串匹配，与 rid:/xy:/pct: 一套写法。"""
    assert server._resolve_target("~:同意") == ("contains", "同意")
    got = {}
    monkeypatch.setattr(server, "get_focused_app", lambda: "pkg/A")
    monkeypatch.setattr(server, "_screen_elements", lambda **k: (
        1080, 2340, [{"text": "您知悉并同意条款", "content_desc": "", "resource_id": ""}]))
    monkeypatch.setattr(server, "tap_element",
                        lambda resource_id="", text="", index=0, contains=False:
                        (got.update(text=text, contains=contains), {"ok": True})[1])
    out = server.run_ui_steps([{"action": "tap_text", "target": "~:同意"}],
                              default_timeout_s=0, poll_ms=1, settle_ms=0)
    assert out["ok"] is True
    assert got == {"text": "同意", "contains": True}


# ---------- tap2 选择器 + 声明式坐标兜底 ----------

def _tw_stub(monkeypatch, calls, els):
    monkeypatch.setattr(server, "get_focused_app", lambda: "pkg/TriggerActivity")
    monkeypatch.setattr(server, "_adb", lambda *a: calls.append(a))
    monkeypatch.setattr(server, "_wm_size", lambda: (1080, 2340))
    monkeypatch.setattr(server, "_screen_elements", lambda **k: (1080, 2340, els))


def test_tap2_selector_resolves_center(monkeypatch):
    calls = []
    _tw_stub(monkeypatch, calls,
             [{"text": "", "content_desc": "取消", "resource_id": "x/g7w", "center": [141, 134]}])
    out = server.tap_wait_tap(-1, -1, "TriggerActivity", 0, 0, timeout_s=0.05, poll_ms=1,
                              settle_ms=1, tap2_target="取消", tap2_fallback="pct:5,5.73")
    assert out["tap2_via"] == "text" and out["tap2_xy"] == [141, 134]
    assert ("shell", "input", "tap", "141", "134") in calls


def test_tap2_falls_back_to_coords_and_says_so(monkeypatch):
    """选择器没命中 → 用声明的兜底坐标，但必须回报 tap2_via=fallback（不静默兜底）。"""
    calls = []
    _tw_stub(monkeypatch, calls, [])                       # 屏上什么都没有
    out = server.tap_wait_tap(-1, -1, "TriggerActivity", 0, 0, timeout_s=0.05, poll_ms=1,
                              settle_ms=1, tap2_target="取消", tap2_fallback="pct:5,5.73")
    assert out["tap2_via"] == "fallback" and out["tap2_xy"] == [54, 134]


def test_tap2_selector_miss_without_fallback_reports_and_stops(monkeypatch):
    """没给兜底又定位不到 → 如实回 tap2_ok=false，一次都不点（别乱点相机页）。"""
    calls = []
    _tw_stub(monkeypatch, calls, [])
    out = server.tap_wait_tap(-1, -1, "TriggerActivity", 0, 0, timeout_s=0.05, poll_ms=1,
                              settle_ms=1, tap2_target="取消")
    assert out["tap2_ok"] is False and out["dismissed"] is False and out["taps"] == 0
    assert not any("input" in c for c in calls)


# ---------- 通用红线：须知里教 AI 传的参数，工具必须认得 ----------

def test_notes_only_use_params_the_tools_actually_expose():
    """扫 xlsx『执行须知』里所有 工具名(参数=…) 写法，逐个参数核对该工具 schema 是否真有。
    这条能拦住整类事故:须知写了个工具不认识的参数 → SDK 静默丢弃 → 行为退化，而测试全绿。"""
    import re
    root = os.path.abspath(__file__)
    for _ in range(5):                       # tests→web→agent→device-mcp→miniapp-pilot
        root = os.path.dirname(root)
    xlsx = os.path.join(root, "examples", "task-list.xlsx")
    if not os.path.exists(xlsx):
        import pytest
        pytest.skip("用例表不在（CI 环境）")
    import openpyxl
    def exposed_params(item):
        schema = getattr(item, "input_schema", {})
        if schema.get("type") == "object" and isinstance(schema.get("properties"), dict):
            return set(schema["properties"])
        return set(schema)

    schemas = {getattr(t, "name", ""): exposed_params(t) for t in tools._ALL}
    ws = openpyxl.load_workbook(xlsx)["执行须知"]
    text = "\n".join(str(r[1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[1])
    bad = []
    for name, argstr in re.findall(r"\b(\w+)\(([^()]*)\)", text):
        if name not in schemas:
            continue                                     # 非工具的括号（如 check_environment(场景)）跳过
        for kw in re.findall(r"(\w+)\s*=", argstr):
            if kw not in schemas[name]:
                bad.append(f"{name}(...{kw}=...) —— 该工具 schema 里没有 {kw}")
    assert not bad, "须知教 AI 传了工具不认识的参数（会被静默丢弃）:\n  " + "\n  ".join(bad)
