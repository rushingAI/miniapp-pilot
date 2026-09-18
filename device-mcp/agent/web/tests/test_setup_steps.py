import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
import suite_data  # noqa: E402
import openpyxl


def _make(path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "v"
    ws.append(["序号", "操作", "观测点", "观测效果", "截图"])
    ws.append(["登录要求", "检查帐号是否为demo_user_alpha，如不是则退出重登，密码：p", "登录成功", None, None])
    ws.append(["测试数据要求", "为演示用户准备测试数据：创建两个未完成任务", "分别成功", None, None])
    ws.append([1, "做事", "看", None, None])
    ws.append(["问题点：需要CI测试数据构造", None, None, None, None])   # 非"要求"的自由说明
    wb.save(path)


def test_read_setup_steps_in_order(tmp_path):
    p = str(tmp_path / "v.xlsx"); _make(p)
    steps = suite_data.read_setup_steps(p, "v")
    assert len(steps) == 2
    assert "退出重登" in steps[0][0]        # 登录要求
    assert "测试数据" in steps[1][0]            # 测试数据要求
    assert steps[0][1] == "登录成功"        # 观测点带上


def test_read_setup_steps_parses_inline_flow_binding(tmp_path):
    p = str(tmp_path / "bound.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "v"
    ws.append(["序号", "操作", "观测点", "观测效果", "截图"])
    ws.append(["帐号要求", "核对测试账号【绑定流程:核对并切号】", "登录成功", None, None])
    wb.save(p)

    steps = suite_data.read_setup_steps(p, "v")

    assert steps == [("核对测试账号", "登录成功", "核对并切号")]


def test_read_setup_steps_absent(tmp_path):
    p = str(tmp_path / "plain.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "s"
    ws.append(["序号", "操作", "观测点", "观测效果", "截图"])
    ws.append([1, "做事", "看", None, None])
    wb.save(p)
    assert suite_data.read_setup_steps(p, "s") == []


def test_read_notes_excludes_requirement_rows(tmp_path):
    p = str(tmp_path / "v.xlsx"); _make(p)
    notes = suite_data.read_notes(p, "v")
    assert "demo_user_alpha" not in notes        # 登录要求不进须知
    assert "创建两个未完成任务" not in notes      # 测试数据要求不进须知
    assert "问题点" in notes                # 非"要求"的说明仍作须知
