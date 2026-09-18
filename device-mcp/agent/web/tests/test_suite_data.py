import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
import suite_data


def test_suite_data_loads_cases_and_persists_completed_verdict(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "用例"
    sheet.append(["序号", "操作", "观测点", "观测效果", "截图"])
    sheet.append([1, "打开入口", "显示首页", "", ""])
    workbook.save(source)

    cases = suite_data.read_cases(str(source), "用例")
    result = suite_data.CaseResult(seq=1, verdict="pass", observed_effect="显示首页")
    output = tmp_path / "result.xlsx"
    summary = suite_data.persist_results(
        str(tmp_path), str(output), str(source), "用例", "", [result], cases,
        final=True,
    )

    assert [case.operation for case in cases] == ["打开入口"]
    assert summary["verdict"] == "pass"
    assert output.exists()


def test_read_cases_preserves_sequence_and_columns(tmp_path):
    source = tmp_path / "sequence.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "用例"
    sheet.append(["序号", "操作", "观测点", "观测效果"])
    for seq in (1, 2, 4, 7, 8):
        sheet.append([seq, f"op{seq}", f"observe{seq}", ""])
    workbook.save(source)

    cases = suite_data.read_cases(str(source), "用例")

    assert [case.seq for case in cases] == [1, 2, 4, 7, 8]
    assert [case.seq for case in cases if case.seq >= 7] == [7, 8]
    assert cases[0].operation == "op1"
    assert cases[0].observe == "observe1"
    assert cases[0].scenario == "用例"


@pytest.mark.parametrize("operation", [
    "【绑定流程:入口】进入目标页",
    "进入【绑定流程：入口】目标页",
    "进入目标页【绑定流程:入口】",
])
def test_read_cases_extracts_inline_flow_binding(tmp_path, operation):
    source = tmp_path / "binding.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "用例"
    sheet.append(["序号", "操作", "观测点", "观测效果"])
    sheet.append([1, operation, "目标页出现", ""])
    workbook.save(source)

    case = suite_data.read_cases(str(source), "用例")[0]

    assert case.required_flow == "入口"
    assert case.operation == "进入目标页"


@pytest.mark.parametrize("operation", [
    "进入目标页【绑定流程:】",
    "进入目标页【绑定流程入口】",
    "【绑定流程:入口A】进入目标页【绑定流程:入口B】",
])
def test_read_cases_rejects_bad_inline_flow_binding(tmp_path, operation):
    source = tmp_path / "bad-binding.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "用例"
    sheet.append(["序号", "操作", "观测点", "观测效果"])
    sheet.append([1, operation, "目标页出现", ""])
    workbook.save(source)

    with pytest.raises(ValueError, match="绑定流程标记"):
        suite_data.read_cases(str(source), "用例")


def test_validate_case_flow_bindings_rejects_unknown_flow():
    case = suite_data.Case(
        1, 2, "用例", "进入目标页", "目标页出现", required_flow="不存在",
    )

    with pytest.raises(ValueError, match="引用不存在的流程"):
        suite_data.validate_case_flow_bindings([case], {"入口": []})


def test_read_notes_collects_rows_and_named_sheet(tmp_path):
    source = tmp_path / "notes.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "用例"
    sheet.append(["序号", "操作", "观测点", "观测效果"])
    sheet.append([1, "op", "observe", ""])
    sheet.append(["", "ROW_NOTE 环境说明", "", ""])
    notes = workbook.create_sheet("执行需知")
    notes.append(["SHEET_NOTE 整表说明"])
    workbook.save(source)

    value = suite_data.read_notes(str(source), "用例")

    assert "ROW_NOTE" in value
    assert "SHEET_NOTE" in value
    assert "op" not in value
