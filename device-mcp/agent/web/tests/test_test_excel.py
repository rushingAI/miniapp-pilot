import asyncio
import json
import os
import sys

import openpyxl
import pytest
from mcp.types import ListToolsRequest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from harness_workspace import HarnessWorkspace  # noqa: E402
from test_excel_mcp import TestExcelService, _error  # noqa: E402
from run_store import RunCoordinator  # noqa: E402
import tools  # noqa: E402


class _Control:
    def __init__(self):
        self.stopped = False
        self.stop_reason = ""
        self.control_lock = asyncio.Lock()


def _book(path, *, setup=False, setup_binding=False, ui_config=False, binding=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "用例"
    ws.append(["序号", "操作", "观测点", "观测效果"])
    if setup:
        setup_operation = (
            "登录测试账号【绑定流程:入口】" if setup_binding else "登录测试账号"
        )
        ws.append(["登录要求", setup_operation, "登录完成", ""])
    operation = "点击 A【绑定流程:入口】" if binding else "点击 A"
    ws.append([1, operation, "显示是否确认？", ""])
    ws.append([2, "点击 B", "显示是否继续阅读？", ""])
    if ui_config:
        scripts = wb.create_sheet("步骤脚本")
        scripts.append(["脚本名", "序", "动作", "目标", "期望出现", "超时"])
        scripts.append(["打开入口", 1, "press_key", "home", "", 3])
        flows = wb.create_sheet("流程映射")
        flows.append(["流程名", "阶段", "序", "匹配", "脚本名", "下一阶段", "失败后继续"])
        flows.append(["入口", "开始", 1, "always", "打开入口", "", "否"])
    wb.save(path)


def _service(tmp_path, *, setup=False, setup_binding=False, ui_config=False, binding=False):
    source = tmp_path / "cases.xlsx"
    _book(source, setup=setup, setup_binding=setup_binding,
          ui_config=ui_config, binding=binding)
    workspace = HarnessWorkspace.create(tmp_path / "attempt")
    imported = workspace.import_file(source)
    return TestExcelService(workspace), imported


def test_record_test_step_exposes_evidence_paths_as_string_array(tmp_path):
    service, _imported = _service(tmp_path)
    server_config, _tool_names = service.server()

    async def exposed_schema():
        server = server_config["instance"]
        response = await server.request_handlers[ListToolsRequest](
            ListToolsRequest(method="tools/list")
        )
        record = next(tool for tool in response.root.tools if tool.name == "record_test_step")
        return record.inputSchema

    schema = asyncio.run(exposed_schema())

    assert schema["properties"]["evidence_paths"] == {
        "type": "array",
        "items": {"type": "string"},
    }


def test_native_suite_loads_and_releases_ui_scripts_and_flows(tmp_path):
    service, imported = _service(tmp_path, ui_config=True)
    tools.set_ui_scripts({})
    tools.set_ui_flows({})

    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))

    assert list(tools.UI_SCRIPTS) == ["打开入口"]
    assert list(tools.UI_FLOWS) == ["入口"]

    asyncio.run(service.finalize_active("测试结束"))
    assert tools.UI_SCRIPTS == {}
    assert tools.UI_FLOWS == {}


def test_open_returns_first_step_and_record_returns_the_next_step(tmp_path):
    service, imported = _service(tmp_path)

    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))

    assert opened["current_step"]["step_id"] == "case-1"
    recorded = asyncio.run(service.record({
        "suite_id": opened["suite_id"], "step_id": "case-1",
        "verdict": "needs_review", "observed_effect": "已执行",
        "evidence_paths": [],
    }))

    assert recorded["recorded_verdict"] == "needs_review"
    assert recorded["current_step"]["step_id"] == "case-2"
    assert recorded["ready_to_finish"] is False


def test_normal_step_handoff_keeps_latest_visual_frame_and_omits_repeated_notes(tmp_path):
    service, imported = _service(tmp_path)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "整套测试必须遵循这段只需读取一次的须知",
    }))
    assert "只需读取一次" in opened["current_step"]["notes"]
    frame = {
        "image_wh": [108, 234], "device_wh": [108, 234],
        "scale": 1.0, "path": str(tmp_path / "latest.png"),
    }
    tools.LAST_SCREENSHOT_META = frame

    recorded = asyncio.run(service.record({
        "suite_id": opened["suite_id"], "step_id": "case-1",
        "verdict": "needs_review", "observed_effect": "已执行",
        "evidence_paths": [],
    }))

    assert "notes" not in recorded["current_step"]
    assert recorded["current_step"]["notes_unchanged"] is True
    assert tools.LAST_SCREENSHOT_META == frame


def test_resume_current_resends_full_suite_notes(tmp_path):
    service, imported = _service(tmp_path)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "恢复新 Turn 时必须重新提供的须知",
    }))

    resumed = asyncio.run(service.next({"suite_id": opened["suite_id"]}))

    assert resumed["status"] == "resume_current"
    assert "恢复新 Turn" in resumed["notes"]


def test_last_record_atomically_completes_suite_and_explicit_finish_is_idempotent(tmp_path):
    emitted = []

    async def emit(event):
        emitted.append(event)

    service, imported = _service(tmp_path)
    service.emit = emit
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 1, "notes": "",
    }))
    evidence = service.suites[opened["suite_id"]].out_dir / "case-1.json"
    evidence.write_text("{}", encoding="utf-8")

    recorded = asyncio.run(service.record({
        "suite_id": opened["suite_id"], "step_id": "case-1",
        "verdict": "pass", "observed_effect": "通过",
        "evidence_paths": [str(evidence)],
    }))

    state = service.suites[opened["suite_id"]]
    assert recorded["suite_completed"] is True
    assert recorded["summary"]["verdict"] == "pass"
    assert state.status == "completed"
    assert service.active_suite_id == ""
    assert [event["type"] for event in emitted[-2:]] == ["suite_done", "download_ready"]
    assert emitted[-2]["suite_id"] == opened["suite_id"]

    repeated = asyncio.run(service.finish({"suite_id": opened["suite_id"]}))
    assert repeated["verdict"] == "pass"
    assert repeated["suite_completed"] is True


def test_stop_winning_control_lock_rejects_record_without_advancing_cursor(tmp_path):
    async def scenario():
        service, imported = _service(tmp_path)
        control = _Control()
        service.control = control
        opened = await service.open({
            "filename": imported["path"], "sheet": "用例", "start_seq": 0,
            "max_cases": 0, "notes": "",
        })
        suite_id = opened["suite_id"]
        step = await service.next({"suite_id": suite_id})
        control.stopped = True
        control.stop_reason = "user_stop"

        result = await service.record({
            "suite_id": suite_id, "step_id": step["step_id"], "verdict": "fail",
            "observed_effect": "不应提交", "evidence_paths": [],
        })

        payload = json.loads(result["content"][0]["text"])
        state = service.suites[suite_id]
        assert payload["stopped"] is True
        assert state.case_index == 0
        assert state.current_step["step_id"] == "case-1"
        assert state.results == []
        await service.finalize_active("测试清理", status="interrupted")

    asyncio.run(scenario())


def test_record_winning_control_lock_commits_atomically_before_stop(tmp_path):
    async def scenario():
        entered_emit = asyncio.Event()
        release_emit = asyncio.Event()

        async def emit(event):
            if event.get("type") == "case_result":
                entered_emit.set()
                await release_emit.wait()

        service, imported = _service(tmp_path)
        control = _Control()
        service.control = control
        service.emit = emit
        opened = await service.open({
            "filename": imported["path"], "sheet": "用例", "start_seq": 0,
            "max_cases": 0, "notes": "",
        })
        suite_id = opened["suite_id"]
        step = await service.next({"suite_id": suite_id})
        record_task = asyncio.create_task(service.record({
            "suite_id": suite_id, "step_id": step["step_id"], "verdict": "fail",
            "observed_effect": "已完整提交", "evidence_paths": [],
        }))
        await entered_emit.wait()

        async def stop_after_lock():
            async with control.control_lock:
                control.stopped = True
                control.stop_reason = "user_stop"

        stop_task = asyncio.create_task(stop_after_lock())
        await asyncio.sleep(0)
        assert not stop_task.done()
        release_emit.set()
        await record_task
        await stop_task

        state = service.suites[suite_id]
        assert state.case_index == 2
        assert state.current_step is None
        assert state.results[0].observed_effect == "已完整提交"
        await service.finalize_active("测试清理", status="interrupted")

    asyncio.run(scenario())


def test_native_bound_case_cannot_record_until_required_flow_succeeds(tmp_path, monkeypatch):
    """Native 路径也必须执行 Excel 绑定，不能绕过 flow 直接记 pass。"""
    service, imported = _service(tmp_path, ui_config=True, binding=True)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    suite_id = opened["suite_id"]
    step = asyncio.run(service.next({"suite_id": suite_id}))
    evidence = service.suites[suite_id].out_dir / "case-1.json"
    evidence.write_text("{}", encoding="utf-8")

    assert step["required_flow"] == "入口"
    with pytest.raises(ValueError, match="route_required"):
        asyncio.run(service.record({
            "suite_id": suite_id, "step_id": step["step_id"], "verdict": "pass",
            "observed_effect": "通过", "evidence_paths": [str(evidence)],
        }))

    monkeypatch.setattr(tools.server, "run_ui_flow", lambda *args, **kwargs: {
        "ok": True, "done_routes": 1, "stage": "开始", "route_log": [],
    })
    flow = json.loads(asyncio.run(tools.run_ui_flow.handler({
        "flow": "入口", "vars": {},
    }))["content"][0]["text"])
    recorded = asyncio.run(service.record({
        "suite_id": suite_id, "step_id": step["step_id"], "verdict": "pass",
        "observed_effect": "通过", "evidence_paths": [str(evidence)],
    }))

    assert flow["ok"] is True
    assert recorded["verdict"] == "pass"
    asyncio.run(service.finalize_active("测试结束"))


def test_native_bound_setup_can_record_after_failed_flow_and_manual_recovery(tmp_path, monkeypatch):
    service, imported = _service(
        tmp_path, setup=True, setup_binding=True, ui_config=True,
    )
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    suite_id = opened["suite_id"]
    step = asyncio.run(service.next({"suite_id": suite_id}))
    evidence = service.suites[suite_id].out_dir / "setup-0.json"
    evidence.write_text("{}", encoding="utf-8")

    assert step["kind"] == "setup"
    assert step["required_flow"] == "入口"
    monkeypatch.setattr(tools.server, "run_ui_flow", lambda *args, **kwargs: {
        "ok": False, "reason": "入口不可达", "stage": "开始", "route_log": [],
    })
    failed = json.loads(asyncio.run(tools.run_ui_flow.handler({
        "flow": "入口", "vars": {},
    }))["content"][0]["text"])
    recorded = asyncio.run(service.record({
        "suite_id": suite_id, "step_id": step["step_id"], "verdict": "pass",
        "observed_effect": "手工恢复后已核验", "evidence_paths": [str(evidence)],
    }))

    assert failed["ok"] is False
    assert recorded["verdict"] == "pass"
    asyncio.run(service.finalize_active("测试结束", status="interrupted"))


def test_test_excel_enforces_order_and_downgrades_pass_without_evidence(tmp_path):
    service, imported = _service(tmp_path)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "按实际截图记录",
    }))
    suite_id = opened["suite_id"]
    first = asyncio.run(service.next({"suite_id": suite_id}))

    assert first["step_id"] == "case-1"
    assert "按实际截图记录" in first["notes"]
    resumed = asyncio.run(service.next({"suite_id": suite_id}))
    assert resumed["status"] == "resume_current"
    assert resumed["step_id"] == first["step_id"]
    assert service.suites[suite_id].case_index == 0

    result = asyncio.run(service.record({
        "suite_id": suite_id, "step_id": "case-1", "verdict": "pass",
        "observed_effect": "截图显示是否确认？", "evidence_paths": [],
    }))
    assert result["verdict"] == "needs_review"
    assert asyncio.run(service.next({"suite_id": suite_id}))["step_id"] == "case-2"


def test_missing_sheet_returns_real_available_sheet_names_without_opening_suite(tmp_path):
    service, imported = _service(tmp_path)

    result = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "", "start_seq": 0,
        "max_cases": 0, "notes": "执行v1",
    }))

    assert result["status"] == "choose_sheet"
    assert result["available_sheets"] == ["用例"]
    assert service.suites == {}


def test_unique_sheet_prefix_is_resolved_without_guessing(tmp_path):
    source = tmp_path / "versions.xlsx"
    _book(source)
    wb = openpyxl.load_workbook(source)
    wb["用例"].title = "v1创建任务上传身份证刷脸通过"
    wb.create_sheet("执行须知")
    wb.save(source)
    workspace = HarnessWorkspace.create(tmp_path / "attempt")
    imported = workspace.import_file(source)
    service = TestExcelService(workspace)

    result = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "v1", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))

    assert result["sheet"] == "v1创建任务上传身份证刷脸通过"


def test_test_excel_setup_failure_blocks_cases_and_persists_workbook(tmp_path):
    service, imported = _service(tmp_path, setup=True)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    suite_id = opened["suite_id"]
    setup = asyncio.run(service.next({"suite_id": suite_id}))
    assert setup["kind"] == "setup"
    asyncio.run(service.record({
        "suite_id": suite_id, "step_id": setup["step_id"], "verdict": "fail",
        "observed_effect": "登录失败", "evidence_paths": [],
    }))
    assert asyncio.run(service.next({"suite_id": suite_id}))["status"] == "complete"
    finished = asyncio.run(service.finish({"suite_id": suite_id}))

    assert finished["summary"]["verdicts"]["blocked"] == 2
    assert finished["verdict"] == "fail"
    assert os.path.isfile(finished["out_xlsx"])


def test_setup_needs_review_continues_and_becomes_suite_verdict(tmp_path):
    service, imported = _service(tmp_path, setup=True, ui_config=True)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    suite_id = opened["suite_id"]
    setup = asyncio.run(service.next({"suite_id": suite_id}))
    recorded = asyncio.run(service.record({
        "suite_id": suite_id, "step_id": setup["step_id"], "verdict": "pass",
        "observed_effect": "登录状态已核对", "evidence_paths": [],
    }))

    assert recorded["verdict"] == "needs_review"
    assert asyncio.run(service.next({"suite_id": suite_id}))["step_id"] == "case-1"

    for seq in (1, 2):
        evidence = service.suites[suite_id].out_dir / f"case-{seq}.json"
        evidence.write_text("{}", encoding="utf-8")
        asyncio.run(service.record({
            "suite_id": suite_id, "step_id": f"case-{seq}", "verdict": "pass",
            "observed_effect": "通过", "evidence_paths": [str(evidence)],
        }))
        terminal = asyncio.run(service.next({"suite_id": suite_id}))

    assert terminal == {"status": "complete", "verdict": "needs_review",
                        "suite_id": suite_id}
    assert tools.UI_FLOWS == {}
    finished = asyncio.run(service.finish({"suite_id": suite_id}))
    assert finished["verdict"] == "needs_review"
    assert finished["summary"]["setup_verdicts"] == ["needs_review"]
    assert tools.UI_SCRIPTS == {}
    assert tools.UI_FLOWS == {}


def test_blocked_case_becomes_finish_only_terminal_contract(tmp_path):
    service, imported = _service(tmp_path)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    suite_id = opened["suite_id"]
    step = asyncio.run(service.next({"suite_id": suite_id}))

    recorded = asyncio.run(service.record({
        "suite_id": suite_id, "step_id": step["step_id"], "verdict": "blocked",
        "observed_effect": "用户确认该阻碍无法解除", "evidence_paths": [],
    }))

    assert recorded["ok"] is True
    assert recorded["verdict"] == "blocked"
    assert recorded["terminal"] is True
    assert recorded["current_step"] is None
    assert recorded["ready_to_finish"] is False
    assert recorded["aggregate_verdict"] == "blocked"
    state = service.suites[suite_id]
    assert state.ready_to_finish is True
    assert state.current_step is None
    assert state.case_index == len(state.cases)
    assert [result.verdict for result in state.results] == ["blocked", "blocked"]
    assert recorded["suite_completed"] is True
    assert state.status == "completed"
    assert service.device_action_block_reason() == ""

    finished = asyncio.run(service.finish({"suite_id": suite_id}))
    assert finished["verdict"] == "blocked"
    assert state.status == "completed"


def test_start_seq_filters_cases_and_skips_setup(tmp_path):
    service, imported = _service(tmp_path, setup=True)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 2,
        "max_cases": 0, "notes": "",
    }))
    step = asyncio.run(service.next({"suite_id": opened["suite_id"]}))

    assert opened["start_seq_skips_setup"] is True
    assert step["step_id"] == "case-2"


def test_evidence_must_be_inside_suite_run(tmp_path):
    service, imported = _service(tmp_path)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 1, "notes": "",
    }))
    step = asyncio.run(service.next({"suite_id": opened["suite_id"]}))
    outside = tmp_path / "outside.png"
    outside.write_bytes(b"x")

    try:
        asyncio.run(service.record({
            "suite_id": opened["suite_id"], "step_id": step["step_id"], "verdict": "pass",
            "observed_effect": "ok", "evidence_paths": [str(outside)],
        }))
    except ValueError as exc:
        assert "当前 suite Run" in str(exc)
    else:
        raise AssertionError("套件外证据必须拒绝")


def test_open_always_starts_new_suite_and_interrupts_previous(tmp_path):
    service, imported = _service(tmp_path)
    first = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))

    reopened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    assert reopened["suite_id"] != first["suite_id"]
    assert service.suites[first["suite_id"]].status == "interrupted"
    assert service.suites[reopened["suite_id"]].status == "active"
    assert service.active_suite_id == reopened["suite_id"]

    try:
        asyncio.run(service.next({"suite_id": first["suite_id"]}))
    except ValueError as exc:
        assert "suite_superseded" in str(exc)
    else:
        raise AssertionError("新测试创建后旧 suite 必须只读")


def test_native_preflight_failure_preserves_previous_active_suite(tmp_path):
    service, imported = _service(tmp_path)
    first = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    bad = tmp_path / "bad.xlsx"
    _book(bad, setup=True, setup_binding=True, ui_config=False)
    bad_imported = service.workspace.import_file(bad)

    with pytest.raises(ValueError, match="绑定流程"):
        asyncio.run(service.open({
            "filename": bad_imported["path"], "sheet": "用例", "start_seq": 0,
            "max_cases": 0, "notes": "",
        }))

    assert service.active_suite_id == first["suite_id"]
    assert service.suites[first["suite_id"]].status == "active"
    asyncio.run(service.finalize_active("测试清理", status="interrupted"))


def test_finish_cannot_skip_remaining_steps_and_explicit_end_preserves_current_step(tmp_path):
    service, imported = _service(tmp_path)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    suite_id = opened["suite_id"]
    asyncio.run(service.next({"suite_id": suite_id}))
    try:
        asyncio.run(service.finish({"suite_id": suite_id}))
    except ValueError as exc:
        assert "尚未记录" in str(exc)
    else:
        raise AssertionError("不能跳过当前步骤直接结束")

    asyncio.run(service.finalize_active("用户明确结束测试"))
    state = service.suites[suite_id]
    assert state.results == []
    assert state.current_step["step_id"] == "case-1"
    assert state.case_index == 0
    assert state.status == "interrupted"
    assert state.finished is True
    summary = json.load(open(state.out_dir / "summary.json", encoding="utf-8"))
    assert "verdict" not in summary


def test_native_summary_only_exposes_verdict_after_suite_is_completed(tmp_path):
    service, imported = _service(tmp_path)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 1, "notes": "",
    }))
    suite_id = opened["suite_id"]
    step = asyncio.run(service.next({"suite_id": suite_id}))
    evidence = service.suites[suite_id].out_dir / "case-1.json"
    evidence.write_text("{}", encoding="utf-8")
    asyncio.run(service.record({
        "suite_id": suite_id, "step_id": step["step_id"], "verdict": "pass",
        "observed_effect": "通过", "evidence_paths": [str(evidence)],
    }))

    complete = asyncio.run(service.next({"suite_id": suite_id}))
    active_summary = json.load(open(
        service.suites[suite_id].out_dir / "summary.json", encoding="utf-8",
    ))
    assert complete["status"] == "complete"
    assert complete["verdict"] == "pass"
    assert active_summary["verdict"] == "pass"

    finished = asyncio.run(service.finish({"suite_id": suite_id}))
    assert finished["summary"]["verdict"] == "pass"


def test_turn_interrupt_keeps_native_suite_active_and_resumable(tmp_path):
    emitted = []

    async def emit(event):
        emitted.append(event)

    service, imported = _service(tmp_path)
    service.emit = emit
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    first = asyncio.run(service.next({"suite_id": opened["suite_id"]}))

    asyncio.run(service.interrupt_turn())
    resumed = asyncio.run(service.next({"suite_id": opened["suite_id"]}))

    state = service.suites[opened["suite_id"]]
    assert state.status == "active"
    assert state.case_index == 0
    assert state.results == []
    assert resumed["status"] == "resume_current"
    assert resumed["step_id"] == first["step_id"]
    interrupted_index = next(
        index for index, event in enumerate(emitted)
        if event["type"] == "turn_interrupted"
    )
    assert interrupted_index < len(emitted) - 1
    assert emitted[-1]["type"] == "step_timing"
    assert emitted[-1]["point"] == "next_returned"


def test_native_step_durations_are_emitted_and_persisted(tmp_path, monkeypatch):
    emitted = []

    async def emit(event):
        emitted.append(event)

    ticks = iter([100.0, 100.0, 112.345, 130.0, 130.0])
    service, imported = _service(tmp_path)
    service.clock = lambda: next(ticks)
    service.emit = emit
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 1, "notes": "",
    }))
    step = asyncio.run(service.next({"suite_id": opened["suite_id"]}))
    asyncio.run(service.record({
        "suite_id": opened["suite_id"], "step_id": step["step_id"], "verdict": "fail",
        "observed_effect": "失败", "evidence_paths": [],
    }))

    result_event = next(event for event in emitted if event["type"] == "case_result")
    assert result_event["duration_s"] == 12.345
    summary = json.load(open(
        service.suites[opened["suite_id"]].out_dir / "summary.json", encoding="utf-8",
    ))
    assert summary["per_case"][0]["duration_s"] == 12
    assert summary["total_duration_s"] == 30.0
    assert summary["total_cost_usd"] is None
    assert summary["tokens"] is None
    assert summary["per_case"][0]["turns"] is None
    assert summary["metrics_source"]["billing"] == "parent_run"


def test_native_emits_step_boundary_timing_with_actual_actor(tmp_path):
    emitted = []

    async def emit(event):
        emitted.append(event)

    service, imported = _service(tmp_path)
    service.emit = emit
    service.note_tool_actor(
        "mcp__test_excel__open_test_suite", "device_operator", "agent-1",
    )
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    suite_id = opened["suite_id"]
    step = opened["current_step"]
    asyncio.run(service.observe_device_action(
        "mcp__device__get_screen", "device_operator", "agent-1",
    ))
    evidence = service.suites[suite_id].out_dir / "case-1.json"
    evidence.write_text("{}", encoding="utf-8")
    service.note_tool_actor(
        "mcp__test_excel__record_test_step", "device_operator", "agent-1",
    )
    asyncio.run(service.record({
        "suite_id": suite_id, "step_id": step["step_id"], "verdict": "pass",
        "observed_effect": "通过", "evidence_paths": [str(evidence)],
    }))

    timing = [event for event in emitted if event["type"] == "step_timing"]
    assert [event["point"] for event in timing] == [
        "next_requested", "next_returned", "first_device_action", "record_finished",
        "next_requested", "next_returned",
    ]
    assert {event["actor"] for event in timing} == {"device_operator"}
    assert {event["agent_id"] for event in timing} == {"agent-1"}
    assert {event["step_id"] for event in timing} == {"case-1", "case-2"}


def test_transport_close_interrupts_suite_without_inventing_verdict(tmp_path):
    service, imported = _service(tmp_path)
    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))
    asyncio.run(service.next({"suite_id": opened["suite_id"]}))

    asyncio.run(service.finalize_active("transport failed"))

    state = service.suites[opened["suite_id"]]
    assert state.results == []
    assert state.current_step["step_id"] == "case-1"
    assert state.status == "interrupted"


def test_test_excel_infrastructure_error_is_real_mcp_error():
    result = _error(RuntimeError("父 Run 租约失效"))

    assert result["is_error"] is True
    assert "停止重试" in result["content"][0]["text"]
    assert "父 Run 租约失效" in result["content"][0]["text"]


def test_test_excel_open_works_with_spawned_borrowed_lease(tmp_path):
    source = tmp_path / "cases.xlsx"
    _book(source)
    workspace = HarnessWorkspace.create(tmp_path / "attempt")
    imported = workspace.import_file(source)
    device = {"status": "ready", "selected_serial": "phone-a", "detail": {}}
    parent_coordinator = RunCoordinator(str(tmp_path / "runs"), lambda event: None, lambda: device)
    parent = parent_coordinator.begin({"source": "interactive"})
    child_coordinator = RunCoordinator(
        str(tmp_path / "runs"), lambda event: None, lambda: device, recover=False,
        borrowed_lease={
            "token": parent.lease.token, "serial": parent.lease.serial,
            "run_id": parent.lease.run_id, "acquired_at": parent.lease.acquired_at,
        },
    )
    service = TestExcelService(
        workspace, coordinator=child_coordinator, parent_run_id=parent.run_id,
    )

    opened = asyncio.run(service.open({
        "filename": imported["path"], "sheet": "用例", "start_seq": 0,
        "max_cases": 0, "notes": "",
    }))

    assert opened["run_id"]
    assert opened["setup_count"] == 0
    asyncio.run(service.finalize_active("测试结束"))
    assert parent_coordinator.leases.active("phone-a") is not None
    parent.finish("completed")
