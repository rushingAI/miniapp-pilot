import importlib.util
from pathlib import Path
import zipfile
import openpyxl
import hashlib
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('source_review', ROOT/'scripts/release_check.py')
review = importlib.util.module_from_spec(spec)
spec.loader.exec_module(review)

def test_secret_scan_reports_location_not_value():
    marker = 'gh' + 'p_' + 'A'*36
    findings = review.inspect_bytes('example.txt', marker.encode())
    assert findings[0]['category'] == 'github-token'
    assert marker not in str(findings)


def test_license_and_public_author_are_in_distribution():
    names = review.inventory()
    assert {'LICENSE', 'NOTICE', 'THIRD_PARTY.md'} <= set(names)
    assert 'LICENSE-REVIEW.md' not in names
    license_text = (ROOT/'LICENSE').read_text()
    assert 'Apache License' in license_text
    assert 'Version 2.0, January 2004' in license_text
    assert 'END OF TERMS AND CONDITIONS' in license_text
    assert 'Copyright 2026 rushingAI' in (ROOT/'NOTICE').read_text()

def test_workbook_has_only_synthetic_visible_sheets():
    path = ROOT/'examples/task-list.xlsx'
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ['任务清单', '执行须知']
    assert all(sheet.sheet_state == 'visible' for sheet in wb)
    assert not wb._external_links
    from suite_data import read_cases, read_setup_steps
    assert len(read_cases(str(path), '任务清单')) == 5
    assert len(read_setup_steps(str(path), '任务清单')) == 1
    with zipfile.ZipFile(path) as archive:
        assert not any('externalLinks' in name for name in archive.namelist())

def test_candidate_inventory_and_bytes_pass_review():
    _, issues = review.check()
    assert not issues


def test_basic_interactions_import_as_five_independent_unexecuted_cases():
    from suite_data import read_cases, read_notes, read_setup_steps
    path = ROOT/'examples/basic-interactions.xlsx'
    cases = read_cases(str(path), '基础交互')
    assert [case.seq for case in cases] == [1, 2, 3, 4, 5]
    assert all(case.operation and case.observe and not case.required_flow for case in cases)
    assert all('起始页面：' in case.operation and '本例准备：' in case.operation for case in cases)
    assert len(read_setup_steps(str(path), '基础交互')) == 1
    assert '待真机验证' in read_notes(str(path), '基础交互')
    assert '预检' in read_setup_steps(str(path), '基础交互')[0][0]
    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ['基础交互', '执行须知']
    assert all(sheet.sheet_state == 'visible' for sheet in workbook)
    sheet = workbook['基础交互']
    assert [sheet.cell(1, column).value for column in (5, 6)] == ['实际结果', '证据']
    assert all(sheet.cell(case.row, column).value is None for case in cases for column in (5, 6))
    assert not workbook._external_links
    assert all(cell.data_type != 'f' for tab in workbook for row in tab for cell in row)
    props = workbook.properties
    assert props.creator in (None, '', 'openpyxl', 'MiniApp Pilot')
    assert props.lastModifiedBy in (None, '', 'MiniApp Pilot')
    assert all(getattr(props, name) in (None, '') for name in ('title', 'subject', 'description', 'keywords', 'category', 'identifier'))
    with zipfile.ZipFile(path) as archive:
        assert not any('externalLinks' in name or 'embeddings' in name or 'vbaProject' in name for name in archive.namelist())
        assert 'docProps/custom.xml' not in archive.namelist()
    assert not review.inspect_bytes(path.name, path.read_bytes())


def test_readable_example_matches_workbook_without_modifying_it():
    source = ROOT/'examples/basic-interactions.xlsx'
    before = source.read_bytes()
    result = subprocess.run([sys.executable, str(ROOT/'scripts/render_example.py')],
                            capture_output=True, text=True, check=True)
    assert result.stdout == (ROOT/'examples/basic-interactions.md').read_text()
    assert source.read_bytes() == before
    workbook = openpyxl.load_workbook(source, read_only=True)
    try:
        from html import escape
        for sheet in workbook:
            assert sheet.title in result.stdout
            for row in sheet.values:
                for value in row:
                    if value is not None:
                        rendered = escape(str(value), quote=False).replace('|', '&#124;').replace('\n', '<br>')
                        assert rendered in result.stdout
    finally:
        workbook.close()


def test_reviewed_video_is_exact_and_binary_changes_fail_closed():
    name = 'docs/media/basic-interactions.mp4'
    raw = (ROOT/name).read_bytes()
    assert hashlib.sha256(raw).hexdigest() == 'aff19d2c4b69328c2f0647c28d0a0ec240a95a784824b32e0772a3d0ea68a2d7'
    assert not review.inspect_bytes(name, raw)
    assert review.inspect_bytes(name, raw+b'changed')
    assert review.inspect_bytes('docs/media/unreviewed.mp4', raw)
